#!/usr/bin/env python3
"""Лист «WB - <месяц>» в утренней книге Ozon: строки — функциями scripts/report_wb_month.py (импорт, не копия), запись — в
уже собранную книгу генератора, вторым листом после «Ozon - <месяц>».

    venv/bin/python3 scripts/book_wb_sheet.py --book data/reports/ozon_2026-09_to_2026-09-24.xlsx --month 2026-09 --date-to 2026-09-24

Что импортируется из report_wb_month: month_days, load_rows_db, load_costs, unit_cost_for, load_ads_db, build_daily, total_row,
COLS, YOUNG_DAYS, SNAP — те же строки и та же раскладка колонок, что у отдельной книги WB (реклама — из таблицы
wb_ad_spend_daily, СС — снимок 1С, дни моложе двух суток — жёлтые). Не импортируется только write_xlsx: он создаёт свою
книгу (openpyxl.Workbook()) и не умеет писать в чужую — поэтому лист пишется здесь тем же списком колонок COLS.
Обращений к WB API — 0: всё из базы. В БД не пишет.
"""
import argparse
import os
import sys
from datetime import date, timedelta
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(ROOT, ".env"))
import report_wb_month as wb_report  # noqa: E402
import loaders.wb_sales_report_loader as wb_loader  # noqa: E402

Z = Decimal(0)


def build_wb_daily(month, date_to, sb=None, snapshot=None, today=None):
    """(строки по дням, итог, заметки, сводка обращений) — функциями report_wb_month; API не зовётся."""
    sb = sb or wb_loader._client()
    today = today or date.today()
    days = wb_report.month_days(month, date_to)
    d1, d2 = days[0], days[-1]
    rows = wb_report.load_rows_db(sb, d1, d2)
    exact, uniform = wb_report.load_costs(sb, snapshot or wb_report.SNAP)
    ads_note, ads_known, ads_by_day = "", False, {}
    try:
        ads_by_day, undated = wb_report.load_ads_db(sb, d1, d2)
        ads_known = True
        ads_note = f"реклама: {wb_report.ads_loader.TABLE} за {d1}…{d2} — списаний {sum(ads_by_day.values(), Z):,.2f} с НДС" + (f", без даты {undated:,.2f}" if undated else "")
    except Exception as error:  # реклама не обязана останавливать лист; пусто — не ноль (как у report_wb_month)
        ads_note = f"реклама не получена: {error}"
    cost_fn = lambda code, size: wb_report.unit_cost_for(exact, uniform, code, size, "base")  # noqa: E731
    outside = {}
    daily = wb_report.build_daily(rows, days, cost_fn, ads_by_day, today, ads_known, outside)
    total = wb_report.total_row(daily)
    young = [r["date"] for r in daily if r.get("young")]
    notes = [f"Источник: отчёт реализации WB (finance-api sales-reports/detailed, period=daily) из таблицы {wb_loader.TABLE}; строк {len(rows)}; "
             "день строки — saleDt в московском времени, без saleDt — rrDate. Строки и колонки — те же функции, что у отдельной книги WB (scripts/report_wb_month.py).",
             f"Комиссия = Σ цена × кВВ строки (commissionPercent), как у владельца; справочно «удержано из выплаты всего» = Оборот − forPay. НДС {total['vat']}.",
             f"Себестоимость: снимок 1С {snapshot or wb_report.SNAP}, стыковка base; без СС {total['no_cost_positions']} из {total['positions']} позиций — оборот {total['no_cost_turnover']:,.2f}.",
             ads_note,
             f"Жёлтым — дни моложе {wb_report.YOUNG_DAYS} суток: строки доезжают." + (f" Таких дней: {', '.join(young)}." if young else "")]
    return daily, total, notes, {"rows": len(rows), "days": len(days), "young": young}


def write_wb_sheet(wb, month, daily, total, notes, position=1):
    """Лист «WB - <месяц>» в книгу wb (openpyxl.Workbook) на позицию position. Колонки — report_wb_month.COLS."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fmt = {"money": "#,##0.00", "pct": "0.0%", "int": "#,##0"}
    bold, young_fill, head_fill = Font(bold=True), PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9E1F2")
    title = f"WB - {wb_report.MONTHS[int(month[5:7]) - 1]}"
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, index=position)
    ws["A1"] = "Выкупы WB"; ws["A1"].font = bold
    ws["E2"] = "Вычитается НДС за возмещение из выручки! (как на листе владельца)"
    cols = wb_report.COLS
    for j, (head, _k, _f) in enumerate(cols, 1):
        if head:
            c = ws.cell(row=3, column=j, value=head); c.font = bold; c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for i, r in enumerate(list(daily) + [total], 4):
        r = dict(r)
        r["note"] = "; ".join(x for x in ("моложе двух суток — строки отчёта доезжают" if r.get("young") else "",
                                          f"без СС {r['no_cost_positions']} из {r['positions']} позиций" if r.get("no_cost_positions") else "",
                                          "реклама не запрашивалась" if r.get("ads") is None else "") if x)
        for j, (head, k, f) in enumerate(cols, 1):
            if not head:
                continue
            v = r.get(k)
            if k == "date" and v != "Итого":
                v = date.fromisoformat(v)
            elif isinstance(v, Decimal):
                v = float(v)
            c = ws.cell(row=i, column=j, value=v)
            if f:
                c.number_format = fmt[f]
            if k == "date" and r["date"] != "Итого":
                c.number_format = "DD.MM.YYYY"
            if r.get("young"):
                c.fill = young_fill
            if r["date"] == "Итого":
                c.font = bold
    base = 4 + len(daily) + 2
    for n, line in enumerate(notes):
        ws.cell(row=base + n, column=1, value=line)
    ws.freeze_panes = "B4"; ws.row_dimensions[3].height = 48
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15 if j > 1 else 16
    ws.column_dimensions[get_column_letter(len(cols))].width = 60
    return ws


def add_wb_sheet(book_path, month, date_to, sb=None, snapshot=None, today=None):
    """Добавить лист WB в книгу на диске. Возвращает (итог, сводка). Книга перезаписывается на месте."""
    import openpyxl
    daily, total, notes, info = build_wb_daily(month, date_to, sb=sb, snapshot=snapshot, today=today)
    wb = openpyxl.load_workbook(book_path)
    write_wb_sheet(wb, month, daily, total, notes)
    wb.save(book_path)
    return total, info


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--book", required=True)
    ap.add_argument("--month", required=True)
    ap.add_argument("--date-to", required=True)
    args = ap.parse_args(argv)
    total, info = add_wb_sheet(args.book, args.month, args.date_to)
    ads = "—" if total["ads"] is None else f"{total['ads']:,.2f}"
    print(f"лист WB добавлен в {args.book}: строк отчёта {info['rows']}, дней {info['days']}; итого оборот {total['turnover']:,.2f}, комиссия {total['commission']:,.2f}, "
          f"логистика {total['logistics']:,.2f}, реклама {ads}, фин. рез. {total['fin_result']:,.2f}; обращений к WB API 0; db_writes = 0")
    return 0


if __name__ == "__main__":
    sys.exit(main())
