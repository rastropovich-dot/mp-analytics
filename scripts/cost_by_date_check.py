#!/usr/bin/env python3
"""Себестоимость по дате продажи — приёмка «было → стало» по месяцам (сорок шестая §3). Только чтение; db_writes = 0.

    venv/bin/python3 scripts/cost_by_date_check.py --snapshots "СС 30.03.26.xlsx=2026-03-30,СС серебро 30.03.26.xlsx=2026-03-30,…" \\
        --dir "~/Downloads/Telegram Desktop/ChatExport_2026-09-29 (2)/files" --book data/reports/finrez_2026-04_2026-09.xlsx \\
        --manual data/manual_report_september_20260922_v2.xlsx [--date-from 2026-04-01 --date-to 2026-09-29]

Было — один снимок 20.05 на все даты (так читатели считали до сорок шестой; в базе он и лежит); стало — правило loaders/unit_cost_history
на всех снимках из файлов (те же, что пойдут в article_unit_costs по слову). Себестоимость = штуки (buyouts_units, иначе позиции) × цена
снимка × индекс СС по дате (как в «Выкупы Ozon»). Фин. рез. «стало» = фин. рез. книги − Δ себестоимости (остальные статьи не зависят от СС).
Против ручного отчёта владельца — его «Итого» листов «Ozon - <месяц>» (колонки: 5 — СС, 15 — фин. рез.).
"""
import argparse
import os
import sys
from collections import defaultdict
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import load_article_unit_costs as L  # noqa: E402
import report_finrez as fr  # noqa: E402
import report_ozon_month as rep  # noqa: E402
from loaders import unit_cost_history as uch  # noqa: E402

D, Z = Decimal, Decimal(0)
RU = {"04": "апрель", "05": "май", "06": "июнь", "07": "июль", "08": "август", "09": "сентябрь"}
LAB = {"04": "апр", "05": "май", "06": "июн", "07": "июл", "08": "авг", "09": "сен"}


def histories(directory, specs):
    """(было: только последний снимок, стало: все) из файлов; несколько файлов одной даты сливаются (серебро 30.03 к основному 30.03)."""
    snaps = defaultdict(dict)
    for name, day in specs:
        rows, _conf, st = L.read_snapshot(os.path.join(directory, name), day, "ozon")
        for r in rows:
            snaps[day][r["offer_id_norm"]] = D(r["unit_cost"])
        print(f"  снимок {day} ← {name}: ключей {len(rows)} (main {st['by_variant_seen'].get('main', 0)}), в снимке теперь {len(snaps[day])}")
    last = max(snaps)
    return uch.CostHistory({last: snaps[last]}), uch.CostHistory(dict(snaps))


def book_by_month(path):
    """«Выкупы Ozon» книги: {мм: (себестоимость, фин. рез.)} по строкам месяцев."""
    import openpyxl
    import warnings
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Выкупы Ozon"].iter_rows(min_row=1, max_row=200, values_only=True)); wb.close()
    ix = {str(v).strip(): j for j, v in enumerate(rows[2]) if v}
    inv = {v: k for k, v in LAB.items()}
    out = {}
    for r in rows[3:]:
        if r and r[0] in inv:
            out[inv[r[0]]] = (D(str(r[ix["Себестоимость, руб."]])), D(str(r[ix["Фин. рез., руб."]])))
    return out


def manual_by_month(path):
    """Ручной отчёт владельца: {мм: (СС, фин. рез.)} из строки «Итого» листов «Ozon - <месяц>»."""
    import openpyxl
    import warnings
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for mm, name in RU.items():
        sheet = f"Ozon - {name}"
        if sheet not in wb.sheetnames:
            continue
        for r in wb[sheet].iter_rows(values_only=True):
            if r and isinstance(r[0], str) and r[0].strip().lower().startswith("итог"):
                out[mm] = (D(str(r[5])), D(str(r[15])))
                break
    wb.close()
    return out


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    ap.add_argument("--snapshots", required=True, help="имя=дата через запятую")
    ap.add_argument("--book", default=os.path.join(ROOT, "data", "reports", "finrez_2026-04_2026-09.xlsx"))
    ap.add_argument("--manual", default=os.path.join(ROOT, "data", "manual_report_september_20260922_v2.xlsx"))
    ap.add_argument("--date-from", default="2026-04-01")
    ap.add_argument("--date-to", default="2026-09-29")
    args = ap.parse_args(argv)
    specs = [tuple(x.rsplit("=", 1)) for x in args.snapshots.split(",")]
    print("снимки из файлов:")
    old_h, new_h = histories(os.path.expanduser(args.dir), specs)
    print(f"было: {old_h.dates}; стало: {new_h.dates}")

    from dotenv import load_dotenv
    from supabase import create_client
    load_dotenv(os.path.join(ROOT, ".env"))
    sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])
    sku2art, _uc, _f, _n, _orders = rep.load_costs(sb, rep.SNAP)
    buyouts = rep.fetch(sb, "marketplace_buyouts", "id,buyout_date,marketplace_sku,buyouts_qty,buyouts_units",
                        [("eq", "marketplace_code", "ozon"), ("gte", "buyout_date", args.date_from), ("lte", "buyout_date", args.date_to)],
                        ["buyout_date", "marketplace_sku"])
    print(f"выкупов Ozon {args.date_from} … {args.date_to}: строк {len(buyouts)}; sku → артикул {len(sku2art)}")
    old_fn, new_fn = uch.unit_cost_fn(old_h, sku2art), uch.unit_cost_fn(new_h, sku2art)
    cogs_old, cogs_new, no_cost = defaultdict(Decimal), defaultdict(Decimal), defaultdict(int)
    later_rub, later_rows = defaultdict(Decimal), defaultdict(int)
    for r in buyouts:
        d, sku = str(r["buyout_date"]), str(r["marketplace_sku"] or "")
        mm = d[5:7]
        units = D(str(r["buyouts_units"])) if r.get("buyouts_units") is not None else D(str(r["buyouts_qty"]))
        idx = fr.cogs_index(d)
        co = old_fn(sku, d, units)
        art = sku2art.get(sku)
        cn, snap, kind = new_h.lookup(art.lower(), d) if art else (None, None, "none")
        new_h.counters[kind] += 1
        if cn is not None:
            new_h.rubles[kind] += cn * units
        if co is None or cn is None:
            no_cost[mm] += int(D(str(r["buyouts_qty"])))
            continue
        cogs_old[mm] += units * co * idx
        cogs_new[mm] += units * cn * idx
        if kind == "later":
            later_rub[mm] += units * cn * idx; later_rows[mm] += 1
    book, manual = (book_by_month(args.book) if os.path.exists(args.book) else {}), (manual_by_month(args.manual) if os.path.exists(args.manual) else {})
    print("\nмесяц    | СС было (20.05)   | СС стало (по дате) | Δ СС           | Δ %    | фин.рез книги | фин.рез стало  | СС владельца   | расх. СС было → стало       | фин.рез владельца | расх. фин.рез было → стало | «более поздний» строк / ₽")
    tot = defaultdict(Decimal)
    for mm in sorted(cogs_old):
        o, n = cogs_old[mm], cogs_new[mm]
        b_cogs, b_fin = book.get(mm, (None, None))
        m_cogs, m_fin = manual.get(mm, (None, None))
        fin_new = (b_fin - (n - o)) if b_fin is not None else None
        line = f"{RU[mm]:8} | {o:>17,.2f} | {n:>18,.2f} | {n - o:>14,.2f} | {(n / o - 1) * 100 if o else 0:>+5.1f}% | "
        line += (f"{b_fin:>13,.2f} | {fin_new:>14,.2f} | " if b_fin is not None else f"{'—':>13} | {'—':>14} | ")
        line += (f"{m_cogs:>14,.2f} | {o - m_cogs:>13,.2f} → {n - m_cogs:>13,.2f} | " if m_cogs is not None else f"{'—':>14} | {'—':>29} | ")
        line += (f"{m_fin:>17,.2f} | {b_fin - m_fin:>12,.2f} → {fin_new - m_fin:>12,.2f} | " if (m_fin is not None and b_fin is not None) else f"{'—':>17} | {'—':>27} | ")
        line += f"{later_rows[mm]} / {later_rub[mm]:,.2f}"
        print(line)
        for k, v in (("o", o), ("n", n)):
            tot[k] += v
        if b_cogs is not None and abs(b_cogs - o) > D("0.5"):
            print(f"         ⚠ книга: себестоимость {b_cogs:,.2f} ≠ «было» {o:,.2f} (разница {b_cogs - o:,.2f}) — книга старее окна или другой снимок")
    print(f"итого    | {tot['o']:>17,.2f} | {tot['n']:>18,.2f} | {tot['n'] - tot['o']:>14,.2f} | {(tot['n'] / tot['o'] - 1) * 100 if tot['o'] else 0:>+5.1f}%")
    print("позиций без СС по месяцам:", dict(no_cost))
    print("подбор снимка (стало):", new_h.note())
    print("db_writes = 0; обращений к API — 0")


if __name__ == "__main__":
    main()
