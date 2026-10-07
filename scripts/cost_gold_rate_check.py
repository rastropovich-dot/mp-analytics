#!/usr/bin/env python3
"""Курс золота 1С и себестоимость (сорок седьмая §2–§3): связь составляющих СС в снимках 1С с курсом 1С и оценка СС выкупов Ozon
за месяцы без снимка по истории курса. Только чтение; db_writes = 0; обращений к Ozon / WB — 0.

    venv/bin/python3 scripts/cost_gold_rate_check.py \\
        --dir "~/Downloads/Telegram Desktop/ChatExport_2026-09-29 (2)/files" \\
        --snapshots "СС 30.03.26.xlsx=2026-03-30,СС серебро 30.03.26.xlsx=2026-03-30,СС 06.04.26.xlsx=2026-04-06,СС 20.04.26.xlsx=2026-04-20,СС 29.04.26.xlsx=2026-04-29,СС 12.05.26.xlsx=2026-05-12,СС 20.05.26.xlsx=2026-05-20" \\
        [--rates data/ll_rates/ll_rates_1c_history.csv] [--book data/reports/finrez_2026-04_2026-10.xlsx] [--manual data/manual_report_september_20260922_v2.xlsx] \\
        [--date-from 2026-06-01 --date-to 2026-10-05] [--no-db]

Что проверяет (§2). В файлах 1С себестоимость разложена: СС_СтоимостьВставок + СС_СтоимостьМеталла + СС_СтоимостьПотерь +
СС_СтоимостьРаботы [− СС_ЭкономияНаВставках]. Скрипт считает: (а) у какой доли строк итог сходится с суммой четырёх / с вычетом
экономии; (б) металл = вес × курс 1С металла (585 — курс золота 585 из истории чата «LL Курсы», 925 — курс серебра) — точно ли и с
каким курсом; (в) k_metal = (металл + потери) / СС по пробе и виду изделия — на 1 % курса 1С СС меняется на k %; (г) по парам соседних
снимков — предсказание СС следующего снимка по k категории × Δкурса (и построчно: металл+потери × Δзолота + вставки × Δдоллара 1С)
против факта.

Что считает (§3). СС выкупов Ozon помесячно: «было» — себестоимость по правилу читателей (снимок ≤ дате, после последнего —
последний; из article_unit_costs), «стало» — СС последнего снимка × (1 + k_категории × Δкурса 1С с даты снимка до дня продажи) и
построчный вариант с долларом; против книги «Фин рез» (фин. рез. «стало» = фин. рез. книги + (СС книги − СС стало); в сентябре СС
книги уже × индекс 1,150) и против ручного отчёта владельца (сентябрь — только 1–21, как в его листе). Индекс 1,150 — против роста
курса 1С с 20.05.

История курса 1С — из `scripts/ll_rates_ocr.py` (data/ll_rates/ll_rates_1c_history.csv): дата установки, значение. Дата снимка 1С
может совпасть с датой смены курса (20.05: файл сделан уже на новом курсе 6 072,85, картинка за 20.05 отсутствует) — курс на
снимок берётся из самого файла: медиана металл / вес по строкам 585 и 925, доллар — по Δвставок к предыдущему снимку.
"""
import argparse
import bisect
import csv
import os
import sys
from collections import defaultdict
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import cabinet  # noqa: E402

_PROFILE = cabinet.profile()   # каталоги кабинета; guard assert_env — перед созданием клиента в main

D, Z = Decimal, Decimal(0)
COMP = ("c_ins", "c_metal", "c_loss", "c_work", "c_save")
WANT = {"kind": ("ВидИзделия",), "article": ("Артикул",), "size": ("Размер",), "weight": ("Вес",),
        "offer": ("Идентификатор Ozon", "ОЗОН основной"), "ins_cat": ("КатегорияКамней",), "fineness": ("Проба",),
        "cost": ("Себестоимость", "Себестоимость Озон и ЛК ВБ"),
        "c_ins": ("СС_СтоимостьВставок",), "c_metal": ("СС_СтоимостьМеталла",), "c_loss": ("СС_СтоимостьПотерь",),
        "c_work": ("СС_СтоимостьРаботы",), "c_save": ("СС_ЭкономияНаВставках",)}
RU = {"06": "июнь", "07": "июль", "08": "август", "09": "сентябрь", "10": "октябрь", "04": "апрель", "05": "май"}


def _norm(s):
    return " ".join(str(s).split()).casefold() if s is not None else ""


def read_components(path, snap):
    """Строки файла 1С с составляющими; шапка в одной или двух строках (СС_* во второй — файлы до апреля)."""
    import openpyxl
    import warnings
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h1, h2 = next(it), next(it)
    hdr = {}
    for j, (a, b) in enumerate(zip(h1, h2)):
        if a is not None:
            hdr.setdefault(_norm(a), j)
        if b is not None and _norm(b).startswith("сс_"):
            hdr.setdefault(_norm(b), j)
    col = {k: next((hdr[_norm(a)] for a in als if _norm(a) in hdr), None) for k, als in WANT.items()}
    missing = [k for k, v in col.items() if v is None and k != "c_save"]
    if missing:
        raise RuntimeError(f"{os.path.basename(path)}: в шапке нет {missing}")
    two_header = any(_norm(x).startswith("сс_") for x in h2 if x)
    rows = ([] if two_header else [h2]) + list(it)
    out = []
    for r in rows:
        if r is None or all(v is None for v in r):
            continue
        rec = {k: (r[col[k]] if col[k] is not None and col[k] < len(r) else None) for k in WANT}
        for k in ("cost", "weight") + COMP:
            v = rec[k]
            rec[k] = D(str(v)) if v not in (None, "") else Z
        rec["fineness"] = str(rec["fineness"] or "").strip()
        rec["okey"] = str(rec["offer"]).strip().lower() if rec["offer"] else None
        rec["snap"] = snap
        out.append(rec)
    wb.close()
    return out


def median(xs):
    xs = sorted(xs)
    return xs[len(xs) // 2] if xs else None


def quart(xs):
    xs = sorted(xs)
    return (xs[len(xs) // 4], xs[len(xs) // 2], xs[3 * len(xs) // 4]) if xs else (None, None, None)


def load_rates(path):
    """{rate: ([даты установки], [значения])} — при дребезге одной даты берётся последняя виденная запись."""
    by = defaultdict(dict)
    with open(path) as fh:
        for r in csv.DictReader(fh):
            by[r["rate"]][r["set_date"]] = D(r["value_1c"])
    return {k: (sorted(v), [v[d] for d in sorted(v)]) for k, v in by.items()}


def rate_on(series, day):
    ds, vs = series
    i = bisect.bisect_right(ds, day) - 1
    return vs[i] if i >= 0 else None


def k_profile_text(kcat, kfin, by_cat, by_fin, snap, digits, min_rows=0):
    """COST_RATE_K для cabinets/<код>.py: по пробе — «*» (медиана пробы) и каждый вид изделия снимка (медиана вида, n строк в комментарии).
    Проба → металл курса: 585 и 375 — золото, 925 — серебро (loaders/unit_cost_history.METAL_OF_FINENESS)."""
    fmt = (lambda v: f"Decimal('{v}')")
    lines = [f"# k по пробам и видам изделия: scripts/cost_gold_rate_check.py --k-profile --k-digits {digits} --k-min-rows {min_rows} (снимок {snap};",
             "# k = (металл + потери) / СС, медиана). Не править руками — пересчитывать скриптом по новому снимку 1С. «*» — медиана пробы: запасной k",
             f"# для вида, которого в снимке не было или у которого строк меньше {min_rows}.",
             "COST_RATE_K = {"]
    for fin in sorted(by_fin):
        lines.append(f"    \"{fin}\": {{  # строк {len(by_fin[fin])}")
        lines.append(f"        \"*\": {fmt(kfin[fin])},")
        for (f, kind), v in sorted(by_cat.items(), key=lambda kv: (-len(kv[1]), kv[0][1])):
            if f == fin and (f, kind) in kcat:
                lines.append(f"        \"{kind}\": {fmt(kcat[(f, kind)])},  # n={len(v)}")
        lines.append("    },")
    lines.append("}")
    return "\n".join(lines)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", required=True)
    ap.add_argument("--snapshots", required=True, help="имя=дата через запятую")
    ap.add_argument("--rates", default=cabinet.data_path("ll_rates", "ll_rates_1c_history.csv", prof=_PROFILE))
    ap.add_argument("--book", default=cabinet.data_path("reports", "finrez_2026-04_2026-10.xlsx", prof=_PROFILE))
    ap.add_argument("--manual", default=cabinet.data_path("manual_report_september_20260922_v2.xlsx", prof=_PROFILE))
    ap.add_argument("--date-from", default="2026-06-01")
    ap.add_argument("--date-to", default="2026-10-05")
    ap.add_argument("--no-db", action="store_true", help="только §2 (без базы)")
    ap.add_argument("--k-profile", action="store_true", help="напечатать k по пробам и видам изделия последнего снимка как COST_RATE_K для профиля кабинета (сорок восьмая §2)")
    ap.add_argument("--k-digits", type=int, default=None, help="округлить k до N знаков и в §3 тоже (как в профиле); без флага — k без округления, как в сорок седьмой")
    ap.add_argument("--k-min-rows", type=int, default=0, help="вид изделия с меньшим числом строк в снимке k не получает (берётся «*» пробы) — и в профиле, и в §3; 0 — все виды, как в сорок седьмой")
    args = ap.parse_args(argv)
    directory = os.path.expanduser(args.dir)
    specs = [tuple(x.rsplit("=", 1)) for x in args.snapshots.split(",")]
    rates = load_rates(args.rates)
    G, AG, U = rates["gold"], rates["silver"], rates["usd"]

    # ---------------- §2
    rows = []
    for name, day in specs:
        part = read_components(os.path.join(directory, name), day)
        rows += part
        print(f"снимок {day} ← {name}: строк {len(part)}")
    snaps = sorted({r["snap"] for r in rows})
    tol = D("0.011")
    both = only4 = only5 = none = 0
    for r in rows:
        s4 = r["c_ins"] + r["c_metal"] + r["c_loss"] + r["c_work"]
        ok4, ok5 = abs(s4 - r["cost"]) <= tol, abs(s4 - r["c_save"] - r["cost"]) <= tol
        both += ok4 and ok5; only4 += ok4 and not ok5; only5 += ok5 and not ok4; none += not ok4 and not ok5
    print(f"\n§2а сходимость составляющих ({len(rows)} строк): вставки+металл+потери+работа = СС и экономия 0 — {both}; "
          f"= СС без вычета экономии (экономия ≠ 0) — {only4}; = СС только с вычетом экономии — {only5}; не сходится — {none}")

    # металл = вес × курс: курс снимка из файла
    snap_rate = {}
    for s in snaps:
        for fin, key in (("585", "gold"), ("925", "silver"), ("375", "375")):
            xs = [r["c_metal"] / r["weight"] for r in rows if r["snap"] == s and r["fineness"] == fin and r["weight"] > 0 and r["c_metal"] > 0]
            if not xs:
                continue
            q1, med, q3 = quart(xs)
            snap_rate[(s, fin)] = med
            chat = rate_on(G if fin == "585" else AG, s) if fin != "375" else None
            exact = sum(1 for r in rows if r["snap"] == s and r["fineness"] == fin and r["weight"] > 0 and abs(r["c_metal"] - r["weight"] * med) <= tol)
            print(f"§2б {s} проба {fin}: металл / вес — медиана {med:.2f} [кв. {q1:.2f}; {q3:.2f}], |металл − вес × медиана| ≤ 0,01 у {exact} из {len(xs)}"
                  + (f"; курс 1С по картинкам на {s}: {chat}" + (" = файлу" if chat is not None and abs(chat - med) < D('0.01') else " ≠ файлу — файл сделан на другом курсе") if chat is not None else ""))

    # k по категориям на последнем снимке
    last = snaps[-1]
    by_cat = defaultdict(list)
    by_fin = defaultdict(list)
    for r in rows:
        if r["snap"] != last or r["cost"] <= 0:
            continue
        k = (r["c_metal"] + r["c_loss"]) / r["cost"]
        by_cat[(r["fineness"], r["kind"])].append((k, r["c_ins"] / r["cost"], r["c_work"] / r["cost"]))
        by_fin[r["fineness"]].append(k)
    kcat = {key: median([x[0] for x in v]) for key, v in by_cat.items()}
    kfin = {f: median(v) for f, v in by_fin.items()}
    if args.k_min_rows:
        kcat = {key: v for key, v in kcat.items() if len(by_cat[key]) >= args.k_min_rows}
    if args.k_digits is not None:
        qk = D(1).scaleb(-args.k_digits)
        kcat = {key: v.quantize(qk) for key, v in kcat.items()}
        kfin = {f: v.quantize(qk) for f, v in kfin.items()}
    if args.k_profile:
        print(k_profile_text(kcat, kfin, by_cat, by_fin, last, args.k_digits, args.k_min_rows))
    print(f"\n§2в k_metal = (металл + потери) / СС, снимок {last}: на 1 % курса 1С металла СС меняется на k %")
    for fin in sorted(by_fin):
        q1, med, q3 = quart(by_fin[fin])
        print(f"  проба {fin}: строк {len(by_fin[fin])}, k медиана {med:.3f} [кв. {q1:.3f}; {q3:.3f}]")
        for (f, kind), v in sorted(by_cat.items(), key=lambda kv: -len(kv[1])):
            if f != fin or len(v) < 50:
                continue
            q1, med, q3 = quart([x[0] for x in v])
            print(f"    {kind:20} n={len(v):6}  k {med:.3f} [кв. {q1:.3f}; {q3:.3f}]  вставки {median([x[1] for x in v]):.3f}  работа {median([x[2] for x in v]):.3f}")

    # пары снимков: предсказание
    print("\n§2г пары соседних снимков (проба 585; ключ — offer Ozon, вес не менялся): предсказание СС_b из СС_a")
    by_snap = {s: {r["okey"]: r for r in rows if r["snap"] == s and r["okey"] and r["fineness"] == "585" and r["cost"] > 0} for s in snaps}
    usd_snap = {}
    for a, b in zip(snaps, snaps[1:]):
        A, B = by_snap[a], by_snap[b]
        common = [k for k in A if k in B and abs(A[k]["weight"] - B[k]["weight"]) < D("0.0005")]
        dg = snap_rate[(b, "585")] / snap_rate[(a, "585")] - 1
        ins = [B[k]["c_ins"] / A[k]["c_ins"] - 1 for k in common if A[k]["c_ins"] > 0]
        du = median(ins) if ins else Z
        usd_snap[(a, b)] = du
        kc = {key: v for key, v in kcat.items()}
        e_cat, e_row = [], []
        sp_c = sp_r = sf = Z
        for k in common:
            ra, rb = A[k], B[k]
            kk = kc.get(("585", ra["kind"]), kfin["585"])
            pc = ra["cost"] * (1 + kk * dg)
            pr = ra["cost"] + (ra["c_metal"] + ra["c_loss"]) * dg + ra["c_ins"] * du
            e_cat.append((pc - rb["cost"]) / rb["cost"] * 100); e_row.append((pr - rb["cost"]) / rb["cost"] * 100)
            sp_c += pc; sp_r += pr; sf += rb["cost"]
        for name, e, sp in (("k категории × Δзолота", e_cat, sp_c), ("построчно: металл+потери × Δзолота + вставки × Δдоллара", e_row, sp_r)):
            within = sum(1 for x in e if abs(x) <= D("0.5")) / len(e) * 100 if e else 0
            print(f"  {a} → {b}: Δзолото 1С {dg * 100:+.2f} %, Δвставок (≈ Δдоллара 1С) {du * 100:+.2f} %; {name}: ошибка медиана {median(e):+.3f} %, "
                  f"|ошибка| ≤ 0,5 % у {within:.1f} % ключей (n={len(e)}), Σ пред / Σ факт {sp / sf:.4f}")

    if args.no_db:
        print("\n--no-db: §3 пропущен; db_writes = 0")
        return

    # ---------------- §3
    from dotenv import load_dotenv
    from supabase import create_client
    import report_ozon_month as rep
    from loaders import unit_cost_history as uch
    import cost_by_date_check as cbc
    load_dotenv(os.path.join(ROOT, ".env"))
    cabinet.assert_env()  # кабинет (MP_CABINET) и база (SUPABASE_URL) должны совпасть — до чтения ключей и создания клиента
    sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])
    sku2art, _uc, _f, _n, _orders = rep.load_costs(sb, rep.SNAP, rates=False)   # здесь курс считается своими руками по csv, читатель — без поправки
    buyouts = rep.fetch(sb, "marketplace_buyouts", "id,buyout_date,marketplace_sku,buyouts_qty,buyouts_units",
                        [("eq", "marketplace_code", "ozon"), ("gte", "buyout_date", args.date_from), ("lte", "buyout_date", args.date_to)],
                        ["buyout_date", "marketplace_sku"])
    norms = sorted({a.lower() for a in sku2art.values() if a})
    hist = uch.load_history(sb, norms)
    print(f"\n§3 выкупов Ozon {args.date_from} … {args.date_to}: строк {len(buyouts)}; sku → артикул {len(sku2art)}; снимков в базе {hist.dates}")
    comp = by_snap_all = {r["okey"]: r for r in rows if r["snap"] == last and r["okey"] and r["cost"] > 0}
    g0, ag0 = snap_rate[(last, "585")], snap_rate.get((last, "925"))
    # доллар 1С на снимок: из истории чата на дату снимка; 20.05 — 75,00 (Δвставок к 12.05 = −6,25 % = 80 → 75; картинка 18.06 показала правку на 74,00 задним числом)
    u0 = rate_on(U, last)
    du_last = usd_snap.get((snaps[-2], last)) if len(snaps) > 1 else None
    if du_last is not None:
        u_prev = rate_on(U, snaps[-2])
        if u_prev:
            u0 = (u_prev * (1 + du_last)).quantize(D("0.01"))
    print(f"   курсы 1С на снимок {last}: золото {g0:.2f} (из файла), серебро {ag0}, доллар {u0} (по Δвставок)")
    for day in ("2026-06-01", "2026-07-01", "2026-08-01", "2026-09-01", "2026-10-06"):
        g, a, u = rate_on(G, day), rate_on(AG, day), rate_on(U, day)
        print(f"   {day}: золото {g} ({(g / g0 - 1) * 100:+.2f} %), серебро {a} ({(a / ag0 - 1) * 100:+.2f} %), доллар {u} ({(u / u0 - 1) * 100:+.2f} %)")
    mon = defaultdict(lambda: defaultdict(Decimal))
    cnt = defaultdict(lambda: defaultdict(int))
    for r in buyouts:
        day = str(r["buyout_date"]); mm = day[:7]
        periods = [mm] + (["2026-09 (1–21)"] if "2026-09-01" <= day <= "2026-09-21" else [])
        units = D(str(r["buyouts_units"])) if r.get("buyouts_units") is not None else D(str(r["buyouts_qty"]))
        art = sku2art.get(str(r["marketplace_sku"]))
        for P in periods:
            cnt[P]["rows"] += 1
        if not art:
            for P in periods:
                cnt[P]["no_art"] += 1
            continue
        c, snap, kind = hist.lookup(art.lower(), day)
        if c is None:
            for P in periods:
                cnt[P]["no_cost"] += 1
            continue
        gd, agd, ud = rate_on(G, day), rate_on(AG, day), rate_on(U, day)
        rc = comp.get(art.lower())
        if rc is not None:
            dm = (agd / ag0 - 1) if rc["fineness"] == "925" else (gd / g0 - 1)
            kk = kcat.get((rc["fineness"], rc["kind"]), kfin.get(rc["fineness"], kfin["585"]))
            new_cat = (rc["cost"] * (1 + kk * dm)).quantize(D("0.01"))
            new_row = (rc["cost"] + (rc["c_metal"] + rc["c_loss"]) * dm + rc["c_ins"] * (ud / u0 - 1)).quantize(D("0.01"))
            tag = "with_comp"
        else:
            dm = gd / g0 - 1
            new_cat = new_row = (c * (1 + kfin["585"] * dm)).quantize(D("0.01"))
            tag = "no_comp"
        for P in periods:
            mon[P]["was"] += units * c; mon[P]["cat"] += units * new_cat; mon[P]["row"] += units * new_row; cnt[P][tag] += 1
    book = cbc.book_by_month(os.path.expanduser(args.book)) if os.path.exists(os.path.expanduser(args.book)) else {}
    manual = cbc.manual_by_month(os.path.expanduser(args.manual)) if os.path.exists(os.path.expanduser(args.manual)) else {}
    print("\nпериод | строк | СС было (правило читателей) | СС стало (k кат. × Δзолота) | Δ % | СС стало (построчно +доллар) | Δ % | фин.рез книги → стало | СС владельца: расх. было → стало | фин.рез владельца: расх. было → стало")
    for P in sorted(mon):
        w, cc, rw = mon[P]["was"], mon[P]["cat"], mon[P]["row"]
        m2 = P[5:7]
        b_cogs, b_fin = book.get(m2, (None, None)) if "(" not in P else (None, None)
        m_cogs, m_fin = manual.get(m2, (None, None)) if (m2 != "09" or "(" in P) else (None, None)
        base = b_cogs if b_cogs is not None else w
        fin_new = (b_fin + (base - cc)) if b_fin is not None else None
        line = f"{P} | {cnt[P]['rows']} | {w:,.2f} | {cc:,.2f} | {(cc / w - 1) * 100:+.2f} % | {rw:,.2f} | {(rw / w - 1) * 100:+.2f} %"
        line += f" | {b_fin:,.2f} → {fin_new:,.2f}" if b_fin is not None else " | —"
        line += f" | {m_cogs:,.2f}: {w - m_cogs:,.2f} → {cc - m_cogs:,.2f}" if m_cogs is not None else " | —"
        line += f" | {m_fin:,.2f}: {b_fin - m_fin:,.2f} → {fin_new - m_fin:,.2f}" if (m_fin is not None and b_fin is not None) else " | —"
        print(line)
        if b_cogs is not None and abs(b_cogs - w) > D("0.5"):
            print(f"      книга: СС {b_cogs:,.2f} ≠ «было» {w:,.2f} — книга / было = {b_cogs / w:.4f} (индекс СС 1,150 с 09-01); фин. рез. «стало» считан от СС книги")
        print(f"      счётчики: {dict(cnt[P])}")
    g_0901, g_1006 = rate_on(G, "2026-09-01"), rate_on(G, "2026-10-06")
    import datetime as dt
    sept = [(dt.date(2026, 9, 1) + dt.timedelta(i)).isoformat() for i in range(30)]
    avg = sum(rate_on(G, d) for d in sept) / 30
    print(f"\nиндекс 1,150 с 09-01 против курса 1С золота: 20.05 {g0:.2f} → 01.09 {g_0901} ({(g_0901 / g0 - 1) * 100:+.2f} %) → 06.10 {g_1006} ({(g_1006 / g0 - 1) * 100:+.2f} %); "
          f"среднее по дням сентября {avg:.2f} ({(avg / g0 - 1) * 100:+.2f} %)")
    print("db_writes = 0; обращений к API — 0")


if __name__ == "__main__":
    main()
