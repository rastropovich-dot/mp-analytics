#!/usr/bin/env python3
"""Две книги «Фин рез» ячейка в ячейку (контрольная сборка main против ветки, сорок восьмая §3). Только чтение.

    venv/bin/python3 scripts/finrez_books_diff.py A.xlsx B.xlsx [--skip "Данные Ozon выкупы,Данные заказы,Данные WB выкупы"] [--tol 0.005] [--show 12]

Листы сравниваются по именам (нет в одной из книг — названо вслух); по каждому листу — ячеек сравнено, числовых разниц, разниц в тексте,
и свод разниц по колонке (заголовок — первая строка листа с ≥ 3 текстами над первой разницей) и по строке (значение колонки A): видно,
«только СС и зависимые» это или нет. Листы сводных (их значения считает Excel, в файле кэш) и длинные «Данные …» по умолчанию пропускаются
(десятки МБ; их сверяет --check самой книги).
"""
import argparse
from collections import Counter, defaultdict
from decimal import Decimal

SKIP_DEFAULT = "Данные Ozon выкупы,Данные заказы,Данные WB выкупы"


def load(path, skip):
    import openpyxl
    import warnings
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in wb.sheetnames:
        if name in skip or name.startswith("Сводная"):
            continue
        rows = []
        for r in wb[name].iter_rows(values_only=True):
            rows.append(r)
        out[name] = rows
    wb.close()
    return out


def header_for(rows, i, j):
    for k in range(i, -1, -1):
        r = rows[k]
        if sum(1 for v in r if isinstance(v, str)) >= 3 and j < len(r) and isinstance(r[j], str):
            return r[j][:40]
    return f"col{j + 1}"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("a"); ap.add_argument("b")
    ap.add_argument("--skip", default=SKIP_DEFAULT)
    ap.add_argument("--tol", default="0.005", help="|a − b| ≤ tol — не разница (округление)")
    ap.add_argument("--show", type=int, default=12, help="сколько первых разниц печатать на лист")
    args = ap.parse_args(argv)
    skip = {s.strip() for s in args.skip.split(",") if s.strip()}
    tol = Decimal(args.tol)
    A, B = load(args.a, skip), load(args.b, skip)
    only_a, only_b = sorted(set(A) - set(B)), sorted(set(B) - set(A))
    if only_a or only_b:
        print(f"листы только в A: {only_a}; только в B: {only_b}")
    total_cells = total_num = total_txt = 0
    for name in sorted(set(A) & set(B)):
        ra, rb = A[name], B[name]
        n = max(len(ra), len(rb))
        cells = num = txt = 0
        by_col, by_row, shown = Counter(), Counter(), 0
        examples = []
        for i in range(n):
            xa = ra[i] if i < len(ra) else ()
            xb = rb[i] if i < len(rb) else ()
            m = max(len(xa), len(xb))
            for j in range(m):
                va = xa[j] if j < len(xa) else None
                vb = xb[j] if j < len(xb) else None
                if va is None and vb is None:
                    continue
                cells += 1
                if isinstance(va, (int, float, Decimal)) and isinstance(vb, (int, float, Decimal)) and not isinstance(va, bool):
                    d = Decimal(str(vb)) - Decimal(str(va))
                    if abs(d) > tol:
                        num += 1
                        by_col[header_for(rb if i < len(rb) else ra, i, j)] += 1
                        by_row[str((xb[0] if xb else xa[0]) if (xb or xa) else "")[:20]] += 1
                        if len(examples) < args.show:
                            examples.append(f"      r{i + 1}c{j + 1} [{header_for(rb if i < len(rb) else ra, i, j)}] {va} → {vb} ({d:+,.2f})")
                elif va != vb:
                    txt += 1
                    if len(examples) < args.show:
                        examples.append(f"      r{i + 1}c{j + 1} текст: {str(va)[:60]!r} → {str(vb)[:60]!r}")
        total_cells += cells; total_num += num; total_txt += txt
        flag = "" if not (num or txt) else "  ← РАЗНИЦЫ"
        print(f"{name}: ячеек {cells:,}, числовых разниц {num:,}, текстовых {txt:,}{flag}")
        if num:
            print("   по колонкам: " + "; ".join(f"{k} {v}" for k, v in by_col.most_common(12)))
            print("   по строкам (A): " + "; ".join(f"{k} {v}" for k, v in sorted(by_row.items())[:40]))
        for e in examples:
            print(e)
    print(f"\nитого: ячеек {total_cells:,}, числовых разниц {total_num:,}, текстовых {total_txt:,}; пропущены листы: {sorted(skip)} и «Сводная …»")


if __name__ == "__main__":
    main()
