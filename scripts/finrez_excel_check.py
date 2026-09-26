#!/usr/bin/env python3
"""Живая проверка книги «Фин рез» в Excel for Mac (сорок первая §2): сводные владельца, срезы, группировки «+/−», формулы «Артикула».

    venv/bin/python3 scripts/finrez_excel_check.py data/reports/finrez_2026-04_2026-09.xlsx [--articles F1,T2,F3] [--mp Ozon] [--timeout 600]

Excel скриптуется через osascript (AppleScript): книга открывается, `refresh all`, читаются области сводных (`table range2`), поля и срезы,
месяц раскрывается в дни (`show detail` у pivot item), у листов формы читаются `summary row` и скрытость строк дней до и после `show levels`;
на «Артикуле» в B1 / B2 подставляются артикул и площадка, лист пересчитывается (`calculate`) и блоки читаются. В файл ничего не пишется:
книга закрывается `close … saving no`; если Excel не ответил (модальное окно, «восстановить книгу?») — текст окна читается через System Events
и печатается дословно, это и есть результат. Каждая проверка — отдельный AppleScript с таймаутом; ошибка Excel печатается как есть.
Все обращения — к своей книге ПО ИМЕНИ (`workbook "<файл>"`), никогда к active workbook: 2026-09-26 проба с active workbook попала в открытую
книгу советника и закрыла её без сохранения. Если книга с таким именем уже открыта — скрипт останавливается, чужие книги не трогает.

Числа сводных сверяются со статичными листами книги (openpyxl, read_only): «Сводная Ozon выкупы» → «Выкупы Ozon» (Начисления по месяцу × статье),
«Сводная WB выкупы» → «Выкупы WB» (семь денег по месяцам), «Сводная заказы» → «Заказы» (шесть денег по (МП, месяц), проценты до 0,01 п. п.);
«Артикул» → Σ по листам «Данные …» тем же артикулом. Код 0 — всё сошлось; 1 — есть расхождения или ошибки проверок; 2 — Excel недоступен.
"""
import argparse
import os
import re
import subprocess
import sys
import time
import warnings
from collections import defaultdict
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MONTHS = ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")
ARTICLES = ("Комиссия", "Логистика", "Прочее", "Реклама", "Товарооборот", "Эквайринг")
WB_MONEY = ("Продажи, ₽", "Комиссия, ₽", "Себес-ть, ₽", "Реклама, ₽", "Хранение, ₽", "Логистика, ₽", "Ост.расходы и компенсации МП, ₽")
ORDER_MONEY = ("Оборот (с НДС)", "Выручка", "Себестоимость", "Маржа", "Реклама", "Гр.фин.рез")
ORDER_PCT = ("Гр.фин.рез, %", "ДДР, %")
PIVOTS = (("Сводная Ozon выкупы", "Выкупы Ozon"), ("Сводная WB выкупы", "Выкупы WB"), ("Сводная заказы", "Заказы"))
Z = Decimal(0)


# ---------- osascript ----------

class ExcelError(RuntimeError):
    pass


class _Stop(Exception):
    """Дальше проверять нечего (книга не открылась) — таблица результатов печатается всё равно."""


def q(s):
    """Строка в литерал AppleScript."""
    return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def osascript(script, timeout=120, structured=True, run=None):
    """Выполнить AppleScript; вернуть stdout (структурный вывод -s s). Ошибка Excel / таймаут — ExcelError с текстом дословно."""
    run = run or subprocess.run
    args = ["osascript"] + (["-s", "s"] if structured else []) + ["-e", script]
    try:
        res = run(args, capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        raise ExcelError(f"osascript не ответил за {timeout} с (Excel занят модальным окном или ждёт разрешения macOS на автоматизацию)")
    if res.returncode != 0:
        raise ExcelError((res.stderr or res.stdout).strip())
    return res.stdout.strip()


def parse_structured(text):
    """Разбор структурного вывода osascript (-s s): списки {…}, строки "…", числа, missing value, true / false, date "…"."""
    text = text.strip()
    pos = 0

    def skip():
        nonlocal pos
        while pos < len(text) and text[pos] in " \t\r\n":
            pos += 1

    def value():
        nonlocal pos
        skip()
        if pos >= len(text):
            return None
        c = text[pos]
        if c == "{":
            pos += 1
            out = []
            while True:
                skip()
                if text[pos] == "}":
                    pos += 1
                    return out
                out.append(value())
                skip()
                if text[pos] == ",":
                    pos += 1
        if c == '"':
            pos += 1
            buf = []
            while text[pos] != '"':
                if text[pos] == "\\":
                    pos += 1
                buf.append(text[pos]); pos += 1
            pos += 1
            return "".join(buf)
        if text.startswith("missing value", pos):
            pos += len("missing value"); return None
        if text.startswith("true", pos):
            pos += 4; return True
        if text.startswith("false", pos):
            pos += 5; return False
        if text.startswith("date ", pos):
            pos += 5
            return "date " + value()
        if c == "«":
            end = text.index("»", pos)
            tok = text[pos:end + 1]; pos = end + 1
            return tok
        m = re.match(r"[-+]?\d+(\.\d+)?(E[-+]?\d+)?(?![.\w])", text[pos:])
        if m:
            pos += m.end()
            tok = m.group(0)
            return Decimal(tok) if ("." in tok or "E" in tok) else int(tok)
        m = re.match(r"[^,}\n]+", text[pos:])
        pos += m.end()
        return m.group(0).strip()

    return value()


def excel(script, timeout=120, run=None):
    return parse_structured(osascript(f'tell application "Microsoft Excel"\n{script}\nend tell', timeout=timeout, run=run))


def dialog_text(run=None):
    """Текст модального окна Excel через System Events (нужно разрешение «Универсальный доступ»); ошибка — как есть."""
    try:
        return parse_structured(osascript('tell application "System Events" to tell process "Microsoft Excel" to get {name of every window, '
                                          'value of every static text of window 1}', timeout=20, run=run))
    except ExcelError as e:
        return f"System Events: {e}"


# ---------- статичные листы (openpyxl) ----------

def read_sheet(path, name, max_row=None):
    import openpyxl
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=max_row, values_only=True)]
    wb.close()
    return rows


def dec(v):
    if v is None or v == "":
        return None
    if isinstance(v, Decimal):
        return v
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    try:
        return Decimal(str(v).replace(" ", "").replace(",", "."))
    except Exception:  # noqa: BLE001
        return None


def is_month(v):
    return isinstance(v, str) and v.strip() in MONTHS


def static_ozon_buyouts(rows):
    """«Выкупы Ozon»: {месяц: {статья: Начисления}} по шапке (строка 2 — статьи, forward-fill; строка 3 — «Начисления» / «Ст-ть»)."""
    heads2, heads3 = rows[1], rows[2]
    article = None
    cols = {}
    for j, (a, b) in enumerate(zip(heads2, heads3)):
        article = a if a else article
        if article in ARTICLES and b and "Начисления" in str(b):
            cols[article] = j
    out = {}
    for r in rows[3:]:
        if r and is_month(r[0]):
            out[r[0].strip()] = {a: dec(r[j]) for a, j in cols.items()}
        if r and r[0] == "Общий итог":
            out["Общий итог"] = {a: dec(r[j]) for a, j in cols.items()}
    return out


def static_by_header(rows, header_row, money):
    """Лист с шапкой в строке header_row (1-based): {метка месяца / «Общий итог»: {колонка: значение}} для колонок money (первое вхождение)."""
    heads = rows[header_row - 1]
    cols = {}
    for j, h in enumerate(heads):
        if h and str(h).strip() in money and str(h).strip() not in cols:
            cols[str(h).strip()] = j
    out = {}
    for r in rows[header_row:]:
        if r and (is_month(r[0]) or r[0] == "Общий итог"):
            out[str(r[0]).strip()] = {k: dec(r[j]) for k, j in cols.items()}
    return out


def static_orders(rows):
    """«Заказы»: {(МП, метка): {колонка: значение}} — два блока с одинаковой шапкой в строке 3, площадки в строке 2."""
    mp_row, heads = rows[1], rows[2]
    blocks = [(j, str(v)) for j, v in enumerate(mp_row) if v in ("Ozon", "WB")]
    out = {}
    for r in rows[3:]:
        if not r or not (is_month(r[0]) or r[0] == "Общий итог"):
            continue
        for bi, (j0, mp) in enumerate(blocks):
            j1 = blocks[bi + 1][0] if bi + 1 < len(blocks) else len(heads)
            vals = {}
            for j in range(j0, j1):
                h = heads[j] if j < len(heads) else None
                if h and str(h).strip() in ORDER_MONEY + ORDER_PCT and str(h).strip() not in vals:
                    vals[str(h).strip()] = dec(r[j]) if j < len(r) else None
            out[(mp, str(r[0]).strip())] = vals
    return out


# ---------- область сводной ----------

def find_labels(grid):
    """В прочитанной области сводной: строки-метки (месяцы / «Общий итог») и шапка (последняя строка до первого месяца, где есть текст)."""
    rows = {}
    header_idx = None
    for i, r in enumerate(grid):
        if r and is_month(r[0]):
            rows[str(r[0]).strip()] = i
            if header_idx is None:
                header_idx = i - 1
        elif r and r[0] == "Общий итог":
            rows["Общий итог"] = i
    return rows, header_idx


def pivot_columns(grid, header_idx, names, article_header=None):
    """Колонки области сводной по именам: шапка в header_idx; для листа выкупов Ozon — статьи строкой выше (article_header), forward-fill."""
    cols = {}
    heads = grid[header_idx] if header_idx is not None and header_idx >= 0 else []
    arts = grid[article_header] if article_header is not None and article_header >= 0 else []
    art = None
    for j, h in enumerate(heads):
        a = arts[j] if j < len(arts) else None
        art = a if a not in (None, "") else art
        hs = str(h).strip() if h is not None else ""
        if article_header is not None:
            if art in names and "Начисления" in hs and art not in cols:
                cols[art] = j
        elif hs in names and hs not in cols:
            cols[hs] = j
    return cols


def compare(static, pivot_grid, names, article_header=None, pct=(), tol_pct=Decimal("0.0001")):
    """[(метка, колонка, статичный, сводная, разница)] по общим меткам; пустая область — одна строка с пояснением."""
    rows, header_idx = find_labels(pivot_grid)
    if header_idx is None:
        return [("—", "—", None, None, "в области сводной не найдено строк-месяцев")]
    art_hdr = (header_idx - 1) if article_header else None
    cols = pivot_columns(pivot_grid, header_idx, names, art_hdr)
    out = []
    for lab, i in rows.items():
        if lab not in static:
            continue
        for name in names:
            j = cols.get(name)
            a = static[lab].get(name)
            b = dec(pivot_grid[i][j]) if (j is not None and j < len(pivot_grid[i])) else None
            if a is None and b is None:
                continue
            diff = ((a or Z) - (b or Z))
            diff = diff.quantize(tol_pct) if name in pct else diff.quantize(Decimal("0.01"))
            out.append((lab, name, a, b, diff))
    if not cols:
        out.append(("—", "—", None, None, f"колонки {names} не найдены в шапке сводной: {pivot_grid[header_idx][:12]}"))
    return out


# ---------- «Артикул»: Σ по листам «Данные» ----------

def sums_for_articles(path, articles):
    """Σ по «Данные Ozon выкупы» (Начисления по статье, Ст-ть, Соинвест, Количество) и «Данные заказы» (₽, шт, Реклама ₽, c СПП, Выручка, Маржа, Фин.рез)
    по (артикул, месяц-метка); площадка заказов — по колонке МП."""
    import openpyxl
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    arts = set(articles)
    buy = defaultdict(lambda: defaultdict(Decimal))
    ws = wb["Данные Ozon выкупы"]
    it = ws.iter_rows(values_only=True)
    heads = [str(h) if h is not None else "" for h in next(it)]
    ix = {h: i for i, h in enumerate(heads)}
    for r in it:
        a = r[ix["Артикул"]]
        if a in arts:
            key = (a, str(r[ix["Месяц"]]))
            buy[key][str(r[ix["Статьи Озон.Вид затрат"]])] += dec(r[ix["Начисления"]]) or Z
            buy[key]["Ст-ть продаж в себ-ти"] += dec(r[ix["Ст-ть продаж в себ-ти"]]) or Z
            buy[key]["Соинвест"] += dec(r[ix["Соинвест (баллы + зелёные цены)"]]) or Z
            buy[key]["Количество"] += dec(r[ix["Количество"]]) or Z
    orders = defaultdict(lambda: defaultdict(Decimal))
    ws = wb["Данные заказы"]
    it = ws.iter_rows(values_only=True)
    heads = [str(h) if h is not None else "" for h in next(it)]
    ix = {h: i for i, h in enumerate(heads)}
    for r in it:
        a = r[ix["Артикул поставщика"]]
        if a in arts:
            key = (a, str(r[ix["МП"]]), int(r[ix["Месяцы"]] or 0))
            for h in ("Заказы, ₽", "Заказы, шт", "Реклама, ₽", "Заказы, руб c СПП", "Выручка, руб без НДС с учетом комиссии", "Маржа, руб без НДС", "Фин.рез"):
                orders[key][h] += dec(r[ix[h]]) or Z
    wb.close()
    return buy, orders


ARTICLE_BLOCK = {"Товарооборот": "Товарооборот", "Комиссия": "Комиссия", "Логистика": "Логистика", "Реклама": "Реклама", "Эквайринг": "Эквайринг", "Прочее": "Прочее",
                 "Ст-ть продаж": "Ст-ть продаж в себ-ти", "Соинвест": "Соинвест", "Штуки": "Количество"}
ORDER_BLOCK = {"Заказы, ₽": "Заказы, ₽", "Заказы, шт": "Заказы, шт", "Реклама, ₽": "Реклама, ₽", "Заказы, руб c СПП": "Заказы, руб c СПП",
               "Выручка": "Выручка, руб без НДС с учетом комиссии", "Маржа": "Маржа, руб без НДС", "Фин.рез": "Фин.рез"}


def compare_article_sheet(grid, article, mp, buy, orders):
    """Блоки листа «Артикул» (значения после пересчёта Excel) против Σ «Данных»: [(блок, показатель, месяц, лист, данные, разница)]."""
    months = None
    out = []
    block = None
    for r in grid:
        if not r:
            continue
        lab = str(r[0]).strip() if r[0] is not None else ""
        if lab == "Показатель":
            months = [(j, str(v).strip()) for j, v in enumerate(r) if is_month(v)]
            continue
        if lab.startswith("Выкупы Ozon ("):
            block = "buy"; continue
        if lab.startswith("Заказы ("):                       # заголовок блока: «Заказы (площадка — B2; …)»; строки блока — «Заказы, ₽», «Заказы, шт»
            block = "orders"; continue
        if lab.startswith("Выкупы WB ("):
            block = "wb"; continue
        if not months or block not in ("buy", "orders"):
            continue
        key = next((k for k in (ARTICLE_BLOCK if block == "buy" else ORDER_BLOCK) if lab == k or lab.startswith(k + " ")), None)
        if key is None:
            continue
        for j, mo in months:
            got = dec(r[j]) if j < len(r) else None
            if block == "buy":
                exp = buy.get((article, mo), {}).get(ARTICLE_BLOCK[key], Z)
            else:
                exp = orders.get((article, mp, MONTHS.index(mo) + 1), {}).get(ORDER_BLOCK[key], Z)
            if not got and exp == 0:
                continue
            out.append((block, key, mo, got, exp, ((got or Z) - exp).quantize(Decimal("0.01"))))
    return out


# ---------- проверки ----------

def run(path, articles=None, mp="Ozon", timeout=600, run_=None, out=print, attach=False):
    """Все проверки; печатает таблицу; возвращает код (0 / 1 / 2). attach — книга уже открыта в Excel (например, открывалась дольше таймаута):
    не открывать, а работать с ней по имени; закрывается в конце всё равно (это наша книга)."""
    path = os.path.abspath(path)
    results = []           # (проверка, статус, детали)
    code = 0
    t0 = time.time()

    def rec(name, ok, detail):
        nonlocal code
        results.append((name, "ок" if ok else "НЕТ", detail))
        if not ok:
            code = 1

    try:
        ver = excel("return version", timeout=30, run=run_)
    except ExcelError as e:
        out(f"Excel недоступен: {e}")
        return 2
    out(f"Excel {ver}; книга {path}")
    opened = False
    book = os.path.basename(path)

    def ex(script, timeout_=120):
        """AppleScript к СВОЕЙ книге: «active workbook» в тексте заменяется ссылкой по имени."""
        return excel(f"set WB to workbook {q(book)}\n" + script.replace("active workbook", "WB"), timeout=timeout_, run=run_)

    try:
        try:
            before = excel("return name of every workbook", timeout=60, run=run_)
            before = [] if before is None else (before if isinstance(before, list) else [before])
            if book in before and not attach:
                raise ExcelError(f"книга «{book}» уже открыта в Excel — закройте её сами или запустите с --attach; чужие книги скрипт не трогает")
            if attach and book not in before:
                raise ExcelError(f"--attach: книги «{book}» среди открытых нет: {before}")
            if [b for b in before if b != book]:
                out(f"в Excel уже открыты: {[b for b in before if b != book]} — не трогаю")
            t = time.time()
            if not attach:
                excel(f"with timeout of {timeout} seconds\nopen (POSIX file {q(path)})\nend timeout", timeout=timeout + 10, run=run_)
            for _ in range(max(1, timeout // 5)):
                names = excel("return name of every workbook", timeout=60, run=run_)
                names = [] if names is None else (names if isinstance(names, list) else [names])
                if book in names:
                    break
                time.sleep(5)
            else:
                raise ExcelError(f"книга «{book}» не появилась среди открытых за {timeout} с: {names}")
            opened = True
            info = ex("return {count of worksheets of WB, name of every sheet of WB, name of active workbook}", 120)
            rec("открытие книги", True, f"{time.time() - t:.1f} с, листов {info[0]}: {info[1]}; активная книга Excel — {info[2]}")
        except ExcelError as e:
            rec("открытие книги", False, f"{e}; окно Excel: {dialog_text(run_)}")
            raise _Stop()
        try:
            t = time.time()
            ex(f"with timeout of {timeout} seconds\nrefresh all WB\nend timeout", timeout + 10)
            rec("refresh all", True, f"{time.time() - t:.1f} с")
        except ExcelError as e:
            rec("refresh all", False, f"{e}; окно Excel: {dialog_text(run_)}")
        # сводные: поля, область, срезы
        grids = {}
        for sheet, static_sheet in PIVOTS:
            try:
                try:
                    n_pt = ex(f'return count of pivot tables of sheet {q(sheet)} of active workbook', 120)
                except ExcelError as e:
                    n_pt = f"ошибка: {e}"
                if n_pt != 1:
                    head = ex(f'return value of range "A1:F14" of sheet {q(sheet)} of active workbook', 120)
                    rec(f"{sheet}: область сводной", False, f"сводных на листе {n_pt} — Excel не видит сводную; A1:F14 листа: {head}")
                    continue
                meta = ex(f'set pt to pivot table 1 of sheet {q(sheet)} of active workbook\n'
                          f'return {{name of pt, get address of table range2 of pt, count of rows of table range2 of pt, name of every pivot field of pt}}', 120)
                name, addr, nrows, fields = meta
                grid = ex(f'return value of table range2 of pivot table 1 of sheet {q(sheet)} of active workbook', 300)
                grid = grid if isinstance(grid, list) and grid and isinstance(grid[0], list) else [grid]
                grids[sheet] = grid
                lab_rows, hdr = find_labels(grid)
                rec(f"{sheet}: область сводной", bool(lab_rows), f"«{name}», {addr}, строк {nrows}, месяцев в области {len([l for l in lab_rows if l != 'Общий итог'])}, "
                                                             f"полей {len(fields)}: {fields}")
            except ExcelError as e:
                rec(f"{sheet}: область сводной", False, str(e))
        # срезы на «Сводная заказы», вычисляемые поля
        try:
            # AppleScript-класс slicer у Excel for Mac без свойств и «slicers of pivot table» всегда 0 — срезы видны как фигуры листа (сорок первая §2)
            sl = ex('return {count of shapes of sheet "Сводная заказы" of active workbook, name of every shape of sheet "Сводная заказы" of active workbook}', 60)
            names = sl[1] if isinstance(sl[1], list) else [sl[1]]
            rec("Сводная заказы: срезы", sl[0] >= 2 and any("День" in str(n) for n in names) and any("МП" in str(n) for n in names), f"фигур на листе {sl[0]}: {names}")
        except ExcelError as e:
            rec("Сводная заказы: срезы", False, str(e))
        try:
            fields = ex('return name of every pivot field of pivot table 1 of sheet "Сводная заказы" of active workbook', 60)
            calc = [f for f in ("Маржа, руб без НДС", "Маржинальность, %", "ДДР, %", "Соинвест, %", "Комиссия сред", "Фин.рез", "Цена", "Фин.рез %")]
            present = [f for f in calc if any(str(x).strip() == f for x in fields)]
            rec("Сводная заказы: вычисляемые поля", len(present) == len(calc), f"есть {len(present)} из {len(calc)}: {present}; нет: {[f for f in calc if f not in present]}")
        except ExcelError as e:
            rec("Сводная заказы: вычисляемые поля", False, str(e))
        # раскрытие месяца на «Сводная WB выкупы»
        try:
            fields = ex('return name of every pivot field of pivot table 1 of sheet "Сводная WB выкупы" of active workbook', 60)
            mf = next((str(f) for f in fields if "Месяц" in str(f)), None)
            if mf is None:
                raise ExcelError(f"поля месяцев нет среди {fields}")
            items = ex(f'return name of every pivot item of pivot field {q(mf)} of pivot table 1 of sheet "Сводная WB выкупы" of active workbook', 60)
            items = [str(i) for i in (items if isinstance(items, list) else [items])]
            month = next((i for i in items if i and i[0] not in "<>"), None)            # имена элементов в AppleScript английские («Apr»), на листе — «апр»
            if month is None:
                raise ExcelError(f"у поля «{mf}» нет элементов-месяцев: {items}")
            before = ex('return count of rows of table range2 of pivot table 1 of sheet "Сводная WB выкупы" of active workbook', 60)
            after = ex(f'set pt to pivot table 1 of sheet "Сводная WB выкупы" of active workbook\n'
                       f'set show detail of pivot item {q(month)} of pivot field {q(mf)} of pt to true\nreturn count of rows of table range2 of pt', 120)
            rec("Сводная WB выкупы: месяц → дни", after > before, f"поле «{mf}», элемент «{month}» (все: {items[:4]} …), строк {before} → {after} (+{after - before})")
        except ExcelError as e:
            rec("Сводная WB выкупы: месяц → дни", False, str(e))
        # числа сводных против статичных листов
        try:
            st = static_ozon_buyouts(read_sheet(path, "Выкупы Ozon"))
            table = compare(st, grids.get("Сводная Ozon выкупы", []), ARTICLES, article_header=True)
            bad = [t for t in table if isinstance(t[4], str) or t[4]]
            rec("Сводная Ozon выкупы = Выкупы Ozon", not bad, f"сравнений {len(table)}, расхождений {len(bad)}" + (f": {bad[:6]}" if bad else ""))
        except Exception as e:  # noqa: BLE001
            rec("Сводная Ozon выкупы = Выкупы Ozon", False, f"{type(e).__name__}: {e}")
        try:
            st = static_by_header(read_sheet(path, "Выкупы WB"), 2, WB_MONEY)
            table = compare(st, grids.get("Сводная WB выкупы", []), WB_MONEY)
            bad = [t for t in table if isinstance(t[4], str) or t[4]]
            rec("Сводная WB выкупы = Выкупы WB", not bad, f"сравнений {len(table)}, расхождений {len(bad)}" + (f": {bad[:6]}" if bad else ""))
        except Exception as e:  # noqa: BLE001
            rec("Сводная WB выкупы = Выкупы WB", False, f"{type(e).__name__}: {e}")
        try:
            st_o = static_orders(read_sheet(path, "Заказы"))
            grid = grids.get("Сводная заказы", [])
            lab_rows, hdr = find_labels(grid)
            # блоки МП — строкой выше шапки: «Ozon» … «WB»
            mp_row = grid[hdr - 1] if hdr and hdr >= 1 else []
            blocks = [(j, str(v)) for j, v in enumerate(mp_row) if v in ("Ozon", "WB")]
            table = []
            for bi, (j0, mp_) in enumerate(blocks):
                j1 = blocks[bi + 1][0] if bi + 1 < len(blocks) else len(grid[hdr])
                sub = [row[j0:j1] if len(row) > j0 else [] for row in grid]
                for row in sub:
                    pass
                sub_grid = [([grid[i][0]] + (grid[i][j0:j1] if len(grid[i]) > j0 else [])) for i in range(len(grid))]
                st_mp = {lab: v for (m_, lab), v in st_o.items() if m_ == mp_}
                for t in compare(st_mp, sub_grid, ORDER_MONEY + ORDER_PCT, pct=ORDER_PCT):
                    table.append((mp_,) + t)
            bad = [t for t in table if isinstance(t[5], str) or t[5]]
            rec("Сводная заказы = Заказы", bool(table) and not bad, f"блоков {len(blocks)}, сравнений {len(table)}, расхождений {len(bad)}" + (f": {bad[:6]}" if bad else ""))
        except Exception as e:  # noqa: BLE001
            rec("Сводная заказы = Заказы", False, f"{type(e).__name__}: {e}")
        # «Артикул»
        if articles:
            try:
                buy, orders = sums_for_articles(path, articles)
                for art in articles:
                    t = time.time()
                    grid = ex(f'set sh to sheet "Артикул" of active workbook\nset value of range "B1" of sh to {q(art)}\nset value of range "B2" of sh to {q(mp)}\n'
                              f'calculate\nreturn value of range "A1:H60" of sh', 300)
                    dt = time.time() - t
                    table = compare_article_sheet(grid, art, mp, buy, orders)
                    bad = [x for x in table if x[5]]
                    rec(f"Артикул {art} ({mp})", bool(table) and not bad, f"пересчёт {dt:.1f} с, сравнений {len(table)}, расхождений {len(bad)}" + (f": {bad[:6]}" if bad else ""))
            except Exception as e:  # noqa: BLE001
                rec("Артикул", False, f"{type(e).__name__}: {e}")
        # плюсики: summary row сверху, дни скрыты, show levels раскрывает
        for sheet in ("Выкупы Ozon", "Заказы"):
            try:
                res = ex(f'set sh to sheet {q(sheet)} of active workbook\nset o to outline object of sh\nset sr to summary row of o\n'
                         f'set n to count of rows of used range of sh\nset h1 to 0\nrepeat with i from 4 to n\nif hidden of row i of sh then set h1 to h1 + 1\nend repeat\n'
                         f'show levels o row levels 2\nset h2 to 0\nrepeat with i from 4 to n\nif hidden of row i of sh then set h2 to h2 + 1\nend repeat\n'
                         f'show levels o row levels 1\nreturn {{sr, n, h1, h2}}', 300)
                sr, n, h1, h2 = res
                rec(f"{sheet}: плюсики", "above" in str(sr).lower() and h1 > 0 and h2 < h1,
                    f"summary row «{sr}»; строк {n}, скрыто при уровне 1: {h1}, при уровне 2: {h2} (дни свёрнуты → раскрыты)")
            except ExcelError as e:
                rec(f"{sheet}: плюсики", False, str(e))
    except _Stop:
        pass
    finally:
        if opened:
            try:
                excel(f"close workbook {q(book)} saving no", timeout=120, run=run_)
                out(f"книга «{book}» закрыта без сохранения; другие книги Excel не тронуты")
            except ExcelError as e:
                out(f"книга НЕ закрыта: {e}")
    out(f"\n{'проверка':44} {'итог':5} детали")
    for name, st, detail in results:
        out(f"{name:44} {st:5} {detail}")
    out(f"проверок {len(results)}, не сошлось {sum(1 for r in results if r[1] != 'ок')}; {time.time() - t0:.1f} с; в файл не писали")
    return code


def main(argv=None):
    ap = argparse.ArgumentParser(description="Живая проверка книги «Фин рез» в Excel (osascript); в файл не пишет.")
    ap.add_argument("book")
    ap.add_argument("--articles", help="артикулы для листа «Артикул» через запятую")
    ap.add_argument("--mp", default="Ozon")
    ap.add_argument("--timeout", type=int, default=600)
    ap.add_argument("--attach", action="store_true", help="книга уже открыта в Excel — не открывать, работать с ней по имени")
    args = ap.parse_args(argv)
    return run(args.book, [a for a in (args.articles or "").split(",") if a], args.mp, args.timeout, attach=args.attach)


if __name__ == "__main__":
    sys.exit(main())
