#!/usr/bin/env python3
"""Выгрузка данных для плана продаж РБХ на ноябрь (inbox_rbh 2026-10-29): 4 кабинета-площадки — РБХ-1 Ozon / WB, РБХ-2 Ozon / WB.

Записей в базы нет: только чтение базы (часть A) и только чтение API (часть B, днём, вне окна ночи).

    MP_CABINET=rbh1 scripts/rbh_sales_plan_export.py db      — таблицы базы → data/<кабинет>/sales_plan_raw/db/*.json
    scripts/rbh_sales_plan_export.py book                      — книга из файлов обоих кабинетов → data/rbh_sales_plan/

Файл есть — повторно не читается (ключ --refresh перечитывает). Период — 23.08 … последний полный день (вчера по МСК).
"""
import argparse
import json
import os
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(ROOT, ".env"))
import cabinet  # noqa: E402

CABINETS = ("rbh1", "rbh2")
CAB_LABEL = {"rbh1": "РБХ-1", "rbh2": "РБХ-2"}
MP_LABEL = {"ozon": "Ozon", "wb": "WB"}
PERIOD_FROM = "2026-08-23"
MSK = timezone(timedelta(hours=3))
PAGE = 1_000


def last_full_day():
    return (datetime.now(MSK).date() - timedelta(days=1)).isoformat()


def raw_path(cab, *parts):
    p = cabinet.data_path("sales_plan_raw", *parts, prof=cabinet.profile(cab))
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


def out_dir():
    """Общий для кабинетов каталог итога: рядом с каталогами данных кабинетов."""
    p = os.path.join(os.path.dirname(cabinet.data_dir(cabinet.profile("rbh1"))), "rbh_sales_plan")
    os.makedirs(p, exist_ok=True)
    return p


# ---------- часть A: чтение базы ----------

def read_keyset(sb, table, select, day_col, id_col, day_from, day_to, filters=()):
    from loaders.keyset import read_keyset as rk
    return rk(sb, table, select, day_from, day_to, day_col=day_col, id_col=id_col, filters=list(filters))


def read_by_day(sb, table, select, day_col, order_cols, day_from, day_to):
    """Таблицы без уникального (день, id): день за днём, страницы по полному ключу с range."""
    from loaders import net_retry
    out, d = [], date.fromisoformat(day_from)
    while d <= date.fromisoformat(day_to):
        off = 0
        while True:
            qb = sb.table(table).select(select).eq(day_col, d.isoformat())
            for c in order_cols:
                qb = qb.order(c)
            data = net_retry.execute(qb.range(off, off + PAGE - 1), f"{table} {d}").data or []
            out.extend(data)
            if len(data) < PAGE:
                break
            off += PAGE
        d += timedelta(days=1)
    return out


def read_all(sb, table, select, order_col):
    from loaders import net_retry
    out, off = [], 0
    while True:
        data = net_retry.execute(sb.table(table).select(select).order(order_col).range(off, off + PAGE - 1), table).data or []
        out.extend(data)
        if len(data) < PAGE:
            return out
        off += PAGE


DB_READS = {
    "orders": lambda sb, f, t: read_keyset(
        sb, "marketplace_orders",
        "id,order_date,marketplace_code,marketplace_sku,article,product_name,order_schema,orders_qty,orders_amount_buyer,"
        "orders_amount_seller,cancelled_orders_qty,cancelled_orders_amount_buyer,cancelled_orders_amount_seller",
        "order_date", "id", f, t),
    "buyouts": lambda sb, f, t: read_keyset(
        sb, "marketplace_buyouts",
        "id,buyout_date,marketplace_code,marketplace_sku,article,product_name,buyouts_qty,buyouts_units,buyouts_amount_buyer,"
        "buyouts_amount_seller", "buyout_date", "id", f, t),
    "stock_daily": lambda sb, f, t: read_keyset(
        sb, "stock_daily",
        "id,stock_date,marketplace_code,marketplace_sku,article,warehouse_name,stock_qty,reserved_qty,available_qty,"
        "stock_marketplace_sku,stock_identity_status", "stock_date", "id", f, t),
    "placement": lambda sb, f, t: read_by_day(
        sb, "ozon_placement_by_product_daily", "day,sku,warehouse,offer_id,units,paid_units,product_type",
        "day", ("sku", "warehouse"), f, t),
    "wb_storage": lambda sb, f, t: read_by_day(
        sb, "wb_paid_storage_daily",
        "day,nm_id,chrt_id,office_id,gi_id,calc_type,barcode,size,vendor_code,subject,brand,warehouse,barcodes_count,"
        "original_date,tariff_fix_date,tariff_lower_date",
        "day", ("nm_id", "chrt_id", "office_id", "gi_id", "calc_type", "barcode"), f, t),
    "wb_funnel": lambda sb, f, t: read_keyset(
        sb, "wb_funnel_products_daily",
        "day,nm_id,vendor_code,title,brand,subject_name,stocks_wb,stocks_mp,order_count,order_sum,buyout_count,buyout_sum,"
        "cancel_count,cancel_sum,avg_price", "day", "nm_id", f, t),
    # продажи и возвраты отчёта реализации: rr_date с запасом на неделю назад, отбор по дню продажи — при сборке
    "wb_sales": lambda sb, f, t: read_keyset(
        sb, "wb_sales_report_rows",
        "rrd_id,rr_date,sale_dt,nm_id,vendor_code,tech_size,seller_oper_name,quantity,retail_price_with_disc,retail_amount,"
        "brand_name,subject_name", "rr_date", "rrd_id",
        (date.fromisoformat(f) - timedelta(days=7)).isoformat(), t,
        filters=[("in_", "seller_oper_name", ["Продажа", "Возврат"])]),
    "sku_catalog": lambda sb, f, t: read_all(sb, "sku_catalog", "marketplace_code,marketplace_sku,article,barcode,product_name,brand,category", "id"),
    "ozon_products": lambda sb, f, t: read_all(
        sb, "ozon_products", "sku,product_id,offer_id,name,type_name,category_path,brand,is_archived,status,observed_at", "sku"),
}


def cmd_db(args):
    prof = cabinet.assert_env()
    cab = os.environ.get("MP_CABINET")
    if cab not in CABINETS:
        raise SystemExit(f"только {', '.join(CABINETS)}; MP_CABINET={cab!r}")
    print(cabinet.banner(prof))
    from supabase import create_client
    sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])
    day_to = args.date_to or last_full_day()
    meta_p = raw_path(cab, "db", "meta.json")
    meta = json.load(open(meta_p)) if os.path.exists(meta_p) else {}
    for name, fn in DB_READS.items():
        if args.only and name not in args.only:
            continue
        p = raw_path(cab, "db", f"{name}.json")
        if os.path.exists(p) and not args.refresh:
            print(f"{name}: файл есть — пропускаю ({p})")
            continue
        t0 = time.time()
        rows = fn(sb, PERIOD_FROM, day_to)
        json.dump(rows, open(p, "w"), ensure_ascii=False, default=str)
        meta[name] = {"rows": len(rows), "from": PERIOD_FROM, "to": day_to, "read_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                      "seconds": round(time.time() - t0, 1)}
        json.dump(meta, open(meta_p, "w"), ensure_ascii=False, indent=1)
        print(f"{name}: {len(rows)} строк за {time.time() - t0:.0f} с")
    print("db_writes 0")


# ---------- общие правила ----------

def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def brand_norm(b):
    """Склейка регистра и точек: DEXclusive = DEXCLUSIVE, CATRICE. = CATRICE; пробелы схлопнуты."""
    s = re.sub(r"\s+", " ", str(b or "")).strip()
    s = re.sub(r"[.\s]+$", "", s)
    return s.upper()


def art_norm(a):
    return re.sub(r"\s+", "", str(a or "")).lower()


def sale_day_msk(ts):
    """saleDt → день продажи по МСК (как в marketplace_buyouts WB)."""
    s = str(ts or "")
    if not s:
        return None
    dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(MSK).date().isoformat()


def load_raw(cab, name):
    p = raw_path(cab, "db", f"{name}.json")
    if not os.path.exists(p):
        raise SystemExit(f"нет {p} — сначала MP_CABINET={cab} … db")
    return json.load(open(p))


# ---------- часть B: чтение API (только днём, вне окна ночи) ----------

OZON = "https://api-seller.ozon.ru"
WB_PRICES = "https://discounts-prices-api.wildberries.ru"
WB_SUPPLIES = "https://supplies-api.wildberries.ru"
WB_CALENDAR = "https://dp-calendar-api.wildberries.ru"
API_FROM_UTC, API_TO_UTC = 4 * 60 + 30, 21 * 60           # 04:30 … 21:00 UTC — правило задачи
NOV_FROM, NOV_TO = "2026-11-01", "2026-11-30"
OZON_ACTIVE_SUPPLY_STATES = ["DATA_FILLING", "READY_TO_SUPPLY", "ACCEPTED_AT_SUPPLY_WAREHOUSE", "IN_TRANSIT",
                             "ACCEPTANCE_AT_STORAGE_WAREHOUSE", "REPORTS_CONFIRMATION_AWAITING", "REPORT_REJECTED"]
LIST_PAGE_CAP = 60                                         # страниц списка заявок за прогон (100 заявок на странице)
BUNDLE_CAP = 1500                                          # составов заявок Ozon за прогон (план задачи: ≤ ~3 000 обращений на кабинет)
WB_ACTIVE_SUPPLY_STATUS = {1, 2, 3, 4, 6}                 # не запланировано / запланировано / отгрузка разрешена / приёмка / на воротах


class Api:
    """Счётчик обращений и ошибок одного кабинета; сырые ответы — в sales_plan_raw/api/ (файл есть — в API не идёт)."""

    def __init__(self, cab):
        self.cab = cab
        self.calls, self.c429, self.retries = Counter(), Counter(), Counter()
        self.errors = []

    def guard(self):
        now = datetime.now(timezone.utc)
        mins = now.hour * 60 + now.minute
        if not (API_FROM_UTC <= mins < API_TO_UTC):
            raise SystemExit(f"{now:%H:%M} UTC — вне окна 04:30 … 21:00 UTC, API не трогаю")

    def cached(self, *parts):
        p = raw_path(self.cab, "api", *parts)
        return json.load(open(p)) if os.path.exists(p) else None

    def save(self, data, *parts):
        json.dump(data, open(raw_path(self.cab, "api", *parts), "w"), ensure_ascii=False)
        return data

    def call(self, method, url, kind, headers, retry_429, **kw):
        """Один логический вызов: http_retry (Retry-After, 5xx, сеть); 429 без повтора внутри — пауза 60 с, не больше 3 раз."""
        from loaders import http_retry
        self.guard()
        for attempt in range(1, 4):
            stats = {}
            resp = http_retry.request(method, url, label=f"{self.cab} {kind}", retry_429=retry_429, headers=headers, timeout=120,
                                      stats=stats, **kw)
            self.calls[kind] += stats.get("requests", 1)
            self.retries[kind] += stats.get("retries", 0)
            self.c429[kind] += http_retry.count_429(stats, resp)
            if resp.status_code == 429 and attempt < 3:
                print(f"  {kind}: 429 — пауза 60 с ({attempt}/3)")
                time.sleep(60)
                continue
            break
        if resp.status_code != 200:
            self.errors.append({"kind": kind, "url": url.split("?")[0], "status": resp.status_code, "body": resp.text[:400]})
            print(f"  {kind}: HTTP {resp.status_code} {resp.text[:200]}")
            return None
        return resp.json()

    def ozon(self, path, body, kind):
        from loaders import http_retry
        h = {"Client-Id": os.environ["OZON_CLIENT_ID"], "Api-Key": os.environ["OZON_API_KEY"], "Content-Type": "application/json"}
        return self.call("POST", OZON + path, kind, h, http_retry.RETRY_429_OZON_RATE_LIMIT, json=body)

    def wb(self, method, url, kind, pause=0.0, **kw):
        from loaders import http_retry
        h = {"Authorization": os.environ["WB_API_KEY"], "Content-Type": "application/json"}
        out = self.call(method, url, kind, h, http_retry.RETRY_429_ANY, **kw)
        if pause:
            time.sleep(pause)
        return out

    def report(self):
        return {"calls": dict(self.calls), "429": dict(self.c429), "retries": dict(self.retries), "errors": self.errors}


def ozon_product_ids(cab):
    return sorted({int(r["product_id"]) for r in load_raw(cab, "ozon_products") if r.get("product_id")})


def ozon_skus(cab):
    skus = {str(r["sku"]) for r in load_raw(cab, "ozon_products")}
    skus |= {str(r["marketplace_sku"]) for r in load_raw(cab, "orders") if r["marketplace_code"] == "ozon"}
    skus |= {str(r["sku"]) for r in load_raw(cab, "placement")[-20000:]}
    return sorted(s for s in skus if s.isdigit())


def api_ozon(api, steps):
    cab = api.cab
    pids = ozon_product_ids(cab)
    if "attrs" in steps and api.cached("ozon_attributes.json") is None:
        out = []
        for vis in ("ALL", "ARCHIVED"):
            last = ""
            while True:
                r = api.ozon("/v4/product/info/attributes", {"filter": {"visibility": vis}, "limit": 1000, "last_id": last, "sort_dir": "ASC"},
                             "ozon /v4/product/info/attributes")
                if r is None:
                    break
                out += [dict(x, _visibility=vis) for x in r.get("result") or []]
                last = r.get("last_id") or ""
                if not last or not r.get("result"):
                    break
        api.save(out, "ozon_attributes.json")
        print(f"  атрибуты: {len(out)} карточек")
    if "prices" in steps and api.cached("ozon_prices.json") is None:
        out = []
        for vis in ("ALL", "ARCHIVED"):
            cursor = ""
            while True:
                r = api.ozon("/v5/product/info/prices", {"filter": {"visibility": vis}, "limit": 1000, "cursor": cursor}, "ozon /v5/product/info/prices")
                if r is None:
                    break
                out += [dict(x, _visibility=vis) for x in r.get("items") or []]
                cursor = r.get("cursor") or ""
                if not cursor or not r.get("items"):
                    break
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "items": out}, "ozon_prices.json")
        print(f"  цены: {len(out)} товаров")
    if "stocks" in steps and api.cached("ozon_stocks.json") is None:
        out, cursor = [], ""
        while True:
            r = api.ozon("/v4/product/info/stocks", {"filter": {"visibility": "ALL"}, "limit": 1000, "cursor": cursor}, "ozon /v4/product/info/stocks")
            if r is None:
                break
            out += r.get("items") or []
            cursor = r.get("cursor") or ""
            if not cursor or not r.get("items"):
                break
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "items": out}, "ozon_stocks.json")
        print(f"  остатки v4: {len(out)} товаров")
    if "analytics_stocks" in steps and api.cached("ozon_analytics_stocks.json") is None:
        skus, out = ozon_skus(cab), []
        for i in range(0, len(skus), 100):
            r = api.ozon("/v1/analytics/stocks", {"skus": skus[i:i + 100]}, "ozon /v1/analytics/stocks")
            if r is not None:
                out += r.get("items") or []
            time.sleep(0.5)
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "skus": len(skus), "items": out}, "ozon_analytics_stocks.json")
        print(f"  аналитика остатков: {len(out)} строк по {len(skus)} SKU")
    if "turnover" in steps and api.cached("ozon_turnover.json") is None:
        out, off = [], 0
        while True:
            r = api.ozon("/v1/analytics/turnover/stocks", {"limit": 1000, "offset": off}, "ozon /v1/analytics/turnover/stocks")
            if r is None:
                break
            items = r.get("items") or []
            out += items
            if len(items) < 1000:
                break
            off += 1000
            time.sleep(61)                                     # 1 запрос в минуту на Client-Id
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "items": out}, "ozon_turnover.json")
        print(f"  оборачиваемость: {len(out)} SKU")
    if "supplies" in steps and api.cached("ozon_supplies.json") is None:
        ids = {}
        for group, states, extra in (("active", OZON_ACTIVE_SUPPLY_STATES, {}),
                                     ("completed", ["COMPLETED"], {"timeslot_from_range": {"from": "2026-01-01T00:00:00Z", "to": "2026-12-31T23:59:59Z", "timeslot_filter_type": "BY_UTC_TIME"}})):
            last = ""
            while True:
                body = {"filter": dict({"states": states}, **extra), "limit": 100, "sort_by": "ORDER_CREATION", "sort_dir": "DESC", "last_id": last}
                r = api.ozon("/v3/supply-order/list", body, "ozon /v3/supply-order/list")
                if r is None:
                    break
                page = [str(x) for x in r.get("order_ids") or []]
                fresh = [x for x in page if x not in ids]
                for oid in page:
                    ids.setdefault(oid, group)
                nxt = r.get("last_id") or ""
                pages = api.calls["ozon /v3/supply-order/list"]
                print(f"  список заявок ({group}): страница {pages}, id {len(page)}, новых {len(fresh)}, last_id {nxt!r}")
                if not nxt or not page or not fresh or nxt == last or pages >= LIST_PAGE_CAP:
                    if fresh and pages >= LIST_PAGE_CAP:
                        api.errors.append({"kind": "ozon /v3/supply-order/list", "status": "cap", "body": f"потолок {LIST_PAGE_CAP} страниц"})
                    break
                last = nxt
        orders, idl = [], sorted(ids)
        for i in range(0, len(idl), 50):
            r = api.ozon("/v3/supply-order/get", {"order_ids": [int(x) for x in idl[i:i + 50]]}, "ozon /v3/supply-order/get")
            if r is not None:
                orders += r.get("orders") or []
        bundles, plan = {}, []
        for o in orders:
            group = ids.get(str(o.get("order_id")))
            for sp in o.get("supplies") or []:
                if sp.get("bundle_id"):
                    plan.append((group, sp["bundle_id"]))
        n_active = sum(1 for g, _ in plan if g == "active")
        print(f"  составов к чтению: активных {n_active}, завершённых {len(plan) - n_active} (потолок {BUNDLE_CAP})")
        plan = [x for x in plan if x[0] == "active"] + [x for x in plan if x[0] != "active"][:max(0, BUNDLE_CAP - n_active)]
        for group, bid in plan:
            if bid in bundles:
                continue
            if True:
                items, last = [], ""
                while True:
                    r = api.ozon("/v1/supply-order/bundle", {"bundle_ids": [bid], "limit": 100, "last_id": last, "is_asc": True},
                                 f"ozon /v1/supply-order/bundle ({group})")
                    if r is None:
                        break
                    items += r.get("items") or []
                    last = r.get("last_id") or ""
                    if not r.get("has_next"):
                        break
                bundles[bid] = items
                time.sleep(0.3)
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "ids": ids, "orders": orders, "bundles": bundles},
                 "ozon_supplies.json")
        print(f"  заявки на поставку: {len(ids)} (активных {sum(1 for g in ids.values() if g == 'active')}), составов {len(bundles)}")
    if "actions" in steps and api.cached("ozon_actions.json") is None:
        r = api.ozon_get_actions()
        actions = [a for a in (r or []) if str(a.get("date_start", ""))[:10] <= NOV_TO and str(a.get("date_end", "9999"))[:10] >= NOV_FROM]
        per = {}
        for a in actions:
            aid = a["id"]
            per[str(aid)] = {}
            for kind, path in (("products", "/v1/actions/products"), ("candidates", "/v1/actions/candidates")):
                items, last = [], None
                while True:
                    body = {"action_id": aid, "limit": 100}
                    if last:
                        body["last_id"] = last
                    rr = api.ozon(path, body, f"ozon {path}")
                    if rr is None:
                        break
                    res = rr.get("result") or {}
                    items += res.get("products") or []
                    nxt = res.get("last_id")
                    if not res.get("products") or not nxt or nxt == last:
                        break
                    last = nxt
                per[str(aid)][kind] = items
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "all_actions": r, "november": actions, "products": per},
                 "ozon_actions.json")
        print(f"  акции: всего {len(r or [])}, пересекают ноябрь {len(actions)}")


def _ozon_get_actions(self):
    from loaders import http_retry
    h = {"Client-Id": os.environ["OZON_CLIENT_ID"], "Api-Key": os.environ["OZON_API_KEY"]}
    r = self.call("GET", OZON + "/v1/actions", "ozon /v1/actions", h, http_retry.RETRY_429_OZON_RATE_LIMIT)
    return (r or {}).get("result")


Api.ozon_get_actions = _ozon_get_actions


def api_wb(api, steps):
    if "prices" in steps and api.cached("wb_prices.json") is None:
        out, off = [], 0
        while True:
            r = api.wb("GET", f"{WB_PRICES}/api/v2/list/goods/filter?limit=1000&offset={off}", "wb /api/v2/list/goods/filter", pause=0.7)
            goods = ((r or {}).get("data") or {}).get("listGoods") or []
            out += goods
            if r is None or len(goods) < 1000:
                break
            off += 1000
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "items": out}, "wb_prices.json")
        print(f"  цены WB: {len(out)} nmID")
    if "supplies" in steps and api.cached("wb_supplies.json") is None:
        sup, off = [], 0
        while True:
            body = {"dates": [{"from": "2025-01-01", "till": datetime.now(MSK).date().isoformat(), "type": "createDate"}]}
            r = api.wb("POST", f"{WB_SUPPLIES}/api/v1/supplies?limit=1000&offset={off}", "wb /api/v1/supplies", pause=2.1, json=body)
            page = r if isinstance(r, list) else []
            sup += page
            if r is None or len(page) < 1000:
                break
            off += 1000
        details, goods = {}, {}
        for s in sup:
            if s.get("statusID") not in WB_ACTIVE_SUPPLY_STATUS:
                continue
            sid, pre = (s.get("supplyID"), "false") if s.get("supplyID") else (s.get("preorderID"), "true")
            details[str(sid)] = api.wb("GET", f"{WB_SUPPLIES}/api/v1/supplies/{sid}?isPreorderID={pre}", "wb /api/v1/supplies/{ID}", pause=2.1)
            items, goff = [], 0
            while True:
                r = api.wb("GET", f"{WB_SUPPLIES}/api/v1/supplies/{sid}/goods?limit=1000&offset={goff}&isPreorderID={pre}",
                           "wb /api/v1/supplies/{ID}/goods", pause=2.1)
                page = r if isinstance(r, list) else []
                items += page
                if r is None or len(page) < 1000:
                    break
                goff += 1000
            goods[str(sid)] = items
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "supplies": sup, "details": details, "goods": goods},
                 "wb_supplies.json")
        print(f"  поставки WB: {len(sup)}, активных {len(details)}")
    if "calendar" in steps and api.cached("wb_calendar.json") is None:
        r = api.wb("GET", f"{WB_CALENDAR}/api/v1/calendar/promotions?startDateTime={NOV_FROM}T00:00:00Z&endDateTime={NOV_TO}T23:59:59Z"
                          "&allPromo=true&limit=1000&offset=0", "wb /api/v1/calendar/promotions", pause=0.7)
        promos = ((r or {}).get("data") or {}).get("promotions") or []
        details = []
        ids = [p["id"] for p in promos]
        for i in range(0, len(ids), 100):
            q = "&".join(f"promotionIDs={x}" for x in ids[i:i + 100])
            d = api.wb("GET", f"{WB_CALENDAR}/api/v1/calendar/promotions/details?{q}", "wb /api/v1/calendar/promotions/details", pause=0.7)
            details += ((d or {}).get("data") or {}).get("promotions") or []
        noms = {}
        for p in promos:
            if str(p.get("type", "")).lower() == "auto":
                noms[str(p["id"])] = {"skipped": "автоакция — метод товаров для неё неприменим"}
                continue
            noms[str(p["id"])] = {}
            for in_action in ("true", "false"):
                items, off = [], 0
                while True:
                    d = api.wb("GET", f"{WB_CALENDAR}/api/v1/calendar/promotions/nomenclatures?promotionID={p['id']}&inAction={in_action}"
                                      f"&limit=1000&offset={off}", "wb /api/v1/calendar/promotions/nomenclatures", pause=0.7)
                    page = ((d or {}).get("data") or {}).get("nomenclatures") or []
                    items += page
                    if d is None or len(page) < 1000:
                        break
                    off += 1000
                noms[str(p["id"])][in_action] = items
        api.save({"observed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "promotions": promos, "details": details,
                  "nomenclatures": noms}, "wb_calendar.json")
        print(f"  акции WB: {len(promos)}, деталей {len(details)}")


API_STEPS = {"ozon": ("attrs", "prices", "stocks", "analytics_stocks", "turnover", "supplies", "actions"),
             "wb": ("prices", "supplies", "calendar")}


def cmd_api(args):
    prof = cabinet.assert_env()
    cab = os.environ.get("MP_CABINET")
    if cab not in CABINETS:
        raise SystemExit(f"только {', '.join(CABINETS)}; MP_CABINET={cab!r}")
    print(cabinet.banner(prof))
    import signal
    signal.signal(signal.SIGTERM, lambda *_: sys.exit("остановлен (SIGTERM)"))      # счётчик обращений пишется и при остановке
    api = Api(cab)
    api.guard()
    mps = args.mp or ("ozon", "wb")
    status = "ok"
    try:
        for mp in mps:
            steps = args.steps or API_STEPS[mp]
            print(f"{CAB_LABEL[cab]} {MP_LABEL[mp]}: {', '.join(steps)}")
            (api_ozon if mp == "ozon" else api_wb)(api, steps)
    except BaseException as e:
        status = f"прерван: {e!r}"[:200]
        raise
    finally:
        p = raw_path(cab, "api", "calls.json")
        log = json.load(open(p)) if os.path.exists(p) else []
        log.append(dict(api.report(), at=datetime.now(timezone.utc).isoformat(timespec="seconds"), mps=list(mps),
                        steps=list(args.steps or []), status=status))
        json.dump(log, open(p, "w"), ensure_ascii=False, indent=1)
        print(json.dumps(api.report(), ensure_ascii=False, indent=1))
    print("db_writes 0")


# ---------- сборка данных кабинета-площадки ----------

MONTHS = (("2026-08", "август (23–31)"), ("2026-09", "сентябрь"), ("2026-10", "октябрь"))


def days_between(a, b):
    d, out = date.fromisoformat(a), []
    while d <= date.fromisoformat(b):
        out.append(d.isoformat())
        d += timedelta(days=1)
    return out


class CabMP:
    """Один кабинет-площадка: справочник, остатки по дням, продажи по дням — из файлов части A."""

    def __init__(self, cab, mp, day_to):
        self.cab, self.mp, self.day_to = cab, mp, day_to
        self.label = f"{CAB_LABEL[cab]} {MP_LABEL[mp]}"
        self.days = days_between(PERIOD_FROM, day_to)
        self.items = {}                     # sku → справочник
        self.stock = defaultdict(dict)      # (day, sku) → {источник: шт}
        self.stock_days = set()             # дни, за которые отчёт основного источника остатков загружен
        self.sales = defaultdict(lambda: defaultdict(float))   # (day, sku) → поля
        self.recon = {}

    def item(self, sku):
        return self.items.setdefault(str(sku), {"sku": str(sku), "article": "", "barcodes": set(), "sizes": set(), "chrts": set(),
                                                "name": "", "brand_src": "", "subject": "", "status": "", "archived": None})


def _fill(it, key, value):
    if value not in (None, "") and not it.get(key):
        it[key] = value


def ozon_pid2sku(cab):
    """product_id → SKU: у Ozon РБХ sku_catalog и stock_daily ключуются product_id (проверено 29.10: 829 / 736 ключей из 829 / 736 —
    product_id, ни одного SKU)."""
    out = {str(a["id"]): str(a["sku"]) for a in (Api(cab).cached("ozon_attributes.json") or []) if a.get("sku")}
    out.update({str(r["product_id"]): str(r["sku"]) for r in load_raw(cab, "ozon_products") if r.get("product_id")})
    return out


def build_ozon(cab, day_to):
    m = CabMP(cab, "ozon", day_to)
    pid2sku = ozon_pid2sku(cab)
    m.unmapped = Counter()
    for r in load_raw(cab, "ozon_products"):
        it = m.item(r["sku"])
        _fill(it, "article", r.get("offer_id"))
        _fill(it, "name", r.get("name"))
        _fill(it, "subject", r.get("type_name"))
        _fill(it, "brand_src", r.get("brand"))
        it["status"] = r.get("status") or ""
        it["archived"] = r.get("is_archived")
        it["product_id"] = r.get("product_id")
    for r in load_raw(cab, "sku_catalog"):
        if r["marketplace_code"] != "ozon":
            continue
        if str(r["marketplace_sku"]) not in pid2sku:
            m.unmapped["sku_catalog"] += 1
            continue
        it = m.item(pid2sku[str(r["marketplace_sku"])])
        _fill(it, "article", r.get("article"))
        _fill(it, "name", r.get("product_name"))
        if r.get("barcode"):
            it["barcodes"].update(b.strip() for b in str(r["barcode"]).split(",") if b.strip())
    attrs = Api(cab).cached("ozon_attributes.json")
    if attrs:
        for a in attrs:
            sku = str(a.get("sku") or "")
            if not sku or sku == "0":
                continue
            archived = a.get("_visibility") == "ARCHIVED"
            if sku not in m.items and archived:
                continue                                    # архивная карточка без истории в базе — не в справочник
            it = m.item(sku)
            brand = next((v.get("value") for at in a.get("attributes") or [] if at.get("id") == 85 for v in at.get("values") or []), None)
            it["brand_api"] = brand or ""
            it["api_seen"] = True
            _fill(it, "article", a.get("offer_id"))
            _fill(it, "name", a.get("name"))
            it["barcodes"].update(b for b in (a.get("barcodes") or []) if b)
            it["archived"] = archived
            it["product_id"] = it.get("product_id") or a.get("id")
    for r in load_raw(cab, "placement"):
        if r["day"] > day_to:
            continue
        sku = str(r["sku"])
        _fill(m.item(sku), "article", r.get("offer_id"))
        st = m.stock[(r["day"], sku)]
        st["fbo_placement"] = st.get("fbo_placement", 0) + (r.get("units") or 0)
        m.stock_days.add(r["day"])
    for r in load_raw(cab, "stock_daily"):
        if r["marketplace_code"] != "ozon" or r["stock_date"] > day_to:
            continue
        key = "fbs_stock" if r["warehouse_name"] == "fbs" else "fbo_stock"
        sku = pid2sku.get(str(r["marketplace_sku"]))
        if not sku:
            m.unmapped["stock_daily"] += 1
            continue
        st = m.stock[(r["stock_date"], sku)]
        st[key] = st.get(key, 0) + num(r.get("stock_qty"))
        st[key + "_days"] = 1
    for r in load_raw(cab, "orders"):
        if r["marketplace_code"] != "ozon" or r["order_date"] > day_to:
            continue
        it = m.item(r["marketplace_sku"])
        _fill(it, "article", r.get("article"))
        _fill(it, "name", r.get("product_name"))
        s = m.sales[(r["order_date"], str(r["marketplace_sku"]))]
        _add_orders(s, r)
        if r.get("order_schema") == "fbs":
            s["fbs_created_qty"] += num(r.get("orders_qty")) + num(r.get("cancelled_orders_qty"))
        if r.get("orders_amount_buyer") is None and num(r.get("orders_qty")):
            s["buyer_unmeasured"] = 1
    for r in load_raw(cab, "buyouts"):
        if r["marketplace_code"] != "ozon" or r["buyout_date"] > day_to:
            continue
        it = m.item(r["marketplace_sku"])
        _fill(it, "article", r.get("article"))
        _fill(it, "name", r.get("product_name"))
        s = m.sales[(r["buyout_date"], str(r["marketplace_sku"]))]
        units = r.get("buyouts_units")
        s["buy_qty"] += num(units if units is not None else r.get("buyouts_qty"))
        s["buy_positions"] += num(r.get("buyouts_qty"))
        if units is None:
            s["buy_units_unmeasured"] = 1
        s["buy_seller"] += num(r.get("buyouts_amount_seller"))
        s["buy_buyer"] += num(r.get("buyouts_amount_buyer"))
    return m


def _add_orders(s, r):
    s["ord_qty"] += num(r.get("orders_qty"))
    s["ord_seller"] += num(r.get("orders_amount_seller"))
    s["ord_buyer"] += num(r.get("orders_amount_buyer"))
    s["cnl_qty"] += num(r.get("cancelled_orders_qty"))
    s["cnl_seller"] += num(r.get("cancelled_orders_amount_seller"))
    s["cnl_buyer"] += num(r.get("cancelled_orders_amount_buyer"))


def build_wb(cab, day_to):
    m = CabMP(cab, "wb", day_to)
    funnel = sorted(load_raw(cab, "wb_funnel"), key=lambda r: r["day"], reverse=True)    # свежий день первым — его поля карточки
    for r in funnel:
        it = m.item(r["nm_id"])
        _fill(it, "article", r.get("vendor_code"))
        _fill(it, "name", r.get("title"))
        _fill(it, "brand_src", r.get("brand"))
        _fill(it, "subject", r.get("subject_name"))
    storage = load_raw(cab, "wb_storage")
    m.batches = defaultdict(lambda: {"count": 0, "calc": set(), "fix": set(), "lower": set()})
    m.gi_first = {}
    for r in sorted(storage, key=lambda r: r["day"], reverse=True):
        it = m.item(r["nm_id"])
        _fill(it, "article", r.get("vendor_code"))
        _fill(it, "brand_src", r.get("brand"))
        _fill(it, "subject", r.get("subject"))
        if r.get("barcode"):
            it["barcodes"].add(r["barcode"])
        if r.get("size") not in (None, ""):
            it["sizes"].add(str(r["size"]))
        it["chrts"].add(str(r["chrt_id"]))
    for r in storage:
        if r["day"] > day_to:
            continue
        nm = str(r["nm_id"])
        m.stock_days.add(r["day"])
        st = m.stock[(r["day"], nm)]
        st["wb_storage"] = st.get("wb_storage", 0) + (r.get("barcodes_count") or 0)
        if (r.get("barcodes_count") or 0) > 0:
            gk = (nm, str(r["gi_id"]))
            m.gi_first[gk] = min(m.gi_first.get(gk, r["day"]), r["day"])
        if r["day"] == day_to:
            b = m.batches[(nm, str(r["chrt_id"]), str(r.get("size") or ""), r.get("barcode") or "", str(r["gi_id"]))]
            b["count"] += r.get("barcodes_count") or 0
            b["calc"].add(r.get("calc_type") or "")
            if r.get("tariff_fix_date"):
                b["fix"].add(r["tariff_fix_date"])
            if r.get("tariff_lower_date"):
                b["lower"].add(r["tariff_lower_date"])
    for r in load_raw(cab, "stock_daily"):
        if r["marketplace_code"] != "wb" or r["stock_date"] > day_to:
            continue
        st = m.stock[(r["stock_date"], str(r["marketplace_sku"]))]
        st["wb_stock"] = st.get("wb_stock", 0) + num(r.get("stock_qty"))
    for r in load_raw(cab, "orders"):
        if r["marketplace_code"] != "wb" or r["order_date"] > day_to:
            continue
        it = m.item(r["marketplace_sku"])
        _fill(it, "article", r.get("article"))
        _add_orders(m.sales[(r["order_date"], str(r["marketplace_sku"]))], r)
    for r in load_raw(cab, "buyouts"):
        if r["marketplace_code"] != "wb" or r["buyout_date"] > day_to:
            continue
        _fill(m.item(r["marketplace_sku"]), "article", r.get("article"))
        s = m.sales[(r["buyout_date"], str(r["marketplace_sku"]))]
        s["buy_qty"] += num(r.get("buyouts_qty"))
        s["buy_seller"] += num(r.get("buyouts_amount_seller"))
        s["buy_buyer"] += num(r.get("buyouts_amount_buyer"))
    for r in funnel:
        if r["day"] > day_to:
            continue
        s = m.sales[(r["day"], str(r["nm_id"]))]
        s["funnel_orders"] += num(r.get("order_count"))
        s["funnel_buyouts"] += num(r.get("buyout_count"))
        s["funnel_cancels"] += num(r.get("cancel_count"))
        st = m.stock[(r["day"], str(r["nm_id"]))]
        st["funnel_stocks_wb"] = st.get("funnel_stocks_wb", 0) + (r.get("stocks_wb") or 0)
    for r in load_raw(cab, "wb_sales"):
        d = sale_day_msk(r.get("sale_dt"))
        if not d or d < PERIOD_FROM or d > day_to:
            continue
        sign = 1 if r["seller_oper_name"] == "Продажа" else -1
        it = m.item(r["nm_id"])
        _fill(it, "article", (r.get("vendor_code") or "").upper())
        _fill(it, "brand_src", r.get("brand_name"))
        _fill(it, "subject", r.get("subject_name"))
        if r.get("tech_size") not in (None, ""):
            it["sizes"].add(str(r["tech_size"]))
        s = m.sales[(d, str(r["nm_id"]))]
        s["report_qty"] += sign * num(r.get("quantity"))
        s["report_seller"] += sign * num(r.get("retail_price_with_disc"))
    return m


# ---------- книга ----------

def cmd_book(args):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font
    from loaders import rbh_categories as rbh_cat

    day_to = args.date_to or last_full_day()
    parts = [build_ozon(c, day_to) for c in CABINETS] + [build_wb(c, day_to) for c in CABINETS]
    parts.sort(key=lambda m: (m.cab, m.mp))
    stamp = datetime.now(MSK).strftime("%Y%m%d")

    # бренды, категории, «в каких ещё кабинетах»
    by_art, by_bar, wb_brand_by_art, wb_brand_by_bar = defaultdict(set), defaultdict(set), {}, {}
    for m in parts:
        for it in m.items.values():
            a = art_norm(it["article"])
            if a:
                by_art[a].add(m.label)
            for b in it["barcodes"]:
                by_bar[b].add(m.label)
            if m.mp == "wb" and it["brand_src"]:
                if a:
                    wb_brand_by_art.setdefault((m.cab, a), it["brand_src"])
                    wb_brand_by_art.setdefault(("*", a), it["brand_src"])
                for b in it["barcodes"]:
                    wb_brand_by_bar.setdefault(b, it["brand_src"])
    brand_variants = defaultdict(Counter)
    for m in parts:
        for it in m.items.values():
            a = art_norm(it["article"])
            if m.mp == "wb":
                it["brand"], it["brand_how"] = it["brand_src"], "карточка WB" if it["brand_src"] else "нет в базе"
            else:
                b, how = None, ""
                if a and (m.cab, a) in wb_brand_by_art:
                    b, how = wb_brand_by_art[(m.cab, a)], "временно: WB того же кабинета по артикулу"
                elif a and ("*", a) in wb_brand_by_art:
                    b, how = wb_brand_by_art[("*", a)], "временно: WB другого кабинета по артикулу"
                else:
                    for bc in sorted(it["barcodes"]):
                        if bc in wb_brand_by_bar:
                            b, how = wb_brand_by_bar[bc], "временно: WB по баркоду"
                            break
                it["brand_tmp"] = b or ""
                if it.get("brand_api"):
                    it["brand"], it["brand_how"] = it["brand_api"], "карточка Ozon (атрибут «Бренд», API)"
                else:
                    it["brand"], it["brand_how"] = (b or ""), (how or f"нет (в базе Ozon «{it['brand_src'] or '—'}»)")
            it["brand_n"] = brand_norm(it["brand"])
            if it["brand"]:
                brand_variants[it["brand_n"]][it["brand"]] += 1
            it["group"], it["category"] = rbh_cat.label(it["subject"], it["name"])
            others = set(by_art.get(a, set()))
            for b in it["barcodes"]:
                others |= by_bar.get(b, set())
            it["elsewhere"] = ", ".join(sorted(others - {m.label}))

    wb_ = Workbook(write_only=True)
    bold = Font(bold=True)

    def sheet(title, header, rows, widths=None, note=None):
        ws = wb_.create_sheet(title)
        ws.freeze_panes = "A3" if note else "A2"
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[_col(i)].width = w
        if note:
            c = WriteOnlyCell(ws, value=note)
            c.font = Font(italic=True)
            ws.append([c])
        hdr = []
        for h in header:
            c = WriteOnlyCell(ws, value=h)
            c.font = bold
            c.alignment = Alignment(wrap_text=True, vertical="top")
            hdr.append(c)
        ws.append(hdr)
        n = 0
        for r in rows:
            ws.append(r)
            n += 1
        counts[title] = n
        return ws

    counts = {}
    D = date.fromisoformat
    R = lambda v: round(v, 2)  # noqa: E731

    # 0. Как читать — заполняется в конце (нужны числа строк), лист создаётся первым
    ws_read = wb_.create_sheet("Как читать")
    ws_src = wb_.create_sheet("Источники и сверки")

    # 1. Справочник SKU
    def catalog_rows():
        for m in parts:
            for sku, it in sorted(m.items.items(), key=lambda kv: (kv[1]["article"] or "", kv[0])):
                yield [m.label, sku, it["article"], ", ".join(sorted(it["barcodes"])), ", ".join(sorted(it["sizes"])),
                       ", ".join(sorted(it["chrts"])), it["name"], it["brand"], it["brand_n"], it["brand_how"], it["subject"],
                       it["group"], it["category"],
                       ("архив" if it["archived"] else "не архив") if it["archived"] is not None else ("" if m.mp == "ozon" else "нет в базе"),
                       it["status"], it["elsewhere"]]
    sheet("Справочник SKU",
          ["Кабинет", "SKU Ozon / nmID WB", "Артикул продавца", "Баркоды", "Размер / оттенок (WB techSize)", "chrtID (WB)",
           "Название", "Бренд (исходный)", "Бренд (нормализованный)", "Источник бренда", "Категория площадки (тип Ozon / предмет WB)",
           "Группа (таксономия РБХ)", "Категория (таксономия РБХ)", "Архив", "Статус карточки Ozon", "Этот артикул / баркод есть ещё в"],
          catalog_rows(), [13, 14, 16, 18, 10, 12, 50, 18, 18, 30, 24, 22, 26, 10, 14, 26])

    brand_rows = [[b, sum(v.values()), "; ".join(f"{k} ({n})" for k, n in v.most_common()), "да" if len(v) > 1 else ""]
                  for b, v in sorted(brand_variants.items())]
    sheet("Бренды", ["Бренд нормализованный", "SKU всего (4 кабинета)", "Исходные написания (SKU)", "Склеено несколько написаний"],
          brand_rows, [26, 12, 60, 12])

    # 2. Остатки по дням (SKU × день)
    def stock_rows():
        for m in parts:
            universe = sorted({sku for (d, sku), st in m.stock.items() if any(v for k, v in st.items() if not k.endswith("_days"))}
                              | {sku for (d, sku), s in m.sales.items() if s.get("ord_qty") or s.get("cnl_qty") or s.get("buy_qty")})
            for d in m.days:
                loaded = d in m.stock_days
                for sku in universe:
                    st = m.stock.get((d, sku), {})
                    it = m.items.get(sku, {})
                    if m.mp == "ozon":
                        main = st.get("fbo_placement", 0) if loaded else None
                        if not loaded and not st:
                            continue
                        yield [m.label, D(d), sku, it.get("article", ""), main, st.get("fbo_stock"), st.get("fbs_stock"), None,
                               "да" if loaded else "нет"]
                    else:
                        main = st.get("wb_storage", 0) if loaded else None
                        if not loaded and not st:
                            continue
                        yield [m.label, D(d), sku, it.get("article", ""), main, None, None, st.get("wb_stock"), "да" if loaded else "нет"]
    sheet("Остатки по дням",
          ["Кабинет", "Дата", "SKU / nmID", "Артикул", "Остаток на складах МП, шт (Ozon — FBO из отчёта размещения; WB — штук на тарификации хранения)",
           "Ozon FBO, stock_daily (с 22–23.10)", "Ozon FBS, stock_daily (с 22–23.10)", "WB, stock_daily (с 22–23.10)",
           "Отчёт остатков дня загружен"],
          stock_rows(), [13, 11, 13, 14, 22, 14, 14, 14, 12],
          note="Пусто в колонке остатка МП = отчёт дня не загружен (это не ноль). 0 при загруженном отчёте = товара на складах МП не было. "
               "Ozon по складам — отдельные файлы …_skladi_<кабинет>.xlsx.")

    # 3. Дни без остатка — SKU × месяц
    def nostock_rows():
        for m in parts:
            skus = sorted({sku for (d, sku) in m.stock} | {sku for (d, sku) in m.sales})
            for sku in skus:
                it = m.items.get(sku, {})
                for mon, mon_label in MONTHS:
                    mdays = [d for d in m.days if d.startswith(mon)]
                    loaded = [d for d in mdays if d in m.stock_days]
                    key = "fbo_placement" if m.mp == "ozon" else "wb_storage"
                    with_st = [d for d in loaded if m.stock.get((d, sku), {}).get(key, 0) > 0]
                    zero = [d for d in loaded if d not in set(with_st)]
                    created = lambda d: m.sales.get((d, sku), {}).get("ord_qty", 0) + m.sales.get((d, sku), {}).get("cnl_qty", 0)  # noqa: E731
                    conf = lambda d: m.sales.get((d, sku), {}).get("ord_qty", 0)  # noqa: E731
                    buy = lambda d: m.sales.get((d, sku), {}).get("buy_qty", 0)  # noqa: E731
                    tot_created = sum(created(d) for d in mdays)
                    if not with_st and not tot_created and not sum(buy(d) for d in mdays):
                        continue
                    c_st = sum(created(d) for d in with_st)
                    q_st = sum(conf(d) for d in with_st)
                    b_st = sum(buy(d) for d in with_st)
                    n = len(with_st)
                    yield [m.label, sku, it.get("article", ""), it.get("brand_n", ""), it.get("group", ""), it.get("category", ""), mon_label,
                           len(mdays), len(loaded), n, len(zero), tot_created, c_st, sum(created(d) for d in zero),
                           R(c_st / n) if n else None, R(q_st / n) if n else None, R(b_st / n) if n else None,
                           R(q_st / n * len(zero)) if n else None]
    sheet("Дни без остатка",
          ["Кабинет", "SKU / nmID", "Артикул", "Бренд", "Группа", "Категория", "Месяц", "Дней в периоде", "Дней с загруженным отчётом остатков",
           "Дней с остатком > 0", "Дней с остатком 0 (на всех складах МП)", "Заказы созданные, шт — весь месяц", "Заказы созданные, шт — в дни с остатком",
           "Заказы созданные, шт — в дни без остатка", "Среднее в день с остатком: заказы созданные, шт", "Среднее в день с остатком: заказы подтверждённые, шт",
           "Среднее в день с остатком: выкупы, шт", "Оценка упущенных подтверждённых заказов, шт = среднее × дней без остатка"],
          nostock_rows(), [13, 13, 14, 18, 20, 22, 14, 8, 9, 9, 9, 10, 10, 10, 10, 10, 10, 12],
          note="Остаток МП: Ozon — FBO (отчёт размещения), WB — штуки на тарификации хранения. Заказы в дни без остатка у Ozon — обычно FBS "
               "(со склада продавца). Выкупы — по дню продажи, не по дню заказа. Оценка упущенного — грубая: не учитывает акции и сезон.")

    # 4. Продажи по дням
    def sales_rows():
        for m in parts:
            for (d, sku) in sorted(m.sales):
                s = m.sales[(d, sku)]
                if not any(v for k, v in s.items()):
                    continue
                it = m.items.get(sku, {})
                created = s["ord_qty"] + s["cnl_qty"]
                row = [m.label, D(d), sku, it.get("article", ""), created, s["cnl_qty"], s["ord_qty"], R(s["ord_seller"]),
                       R(s["ord_buyer"]) if not s.get("buyer_unmeasured") else None, R(s["cnl_seller"]),
                       s["buy_qty"], R(s["buy_seller"]), R(s["buy_buyer"])]
                if m.mp == "ozon":
                    row += [s["fbs_created_qty"], s["buy_positions"], None, None, None, None, None]
                else:
                    row += [None, None, s["funnel_orders"], s["funnel_buyouts"], s["funnel_cancels"], s["report_qty"], R(s["report_seller"])]
                yield row
    sheet("Продажи по дням",
          ["Кабинет", "Дата", "SKU / nmID", "Артикул", "Заказы созданные, шт", "из них отменённые, шт", "Заказы подтверждённые, шт",
           "Заказы подтверждённые, ₽ продавца", "Заказы подтверждённые, ₽ покупателя (пусто — не измерено)", "Отменённые, ₽ продавца",
           "Выкупы, шт (Ozon — штуки, WB — продажи − возвраты)", "Выкупы, ₽ продавца", "Выкупы, ₽ покупателя",
           "Ozon: из созданных FBS, шт", "Ozon: выкупы, позиций", "WB воронка: заказали, шт", "WB воронка: выкупили, шт (по дню заказа)",
           "WB воронка: отменили и вернули, шт", "WB отчёт реализации: продажи − возвраты, шт (по saleDt МСК)", "WB отчёт реализации: ₽ продавца"],
          sales_rows(), [13, 11, 13, 14] + [11] * 16,
          note="Заказы — по дню заказа (Ozon: МСК-сутки создания отправления; WB: дата заказа). Выкупы — по дню продажи. "
               "«Подтверждённые» = созданные − отменённые на момент последнего сбора.")

    # 5. Месяцы: итоги и фактические цены по SKU
    def month_rows():
        for m in parts:
            agg = defaultdict(lambda: defaultdict(float))
            for (d, sku), s in m.sales.items():
                for k, v in s.items():
                    agg[(d[:7], sku)][k] += v
            for (mon, sku) in sorted(agg):
                if mon not in ("2026-09", "2026-10"):
                    continue
                a = agg[(mon, sku)]
                it = m.items.get(sku, {})
                div = lambda x, y: R(x / y) if y else None  # noqa: E731
                yield [m.label, sku, it.get("article", ""), it.get("brand_n", ""), it.get("group", ""), it.get("category", ""),
                       "сентябрь" if mon == "2026-09" else f"октябрь (1–{date.fromisoformat(day_to).day})",
                       a["ord_qty"] + a["cnl_qty"], a["ord_qty"], R(a["ord_seller"]), R(a["ord_buyer"]) if not a.get("buyer_unmeasured") else None,
                       div(a["ord_seller"], a["ord_qty"]), div(a["ord_buyer"], a["ord_qty"]) if not a.get("buyer_unmeasured") else None,
                       a["buy_qty"], R(a["buy_seller"]), R(a["buy_buyer"]), div(a["buy_seller"], a["buy_qty"]), div(a["buy_buyer"], a["buy_qty"])]
    sheet("Месяцы и цены",
          ["Кабинет", "SKU / nmID", "Артикул", "Бренд", "Группа", "Категория", "Месяц", "Заказы созданные, шт", "Заказы подтверждённые, шт",
           "Заказы подтв., ₽ продавца", "Заказы подтв., ₽ покупателя", "Средняя цена заказа продавца, ₽/шт", "Средняя цена заказа покупателя, ₽/шт",
           "Выкупы, шт", "Выкупы, ₽ продавца", "Выкупы, ₽ покупателя", "Средняя цена выкупа продавца, ₽/шт", "Средняя цена выкупа покупателя, ₽/шт"],
          month_rows(), [13, 13, 14, 18, 20, 22, 14] + [11] * 11,
          note="Цена продавца — до соинвеста площадки; цена покупателя — что заплатил покупатель. Пусто у Ozon в цене покупателя — в месяце есть "
               "заказы FBO без измеренной цены покупателя (отчёт ЛК), среднее не считается, чтобы не занизить.")

    # итоги кабинетов по месяцам
    def total_rows():
        for m in parts:
            agg = defaultdict(lambda: defaultdict(float))
            for (d, sku), s in m.sales.items():
                for k, v in s.items():
                    agg[d[:7]][k] += v
            for mon, mon_label in MONTHS:
                a = agg[mon]
                yield [m.label, mon_label if mon != "2026-10" else f"октябрь (1–{date.fromisoformat(day_to).day})",
                       a["ord_qty"] + a["cnl_qty"], a["ord_qty"], R(a["ord_seller"]), a["buy_qty"], R(a["buy_seller"]), R(a["buy_buyer"]),
                       len({sku for (d, sku), s in m.sales.items() if d.startswith(mon) and (s.get("buy_qty") or 0) > 0})]
    sheet("Итоги кабинетов",
          ["Кабинет", "Месяц", "Заказы созданные, шт", "Заказы подтверждённые, шт", "Заказы подтв., ₽ продавца", "Выкупы, шт", "Выкупы, ₽ продавца",
           "Выкупы, ₽ покупателя", "SKU с выкупами"], total_rows(), [13, 16] + [14] * 7)

    # 6. Партии WB
    def batch_rows():
        for m in parts:
            if m.mp != "wb":
                continue
            sup = {}
            for x in (Api(m.cab).cached("wb_supplies.json") or {}).get("supplies") or []:
                if x.get("supplyID"):
                    sup[str(x["supplyID"])] = _d10(x.get("factDate")) or _d10(x.get("supplyDate"))
            for (nm, chrt, size, bc, gi), b in sorted(m.batches.items()):
                if b["count"] <= 0:
                    continue
                first = m.gi_first.get((nm, gi))
                it = m.items.get(nm, {})
                yield [m.label, nm, it.get("article", ""), it.get("brand_n", ""), chrt, size, bc, "WB (одна строка «Склад WB РФ»)",
                       int(gi) if gi.isdigit() else gi, b["count"], ", ".join(sorted(b["calc"])), ", ".join(sorted(b["fix"])),
                       ", ".join(sorted(b["lower"])), D(first) if first else None,
                       ("≥ " if first == PERIOD_FROM else "") + str((D(day_to) - D(first)).days) if first else None,
                       sup.get(gi), (D(day_to) - sup[gi]).days if sup.get(gi) else None]
    sheet("Партии WB",
          ["Кабинет", "nmID", "Артикул", "Бренд", "chrtID", "Размер", "Баркод", "Склад", "Поставка (giId; 0 — без поставки)",
           f"Штук на {day_to}", "Вид расчёта хранения", "tariffFixDate", "tariffLowerDate", "Партия в хранении с (по нашим данным, ≥ 23.08)",
           "Возраст по нашим данным, дней", "Дата поставки (supplies API: факт, иначе план)", "Возраст от даты поставки, дней"],
          batch_rows(), [13, 12, 14, 18, 12, 8, 15, 18, 14, 9, 30, 12, 12, 14, 10, 14, 12],
          note="Отчёт платного хранения WB отдаёт склад одной строкой «Склад WB РФ» (officeId 0) — разбивки по складам в нём нет. "
               "originalDate во всех строках = дню отчёта (это дата перерасчёта, не поставки). tariffFixDate — дата фиксации тарифа (гипотеза: "
               "не дата поставки; сверка с датой поставки — часть B).")

    # часть B — листы из сырья API (кабинет без сырья — без строк, причина в «Источниках»)
    api_sheets(sheet, parts, day_to, D, R)

    # 12. Склад РБХ — заготовка
    sheet("Склад РБХ", ["Артикул", "Баркод", "Бренд", "Остаток, шт", "Свободный остаток, шт", "Сопоставлено с кабинетами", "SKU / nmID"], [],
          [16, 16, 18, 10, 10, 30, 20], note="Заготовка: файл остатков из учётной системы РБХ даст владелец; свести по артикулу / баркоду со справочником.")

    # сверки
    recon = reconcile(parts, day_to)
    ws_src.column_dimensions["A"].width = 40
    for row in recon:
        ws_src.append(row)
    ws_read.column_dimensions["A"].width = 120
    for line in read_me(day_to, counts):
        ws_read.append([line])

    path = os.path.join(out_dir(), f"rbh_sales_plan_nov_{stamp}.xlsx")
    wb_.save(path)
    print(f"книга: {path}")
    for k, v in counts.items():
        print(f"  {k}: {v} строк")

    # остатки Ozon по складам — по файлу на базу (строк больше, чем помещается рядом)
    for cab in CABINETS:
        wbk = Workbook(write_only=True)
        ws = wbk.create_sheet("Ozon остатки по складам")
        ws.freeze_panes = "A2"
        ws.append(["Дата", "SKU", "Артикул", "Склад", "Штук (units)", "Из них платных (paid_units)", "Тип товара"])
        n = 0
        for r in load_raw(cab, "placement"):
            if r["day"] <= day_to and (r.get("units") or 0) > 0:
                ws.append([D(r["day"]), int(r["sku"]), r.get("offer_id"), r["warehouse"], r.get("units"), r.get("paid_units"), r.get("product_type")])
                n += 1
        ws2 = wbk.create_sheet("WB хранение по партиям")
        ws2.freeze_panes = "A2"
        ws2.append(["Дата", "nmID", "chrtID", "Размер", "Баркод", "Артикул", "Поставка (giId)", "Штук на тарификации", "Вид расчёта"])
        n2 = 0
        for r in load_raw(cab, "wb_storage"):
            if r["day"] <= day_to and (r.get("barcodes_count") or 0) > 0:
                ws2.append([D(r["day"]), r["nm_id"], r["chrt_id"], r.get("size"), r.get("barcode"), r.get("vendor_code"), r["gi_id"],
                            r["barcodes_count"], r.get("calc_type")])
                n2 += 1
        p = os.path.join(out_dir(), f"rbh_sales_plan_nov_{stamp}_skladi_{cab}.xlsx")
        wbk.save(p)
        print(f"склады {CAB_LABEL[cab]}: {p} — Ozon {n} строк, WB {n2} строк")


def _d10(s):
    s = str(s or "")[:10]
    return date.fromisoformat(s) if re.match(r"\d{4}-\d{2}-\d{2}$", s) else None


def _msk(s):
    """ISO-время с зоной → дата и время МСК строкой (для слотов и дат акций)."""
    if not s:
        return None
    try:
        dt = datetime.fromisoformat(str(s).replace("Z", "+00:00"))
    except ValueError:
        return str(s)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(MSK).strftime("%Y-%m-%d %H:%M")


def api_sheets(sheet, parts, day_to, D, R):
    """Листы части B: цены сейчас, товары в пути, остатки и оборачиваемость Ozon, акции ноября; дата поставки — в «Партиях WB»."""
    ozon = {m.cab: m for m in parts if m.mp == "ozon"}
    wbm = {m.cab: m for m in parts if m.mp == "wb"}
    pid2sku = {}
    for cab, m in ozon.items():
        for a in Api(cab).cached("ozon_attributes.json") or []:
            if a.get("sku"):
                pid2sku[(cab, int(a["id"]))] = str(a["sku"])
        for r in load_raw(cab, "ozon_products"):
            if r.get("product_id"):
                pid2sku.setdefault((cab, int(r["product_id"])), str(r["sku"]))

    def art(m, sku):
        return (m.items.get(str(sku)) or {}).get("article", "")

    def brand(m, sku):
        return (m.items.get(str(sku)) or {}).get("brand_n", "")

    # 10. Цены сейчас
    def price_rows():
        for cab, m in ozon.items():
            d = Api(cab).cached("ozon_prices.json")
            if not d:
                continue
            for x in d["items"]:
                if x.get("_visibility") == "ARCHIVED":
                    continue
                sku = pid2sku.get((cab, int(x["product_id"])), "")
                p = x.get("price") or {}
                com = x.get("commissions") or {}
                acts = [a for a in (x.get("marketing_actions") or {}).get("actions") or []
                        if str(a.get("date_from", ""))[:10] <= NOV_TO and str(a.get("date_to", "9999"))[:10] >= "2026-10-31"]
                yield [m.label, sku, x.get("offer_id"), brand(m, sku), "", num(p.get("old_price")), num(p.get("price")),
                       num(p.get("marketing_seller_price")), num(p.get("min_price")), None, None,
                       com.get("sales_percent_fbo"), com.get("sales_percent_fbs"),
                       "; ".join(f"{a.get('title')} ({a.get('value')}; {_msk(a.get('date_from'))} … {_msk(a.get('date_to'))})" for a in acts),
                       _msk(d["observed_at"])]
        for cab, m in wbm.items():
            d = Api(cab).cached("wb_prices.json")
            if not d:
                continue
            for g in d["items"]:
                for s in g.get("sizes") or [{}]:
                    nm = str(g.get("nmID"))
                    yield [m.label, nm, g.get("vendorCode"), brand(m, nm), s.get("techSizeName"), num(s.get("price")), num(s.get("discountedPrice")),
                           None, None, g.get("discount"), num(s.get("clubDiscountedPrice")) if s.get("clubDiscountedPrice") is not None else None,
                           None, None, "", _msk(d["observed_at"])]
    sheet("Цены сейчас",
          ["Кабинет", "SKU / nmID", "Артикул", "Бренд", "Размер (WB)", "Цена до скидки, ₽ (Ozon old_price / WB price)",
           "Цена на витрине, ₽ (Ozon price / WB discountedPrice)", "Ozon: цена с акциями продавца (marketing_seller_price)", "Ozon: минимальная цена",
           "WB: скидка продавца, %", "WB: цена со скидкой WB Клуба", "Ozon: комиссия FBO, %", "Ozon: комиссия FBS, %",
           "Ozon: акции продавца, действующие в ноябре (значение; период МСК)", "Снято (МСК)"],
          price_rows(), [13, 13, 14, 18, 8, 12, 12, 12, 12, 9, 12, 9, 9, 60, 16],
          note="Цены на момент выгрузки. СПП WB API не отдаёт — не берём. Ozon: «цена на витрине» — цена продавца со скидкой, без скидок Ozon "
               "за свой счёт; архивные карточки не включены.")

    # 8. Товары в пути
    def transit_rows():
        for cab, m in ozon.items():
            d = Api(cab).cached("ozon_supplies.json")
            if not d:
                continue
            for o in d["orders"]:
                if d["ids"].get(str(o.get("order_id"))) != "active":
                    continue
                ts = ((o.get("timeslot") or {}).get("timeslot") or {})
                for sp in o.get("supplies") or []:
                    wh = (sp.get("storage_warehouse") or {}).get("name") or ""
                    for it in d["bundles"].get(sp.get("bundle_id") or "", []):
                        sku = str(it.get("sku"))
                        yield [m.label, "заявка на поставку", o.get("order_number"), str(o.get("order_id")), str(sp.get("supply_id") or ""),
                               _st(o.get("state")), _st(sp.get("state")), _msk(o.get("created_date")),
                               _msk(ts.get("from") or ts.get("from_in_timezone")), _msk(ts.get("to") or ts.get("to_in_timezone")),
                               (o.get("dropoff_warehouse") or {}).get("name"), wh, sku, it.get("offer_id"), brand(m, sku), it.get("barcode"),
                               it.get("quantity"), None, ozon_stage(o.get("state"))]
        for cab, m in wbm.items():
            d = Api(cab).cached("wb_supplies.json")
            if not d:
                continue
            for s in d["supplies"]:
                sid = str(s.get("supplyID") or s.get("preorderID"))
                if sid not in d["goods"]:
                    continue
                det = d["details"].get(sid) or {}
                for g in d["goods"][sid]:
                    nm = str(g.get("nmID"))
                    yield [m.label, "поставка FBW", str(s.get("supplyID") or ""), str(s.get("preorderID") or ""), "",
                           WB_STATUS.get(s.get("statusID"), s.get("statusID")), "", _msk(s.get("createDate")), _msk(s.get("supplyDate")), "",
                           "", det.get("warehouseName") or "", nm, g.get("vendorCode"), brand(m, nm), g.get("barcode"),
                           g.get("quantity"), g.get("acceptedQuantity"), WB_STAGE.get(s.get("statusID"), "")]
    sheet("Товары в пути",
          ["Кабинет", "Вид", "Номер заявки / ID поставки WB", "ID заявки Ozon / ID заказа WB", "ID поставки Ozon", "Статус заявки", "Статус поставки",
           "Создана (МСК)", "Слот с / плановая дата (МСК)", "Слот по (МСК)", "Пункт отгрузки", "Склад размещения", "SKU / nmID", "Артикул", "Бренд",
           "Баркод", "Штук в заявке / поставке", "WB: принято, шт", "Стадия (для плана)"],
          transit_rows(), [13, 16, 18, 16, 14, 22, 18, 16, 16, 16, 22, 22, 13, 14, 18, 15, 10, 9, 28],
          note="Ozon — заявки не в статусах «Завершена / Отменена / Просрочена / Отказано» (/v3/supply-order/list, /get, /v1/supply-order/bundle); "
               "WB — поставки FBW не в статусе «Принято» (supplies API). Штуки «в пути» и «в заявках» по SKU Ozon — ещё на листе «Ozon остатки сейчас».")

    # 9. Ozon: остатки сейчас, оборачиваемость, возраст (FIFO)
    def turnover_rows():
        for cab, m in ozon.items():
            an = Api(cab).cached("ozon_analytics_stocks.json")
            tu = Api(cab).cached("ozon_turnover.json")
            st4 = Api(cab).cached("ozon_stocks.json")
            sup = Api(cab).cached("ozon_supplies.json")
            if not an and not tu:
                continue
            agg = defaultdict(lambda: defaultdict(float))
            grade = {}
            for x in (an or {}).get("items") or []:
                a = agg[str(x["sku"])]
                for k in ("available_stock_count", "valid_stock_count", "transit_stock_count", "requested_stock_count",
                          "return_from_customer_stock_count", "excess_stock_count", "other_stock_count"):
                    a[k] += x.get(k) or 0
                a["ads"] = max(a["ads"], x.get("ads") or 0)
                a["dws"] = max(a["dws"], x.get("days_without_sales") or 0)
                if (x.get("available_stock_count") or 0) > 0:
                    a["wh_with_stock"] += 1
                grade[str(x["sku"])] = x.get("turnover_grade")
            tmap = {str(x["sku"]): x for x in (tu or {}).get("items") or []}
            fbs = defaultdict(float)
            for p in (st4 or {}).get("items") or []:
                for s in p.get("stocks") or []:
                    if s.get("type") in ("fbs", "rfbs"):
                        fbs[str(s.get("sku"))] += s.get("present") or 0
            fifo = fifo_age(sup, {sku: a["available_stock_count"] + a["valid_stock_count"] for sku, a in agg.items()}, day_to) if sup else {}
            for sku in sorted(set(agg) | set(tmap)):
                a, t, f = agg.get(sku, {}), tmap.get(sku, {}), fifo.get(sku, {})
                yield [m.label, sku, art(m, sku), brand(m, sku), a.get("available_stock_count"), a.get("valid_stock_count"),
                       a.get("transit_stock_count"), a.get("requested_stock_count"), a.get("return_from_customer_stock_count"),
                       a.get("wh_with_stock"), fbs.get(sku), R(a["ads"]) if a.get("ads") else None, a.get("dws"), grade.get(sku),
                       t.get("current_stock"), t.get("ads"), t.get("idc"), t.get("turnover"), t.get("idc_grade"), t.get("turnover_grade"),
                       f.get("covered"), D(f["oldest"]) if f.get("oldest") else None, f.get("avg_age"), f.get("note", "FIFO-допущение")]
    sheet("Ozon остатки сейчас",
          ["Кабинет", "SKU", "Артикул", "Бренд", "Доступно к продаже, шт", "Готовим к продаже, шт", "В пути (поставки), шт", "В заявках на поставку, шт",
           "Возвраты от покупателей в пути, шт", "Складов с остатком", "FBS (склад продавца), шт", "Продажи в день за 28 дн. (аналитика остатков)",
           "Дней без продаж", "Статус ликвидности", "Оборачиваемость: остаток, шт", "Оборачиваемость: продажи в день за 60 дн.",
           "Хватит на, дней", "Оборачиваемость, дней", "Уровень остатка", "Уровень оборачиваемости",
           "FIFO: штук остатка, покрытых завершёнными поставками 2026", "FIFO: самая старая поставка в остатке", "FIFO: средний возраст остатка, дней",
           "Как считан возраст"],
          turnover_rows(), [13, 13, 14, 18] + [10] * 19 + [30],
          note="Ozon не отдаёт возраст партий. Возраст — FIFO-допущение: текущий остаток (доступно + готовим) раскладывается на последние завершённые "
               "поставки 2026 г. (/v3/supply-order, состав /v1/supply-order/bundle) от новых к старым; дата поставки — конец слота. "
               "По складам — сырьё /v1/analytics/stocks в data/<кабинет>/sales_plan_raw/api/.")

    # 11. Акции ноября
    def promo_rows():
        for cab, m in ozon.items():
            d = Api(cab).cached("ozon_actions.json")
            if not d:
                continue
            for a in d["november"]:
                per = d["products"].get(str(a["id"]), {})
                for kind, label in (("products", "участвует"), ("candidates", "кандидат")):
                    for p in per.get(kind) or []:
                        sku = pid2sku.get((cab, int(p["id"])), "")
                        price, ap = num(p.get("price")), num(p.get("action_price"))
                        if kind == "candidates" and not ap:
                            ap = num(p.get("max_action_price"))
                        yield [m.label, str(a["id"]), a.get("title"), a.get("action_type"), _msk(a.get("date_start")), _msk(a.get("date_end")),
                               sku, art(m, sku), brand(m, sku), label, price, ap, R((1 - ap / price) * 100) if price and ap else None,
                               num(p.get("max_action_price")) or None, p.get("add_mode")]
        for cab, m in wbm.items():
            d = Api(cab).cached("wb_calendar.json")
            if not d:
                continue
            info = {p["id"]: p for p in d["promotions"]}
            info.update({p["id"]: dict(info.get(p["id"], {}), **p) for p in d["details"]})
            for pid, per in d["nomenclatures"].items():
                p = info.get(int(pid), {})
                if per.get("skipped"):
                    yield [m.label, pid, p.get("name"), p.get("type"), _msk(p.get("startDateTime")), _msk(p.get("endDateTime")), "", "", "",
                           per["skipped"], None, None, None, None, None]
                    continue
                for flag, label in (("true", "участвует"), ("false", "кандидат")):
                    for x in per.get(flag) or []:
                        nm = str(x.get("id"))
                        price, pp = num(x.get("price")), num(x.get("planPrice"))
                        yield [m.label, pid, p.get("name"), p.get("type"), _msk(p.get("startDateTime")), _msk(p.get("endDateTime")),
                               nm, art(m, nm), brand(m, nm), label, price, pp, x.get("planDiscount"), None, None]
    sheet("Акции ноября",
          ["Кабинет", "ID акции", "Акция", "Тип", "Начало (МСК)", "Конец (МСК)", "SKU / nmID", "Артикул", "Бренд", "Участвует / кандидат",
           "Текущая цена, ₽", "Акционная цена, ₽ (Ozon кандидат — максимальная допустимая)", "Скидка к текущей цене, %",
           "Ozon: максимальная цена по акции", "Ozon: как добавлен"],
          promo_rows(), [13, 12, 40, 16, 16, 16, 13, 14, 18, 12, 11, 12, 10, 11, 12],
          note="Ozon — акции /v1/actions, пересекающие ноябрь; участники /v1/actions/products, кандидаты /v1/actions/candidates. "
               "WB — календарь акций с датами в ноябре; товары автоакций метод не отдаёт.")


OZON_STATE_RU = {"DATA_FILLING": "заполнение данных", "READY_TO_SUPPLY": "готова к отгрузке", "ACCEPTED_AT_SUPPLY_WAREHOUSE": "принята на точке отгрузки",
                 "IN_TRANSIT": "в пути", "ACCEPTANCE_AT_STORAGE_WAREHOUSE": "приёмка на складе", "REPORTS_CONFIRMATION_AWAITING": "согласование актов",
                 "REPORT_REJECTED": "спор", "COMPLETED": "завершена", "REJECTED_AT_SUPPLY_WAREHOUSE": "отказано в приёмке", "CANCELLED": "отменена",
                 "OVERDUE": "просрочена"}
WB_STATUS = {1: "не запланировано", 2: "запланировано", 3: "отгрузка разрешена", 4: "идёт приёмка", 5: "принято", 6: "отгружено на воротах"}


OZON_STAGE = {"DATA_FILLING": "планируется (заполнение данных)", "READY_TO_SUPPLY": "планируется (слот назначен)",
              "ACCEPTED_AT_SUPPLY_WAREHOUSE": "в пути", "IN_TRANSIT": "в пути", "ACCEPTANCE_AT_STORAGE_WAREHOUSE": "на приёмке",
              "REPORTS_CONFIRMATION_AWAITING": "принята, согласование актов (обычно уже в остатке)", "REPORT_REJECTED": "принята, спор по акту"}
WB_STAGE = {1: "не запланирована (черновик, товар не отгружен)", 2: "планируется (дата назначена)", 3: "планируется (отгрузка разрешена)",
            4: "на приёмке", 6: "на воротах склада"}


def ozon_stage(state):
    s = str(state or "")
    return OZON_STAGE.get(s.replace("ORDER_STATE_", ""), s)


def _st(s):
    s = str(s or "")
    return OZON_STATE_RU.get(s.replace("ORDER_STATE_", "").replace("SUPPLY_STATE_", ""), s)


def supply_day(o):
    """Дата поставки Ozon: конец слота, иначе дата смены статуса."""
    ts = ((o.get("timeslot") or {}).get("timeslot") or {})
    return _d10(ts.get("to") or ts.get("from") or o.get("state_updated_date") or o.get("created_date"))


def fifo_age(sup, stock_by_sku, day_to):
    """FIFO-допущение: остаток SKU = последние завершённые поставки (от новых к старым)."""
    lots = defaultdict(list)
    for o in sup["orders"]:
        if sup["ids"].get(str(o.get("order_id"))) != "completed":
            continue
        d = supply_day(o)
        for sp in o.get("supplies") or []:
            for it in sup["bundles"].get(sp.get("bundle_id") or "", []):
                if d and it.get("quantity"):
                    lots[str(it["sku"])].append((d, it["quantity"]))
    end = date.fromisoformat(day_to)
    out = {}
    for sku, stock in stock_by_sku.items():
        if stock <= 0:
            continue
        left, covered, wsum, oldest = stock, 0, 0, None
        for d, q in sorted(lots.get(sku, []), reverse=True):
            take = min(q, left)
            covered += take
            wsum += take * (end - d).days
            oldest = d
            left -= take
            if left <= 0:
                break
        out[sku] = {"covered": covered, "oldest": oldest.isoformat() if oldest else None,
                    "avg_age": round(wsum / covered, 1) if covered else None,
                    "note": "FIFO-допущение" + ("" if left <= 0 else f"; {int(left)} шт старше поставок 2026 или без заявки в API")}
    return out


def _col(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


def reconcile(parts, day_to):
    """Сверки источников: строки листа «Источники и сверки» (и печать в консоль)."""
    out = [["Источники"], [""],
           ["Ozon заказы", "marketplace_orders (FBO /v3, FBS /v4 — отправления; подтверждённые и отменённые раздельно), день — МСК-сутки создания"],
           ["Ozon выкупы", "marketplace_buyouts из /v1/finance/accrual/by-day; штуки — buyouts_units (accrual/postings), позиции — buyouts_qty"],
           ["Ozon остатки FBO", "ozon_placement_by_product_daily.units — отчёт размещения по SKU × склад × день (с 23.08)"],
           ["Ozon остатки FBO + FBS", "stock_daily (/v4/product/info/stocks), только итог по схеме, с 22–23.10"],
           ["WB заказы", "marketplace_orders из statistics supplier/orders: без отмены и отменённые (isCancel) раздельно"],
           ["WB выкупы", "marketplace_buyouts из отчёта реализации (wb_sales_report_rows): продажи − возвраты по saleDt МСК"],
           ["WB остатки", "wb_paid_storage_daily.barcodes_count — штук на тарификации платного хранения (nmID × размер × поставка × день)"],
           ["WB воронка", "wb_funnel_products_daily: orderCount (по дню заказа), buyoutCount (по дню заказа), stocks.wb (на момент сбора)"],
           ["Справочник", "ozon_products, sku_catalog (Ozon); воронка, хранение, отчёт реализации (WB); категории — loaders/rbh_categories"],
           [""], ["Сверки"]]
    for m in parts:
        out.append([""])
        out.append([m.label])
        missing = [d for d in m.days if d not in m.stock_days]
        out.append(["Дней с загруженным отчётом остатков", f"{len(m.stock_days)} из {len(m.days)}", "нет: " + (", ".join(missing) or "—")])
        if m.mp == "ozon":
            out.append(["Ключ stock_daily и sku_catalog у Ozon РБХ — product_id; переведён в SKU по карточкам, без пары",
                        f"stock_daily {m.unmapped['stock_daily']} строк, sku_catalog {m.unmapped['sku_catalog']} строк"])
            out.append(["День", "FBO размещение, шт", "FBO stock_daily, шт", "FBS stock_daily, шт", "Разница FBO (stock_daily − размещение)",
                        "SKU с разницей FBO"])
            for d in m.days:
                rows = {sku: st for (dd, sku), st in m.stock.items() if dd == d}
                if not any("fbo_stock_days" in st or "fbs_stock_days" in st for st in rows.values()):
                    continue
                p = sum(st.get("fbo_placement", 0) for st in rows.values())
                f = sum(st.get("fbo_stock", 0) for st in rows.values())
                s = sum(st.get("fbs_stock", 0) for st in rows.values())
                diff_skus = sum(1 for st in rows.values() if abs(st.get("fbo_stock", 0) - st.get("fbo_placement", 0)) > 0.5)
                out.append([d, p, f, s, f - p, diff_skus])
        else:
            out.append(["День", "Хранение, шт", "stock_daily, шт", "Воронка stocks.wb, шт", "Разница (stock_daily − хранение)"])
            for d in m.days:
                rows = [st for (dd, sku), st in m.stock.items() if dd == d]
                if not any("wb_stock" in st for st in rows):
                    continue
                h = sum(st.get("wb_storage", 0) for st in rows)
                s = sum(st.get("wb_stock", 0) for st in rows)
                f = sum(st.get("funnel_stocks_wb", 0) for st in rows)
                out.append([d, h, s, f, s - h])
            out.append([""])
            out.append(["Месяц", "Заказы созданные (supplier/orders), шт", "Воронка orderCount, шт", "Разница", "Выкупы таблицы, шт",
                        "Отчёт реализации продажи − возвраты, шт", "Разница", "Воронка buyoutCount (по дню заказа), шт"])
            for mon, mon_label in MONTHS:
                a = defaultdict(float)
                for (d, sku), s in m.sales.items():
                    if d.startswith(mon):
                        for k, v in s.items():
                            a[k] += v
                created = a["ord_qty"] + a["cnl_qty"]
                out.append([mon_label, created, a["funnel_orders"], created - a["funnel_orders"], a["buy_qty"], a["report_qty"],
                            a["buy_qty"] - a["report_qty"], a["funnel_buyouts"]])
            days_diff = []
            for d in m.days:
                a = defaultdict(float)
                for (dd, sku), s in m.sales.items():
                    if dd == d:
                        for k, v in s.items():
                            a[k] += v
                if abs(a["ord_qty"] + a["cnl_qty"] - a["funnel_orders"]) > 0.5 or abs(a["buy_qty"] - a["report_qty"]) > 0.5:
                    days_diff.append([d, a["ord_qty"] + a["cnl_qty"], a["funnel_orders"], a["buy_qty"], a["report_qty"]])
            out.append(["Дней с разницей (заказы ≠ воронке или выкупы ≠ отчёту)", len(days_diff)])
            if days_diff:
                out.append(["День", "Заказы созданные", "Воронка orderCount", "Выкупы таблицы", "Отчёт реализации"])
                out.extend(days_diff)
    for r in out:
        print("  " + " | ".join("" if v is None else str(v) for v in r))
    return out


def read_me(day_to, counts):
    lines = [
        f"Данные для плана продаж РБХ на ноябрь — выгрузка {datetime.now(MSK).strftime('%d.%m.%Y %H:%M')} МСК, период {PERIOD_FROM} … {day_to} "
        "(вся история проекта по РБХ; последний полный день).",
        "Кабинеты: РБХ-1 Ozon, РБХ-1 WB, РБХ-2 Ozon, РБХ-2 WB — колонка «Кабинет» на каждом листе. SKU — SKU Ozon или nmID WB.",
        "Часть A — из базы mp-analytics (обращений к API 0). Часть B (бренды Ozon, товары в пути, цены, акции) — дополнением в этот же файл.",
        "",
        "Листы:",
        "• Справочник SKU — товар каждого кабинета: артикул, баркоды, название, бренд исходный и нормализованный (регистр и точки склеены), "
        "категория по таксономии РБХ, архив, в каких ещё кабинетах есть тот же артикул / баркод. Бренд Ozon в базе неверный "
        "(«Beautyhome.me» у всех) — временно подставлен бренд WB по артикулу / баркоду, источник в колонке «Источник бренда».",
        "• Бренды — какие написания склеены в один бренд.",
        "• Остатки по дням — SKU × день, остаток на складах маркетплейса. Ozon — FBO (отчёт размещения), WB — штуки на тарификации хранения. "
        "Пусто — отчёт дня не загружен; 0 — товара не было. Ozon по складам и WB по партиям — в файлах …_skladi_rbh1 / rbh2.",
        "• Дни без остатка — SKU × месяц: дней с остатком и без, средние продажи в день с остатком, грубая оценка упущенного.",
        "• Продажи по дням — SKU × день: заказы созданные / отменённые / подтверждённые (шт, ₽ продавца и покупателя), выкупы (шт, ₽); "
        "для WB — сверочные колонки воронки и отчёта реализации.",
        "• Месяцы и цены — SKU × сентябрь / октябрь: итоги и средние фактические цены заказа и выкупа (продавца и покупателя).",
        "• Итоги кабинетов — по месяцам.",
        f"• Партии WB — остаток на {day_to} по поставкам (giId): штуки, даты из отчёта хранения, с какого дня партия видна у нас.",
        "• Склад РБХ — заготовка под файл владельца.",
        "• Источники и сверки — откуда каждая цифра и как источники сходятся между собой.",
        "",
        "Определения:",
        "• Заказы созданные = подтверждённые + отменённые; «подтверждённые» — на момент последнего ночного сбора (свежие дни ещё дозревают: "
        "у Ozon отмены доезжают до ~21 дня).",
        "• Выкупы Ozon — штуки по начислениям (день продажи); выкупы WB — продажи − возвраты отчёта реализации (день продажи, МСК).",
        "• Цена продавца — до соинвеста площадки; цена покупателя — что заплатил покупатель (у Ozon FBO — из отчёта ЛК, не на всех днях).",
        "",
        "Строк на листах: " + "; ".join(f"{k} — {v}" for k, v in counts.items()),
    ]
    return lines


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("db")
    a.add_argument("--date-to")
    a.add_argument("--only", nargs="*")
    a.add_argument("--refresh", action="store_true")
    a.set_defaults(fn=cmd_db)
    c = sub.add_parser("api")
    c.add_argument("--mp", nargs="*", choices=("ozon", "wb"))
    c.add_argument("--steps", nargs="*")
    c.set_defaults(fn=cmd_api)
    b = sub.add_parser("book")
    b.add_argument("--date-to")
    b.set_defaults(fn=cmd_book)
    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
