#!/usr/bin/env python3
"""Лист «Ozon - <месяц>» по выкупам: то, что владелец считал руками, — из наших источников. Только чтение.

    venv/bin/python3 scripts/report_ozon_month.py --month 2026-09                       → data/reports/ozon_2026-09.xlsx
    venv/bin/python3 scripts/report_ozon_month.py --month 2026-09 --date-to 2026-09-15 \
        --check --xlsx data/manual_report_september.xlsx --sheet "Ozon - сентябрь"      приёмка против ручного листа

Лист 1 «Ozon - <месяц>» — по дням и «Итого». Что измеряет каждая колонка:

    Оборот            marketplace_buyouts.buyouts_amount_seller за день реализации, с НДС (как в ручном листе)
    Комиссия          marketplace_buyouts.commission_amount (sale_commission начислений), с НДС
    Выручка           (Оборот − Комиссия) / НДС
    Себестоимость     article_unit_costs (снимок 1С, --snapshot) × ПОЗИЦИИ выкупов; sku → article по заказам
    Логистика, Прочее, Подписка — типы начислений, свёрнутые в статьи через TYPE_TO_EXPENSE (классификация у читателя)
    Эквайринг         тип 1 / НДС (в marketplace_expenses он внутри статьи other)
    Подписка          статья subscription (типы 51, 52, 74) / НДС
    Реклама           типы 41 + 54 / НДС — «на круг», решение владельца 5 от 2026-09-19
    Компенсации       типы 25 + 10, ДОХОД: «+» — деньги нам; в расходы и в «Прочее» не входят (решение 8 от 2026-09-21)
    Фин. рез. с компенсациями   справочно; сопоставим с «Фин. рез.» владельца, у которого компенсации внутри «Прочего»
    ДРР %             Реклама / (Оборот / НДС) — формула ручного листа, «от ТО»
    Прочее            (статьи other + external_promo − тип 1) / НДС; незнакомая статья — сюда же И называется вслух
    Фин. рез.         Маржа − Логистика − Эквайринг − Подписка − Реклама − Прочее
    Ebitda            Фин. рез. − накладные в день (OVERHEAD_PER_DAY: параметр с датой действия, как НДС; Ozon с 2026-09-01 —
                      311 527,00, ячейка T83 ручного листа от 22.09); «Ebitda с компенсациями» сопоставима с «Ebitda» владельца
    справочно         «Реклама по образцу» = (41 + 54 + 96 + вся статья subscription) / НДС — так считает ручной лист
                      (проверено 2026-09-22: 09-20 их «Реклама» − наша = 24 990,00 / 1,22 — подписка сверх Premium);
                      «Логистика + Прочее по образцу» = (logistics + other − тип 1 − тип 96) / НДС —
                      у них это две строки одной группы, граница между ними проведена не по type_id

Листы «Заказы», «Заказы FBO», «Заказы FBS» — по дню ЗАКАЗА, прогноз без зашитых долей (ozon_orders_forecast.py):
создано / подтверждено / отменено из marketplace_orders; день от 21 суток — факт, моложе — подтверждено × (1 − доля,
которая ещё отменится) по кривой дозревания из ozon_posting_status_log; комиссия — измеренная доля за 30 дней по
площадке товара; прочие — доля леджера начислений в обороте выкупов за те же 30 дней; реклама — 41 + 54 дня заказа.
Справочно рядом — формулы владельца на его константах 0,65 / 0,59 / 0,024. Кривая — на листе «Кривая дозревания».
Лист не собрался — остальные выходят, причина пишется в книгу и вслух, код возврата 2. --no-orders — не собирать.

Лист 2 «По SKU» — daily_sku_kpi за месяц; реклама там из Performance API (решение 5), поэтому итог
рекламы двух листов различается — разница и причина написаны внизу листа.

Типы начислений — из базы: леджер ozon_accrual_daily_types (ночной шаг расходов пишет его из того же
ответа by-day). Дня в леджере нет — запасной путь: файл data/accrual_history/<день>.json, нет файла —
один сбор дня из API (1–3 обращения), для дат моложе двух суток файл НЕ сохраняется. --types-from files
возвращает прежний путь целиком; --check читает ОБА источника и требует совпадения по каждому типу и дате — кроме дат моложе
двух суток: их лист берёт живым ответом (начисления рекламы за D−1 доезжают после ночного сбора), файл не сохраняется.
marketplace_expenses в лист не идёт — только в сверку по статьям: расхождение называется по дням.

НДС — параметр с датой действия (VAT_RATES), не константа в формуле. Даты моложе двух суток
помечаются: начисления доезжают. В окне ночного прогона (loaders/pipeline_window.py) в API не идёт.
В БД не пишет: db_writes = 0.
"""
import argparse
import calendar
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from zoneinfo import ZoneInfo

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(ROOT, ".env"))
from loaders import ozon_finance_accrual as accrual  # noqa: E402
from reconcile_manual_report_ozon import read_manual_sheet, type_sums  # noqa: E402
import ozon_orders_forecast as forecast  # noqa: E402

Z = Decimal(0)
C = Decimal("0.01")
SNAP = "2026-05-20"
OUT_DIR = os.path.join("data", "reports")
RAW_DIR = os.path.join("data", "accrual_history")
YOUNG_DAYS = 2                     # сверка берёт дату не моложе двух суток: начисления доезжают (CLAUDE.md §2)
ACQUIRING_TYPE, PREMIUM_TYPE, REVIEWS_TYPE = 1, 51, 96
KNOWN_ARTICLES = {"logistics", "other", "subscription", "external_promo", "commission"}

# НДС с датой действия: (с какой даты, коэффициент). Лист декабря 2025 считался с 1,2.
VAT_RATES = (("2026-01-01", Decimal("1.22")), ("0001-01-01", Decimal("1.20")))
# Накладные в день по площадке, с датой действия. Источник — ручной лист владельца от 2026-09-22: на листе
# «Ozon - сентябрь» R = P − $T$83, T83 = 311 527,00 (у WB своя константа 74 002,00 — здесь не используется).
# Откуда число и как оно меняется — владелец скажет; до даты действия Ebitda пуста, не ноль.
OVERHEAD_PER_DAY = {"ozon": (("2026-09-01", Decimal("311527.00")),)}
# Индекс себестоимости — справочно, пока нет новой выгрузки 1С: наш снимок 05-20 против цен владельца за 1–21 сентября
# (лист 22.09): 27 401 217,62 / 23 824 962,63 = 1,150; по дням 1,175 (начало месяца) → 1,116 (21-е) — у него величина
# движется внутри месяца, по SKU из его файла не восстановить. С новым снимком 1С индекс с той даты — 1,00.
COST_INDEX = (("2026-09-01", Decimal("1.150")),)
COST_INDEX_STALE_PCT = Decimal("0.03")     # индекс по листу отличается от параметра больше — «протух», --check говорит вслух

MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]


def D(v):
    return Decimal(str(v or 0))


def q(v):
    return Decimal(v).quantize(C)


def vat_for(day):
    for valid_from, rate in VAT_RATES:
        if day >= valid_from:
            return rate
    raise ValueError(f"нет ставки НДС для {day}")


def overhead_for(day, platform="ozon"):
    """Накладные в день на дату или None, если ставка на эту дату не задана."""
    for valid_from, value in OVERHEAD_PER_DAY.get(platform, ()):
        if day >= valid_from:
            return value
    return None


def cost_index_for(day):
    """Индекс СС на дату или None, если на эту дату не задан."""
    for valid_from, value in COST_INDEX:
        if day >= valid_from:
            return value
    return None


def month_days(month, date_to=None):
    y, m = int(month[:4]), int(month[5:7])
    last = date(y, m, calendar.monthrange(y, m)[1])
    if date_to:
        last = min(last, date.fromisoformat(date_to))
    out, d = [], date(y, m, 1)
    while d <= last:
        out.append(d.isoformat()); d += timedelta(days=1)
    return out


def ratio(a, b):
    return (a / b) if b else None


def build_daily(days, buyouts, expenses, types_by_day, unit_cost, today):
    """Строки листа по дням. Чистая функция: на входе строки таблиц и суммы сырья по type_id.

    buyouts      [{buyout_date, marketplace_sku, buyouts_qty, buyouts_amount_seller, commission_amount}]
    expenses     [{expense_date, expense_type, expense_amount}]
    types_by_day {день: {type_id: расход, положителен при списании}}; дня нет — сырья нет
    unit_cost    sku -> Decimal | None
    """
    acc = {d: defaultdict(Decimal) for d in days}
    no_cost = defaultdict(int)
    unknown = defaultdict(Decimal)
    for r in buyouts:
        d = r["buyout_date"]
        if d not in acc:
            continue
        qty = D(r["buyouts_qty"])
        acc[d]["positions"] += qty
        acc[d]["turnover"] += D(r["buyouts_amount_seller"])
        acc[d]["commission"] += D(r["commission_amount"])
        # Себестоимость — по ШТУКАМ, где они измерены (buyouts_units не null), иначе по позициям. null — «не
        # измерено», а не ноль: 0 штук (продажа и возврат в один день) — измеренный ноль и даёт нулевую СС.
        units = r.get("buyouts_units")
        base = qty if units is None else D(units)
        acc[d]["rows_by_units" if units is not None else "rows_by_positions"] += 1
        acc[d]["units"] += base
        uc = unit_cost(r["marketplace_sku"])
        if uc is None:
            no_cost[d] += int(qty)
        else:
            acc[d]["cogs"] += base * uc
    for r in expenses:
        d, t, v = r["expense_date"], str(r["expense_type"] or ""), D(r["expense_amount"])
        if d not in acc:
            continue
        if t.startswith("advertising") or t == "commission":
            continue            # реклама Performance живёт на листе «По SKU»; комиссия уже взята из выкупов
        if t not in KNOWN_ARTICLES:
            unknown[t] += v     # незнакомая статья базы: в запасном пути (нет типов) идёт в прочее И называется вслух
            t = "other"
        acc[d]["exp_" + t] += v

    rows = []
    for d in days:
        a, vat, ts = acc[d], vat_for(d), types_by_day.get(d)
        row = {"date": d, "vat": vat, "has_raw": ts is not None,
               "young": (date.fromisoformat(today) - date.fromisoformat(d)).days < YOUNG_DAYS,
               "positions": a["positions"], "no_cost_positions": no_cost.get(d, 0), "units": a["units"],
               "rows_by_units": int(a["rows_by_units"]), "rows_by_positions": int(a["rows_by_positions"])}
        row["turnover"], row["commission"] = a["turnover"], a["commission"]
        row["revenue"] = (a["turnover"] - a["commission"]) / vat
        row["cogs"] = a["cogs"]
        row["margin"] = row["revenue"] - row["cogs"]
        if ts is None:
            # Типов за день нет ни в леджере, ни в файле: статьи берём из marketplace_expenses, но эквайринг от
            # прочего не отделить, рекламы из начислений и компенсаций нет — не ноль, а пусто.
            row["logistics"] = a["exp_logistics"] / vat
            row["subscription"] = a["exp_subscription"] / vat
            row["other"] = (a["exp_other"] + a["exp_external_promo"]) / vat
            for k in ("acquiring", "ads", "compensations", "ads_like_manual", "log_other_like_manual", "fin_result", "fin_result_with_comp",
                      "ebitda", "ebitda_with_comp", "fin_result_index", "ebitda_index"):
                row[k] = None
            row["overhead"] = overhead_for(d)
            row["cost_index"] = cost_index_for(d)
            row["cogs_index"] = None if row["cost_index"] is None else row["cogs"] * row["cost_index"]
            row["margin_index"] = None if row["cost_index"] is None else row["revenue"] - row["cogs_index"]
            rows.append(row)
            continue
        # Статьи — свёрткой типов через TYPE_TO_EXPENSE: классификация у читателя, леджер её не знает.
        articles = defaultdict(Decimal)
        for type_id, v in ts.items():
            if type_id in accrual.AD_TYPE_IDS or type_id in accrual.UNCLASSIFIED_TYPE_IDS:
                continue
            article = accrual.TYPE_TO_EXPENSE.get(type_id)
            if article is None:
                unknown[f"тип {type_id}"] += v      # незнакомый тип идёт в прочее И называется вслух
                article = "other"
            articles[article] += v
        acq = ts.get(ACQUIRING_TYPE, Z)
        ads = sum((ts.get(t, Z) for t in accrual.AD_TYPE_IDS), Z)
        # types_by_day держит расход положительным; компенсация — деньги НАМ, в леджере она > 0, здесь < 0
        comp = -sum((ts.get(t, Z) for t in accrual.UNCLASSIFIED_TYPE_IDS), Z)
        row["logistics"] = articles["logistics"] / vat
        row["subscription"] = articles["subscription"] / vat
        row["acquiring"] = acq / vat
        row["ads"] = ads / vat
        row["other"] = (articles["other"] + articles["external_promo"] - acq) / vat
        row["compensations"] = comp / vat
        # у владельца в «Рекламе» — вся статья подписки (51, 52, 74), не только Premium: 2026-09-20 разница ровно 24 990,00 / 1,22
        row["ads_like_manual"] = (ads + articles["subscription"] + ts.get(REVIEWS_TYPE, Z)) / vat
        row["log_other_like_manual"] = (articles["logistics"] + articles["other"] - acq - ts.get(REVIEWS_TYPE, Z)) / vat
        row["fin_result"] = row["margin"] - row["logistics"] - row["acquiring"] - row["subscription"] - row["ads"] - row["other"]
        row["fin_result_with_comp"] = row["fin_result"] + row["compensations"]
        row["overhead"] = overhead_for(d)
        row["ebitda"] = None if row["overhead"] is None else row["fin_result"] - row["overhead"]
        row["ebitda_with_comp"] = None if row["overhead"] is None else row["fin_result_with_comp"] - row["overhead"]
        # справочно по индексу СС: та же строка, где себестоимость = снимок × индекс
        row["cost_index"] = cost_index_for(d)
        idx = row["cost_index"]
        row["cogs_index"] = None if idx is None else row["cogs"] * idx
        row["margin_index"] = None if idx is None else row["revenue"] - row["cogs_index"]
        row["fin_result_index"] = None if idx is None else row["fin_result"] + row["cogs"] - row["cogs_index"]
        row["ebitda_index"] = None if idx is None or row["overhead"] is None else row["fin_result_index"] - row["overhead"]
        # Сверка двух таблиц одной базы: статья из типов против той же статьи в marketplace_expenses. Расходятся —
        # называем: upsert расходов строк не удаляет, и начисление, которое Ozon убрал, застревает в базе
        # (найдено 2026-09-21: 09-11, 09-12, 09-15, 09-19 — по одной строке). Лист считает по типам.
        row["db_minus_raw"] = {art: q(a["exp_" + art] - articles[art]) for art in ("logistics", "other", "subscription", "external_promo")
                               if q(a["exp_" + art]) != q(articles[art])}
        rows.append(row)
    return rows, dict(unknown)


MONEY = ["turnover", "commission", "revenue", "cogs", "margin", "logistics", "acquiring", "subscription", "ads", "other",
         "fin_result", "compensations", "fin_result_with_comp", "overhead", "ebitda", "ebitda_with_comp",
         "cogs_index", "margin_index", "fin_result_index", "ebitda_index",
         "ads_like_manual", "log_other_like_manual", "positions", "units"]


def add_ratios(row):
    rev = row.get("revenue") or Z
    row["commission_pct"] = ratio(row["commission"], row["turnover"])
    row["margin_pct"] = ratio(row["margin"], rev)
    row["logistics_pct"] = ratio(row["logistics"], rev)
    row["acquiring_pct"] = None if row.get("acquiring") is None else ratio(row["acquiring"], rev)
    row["drr_pct"] = None if row.get("ads") is None else ratio(row["ads"], row["turnover"] / row["vat"]) if row["turnover"] else None
    row["fin_result_pct"] = None if row.get("fin_result") is None else ratio(row["fin_result"], rev)
    row["fin_result_with_comp_pct"] = None if row.get("fin_result_with_comp") is None else ratio(row["fin_result_with_comp"], rev)
    row["ebitda_pct"] = None if row.get("ebitda") is None else ratio(row["ebitda"], rev)
    row["ebitda_with_comp_pct"] = None if row.get("ebitda_with_comp") is None else ratio(row["ebitda_with_comp"], rev)
    return row


def total_row(rows):
    """«Итого»: деньги — суммой, проценты — от сумм. Колонка пуста хоть за один день — в итоге она неполна, и это видно."""
    t = {"date": "Итого", "vat": rows[-1]["vat"] if rows else vat_for("2026-01-01"), "young": False, "has_raw": all(r["has_raw"] for r in rows),
         "no_cost_positions": sum(r["no_cost_positions"] for r in rows),
         "rows_by_units": sum(r.get("rows_by_units", 0) for r in rows), "rows_by_positions": sum(r.get("rows_by_positions", 0) for r in rows)}
    for k in MONEY:
        vals = [r.get(k) for r in rows]
        t[k] = None if any(v is None for v in vals) else sum(vals, Z)
    # ДРР в итоге — по формуле листа от сумм; при смене НДС внутри месяца оборот без НДС считается по дням
    t["_turnover_net"] = sum((r["turnover"] / r["vat"] for r in rows), Z)
    add_ratios(t)
    t["drr_pct"] = None if t["ads"] is None else ratio(t["ads"], t["_turnover_net"])
    return t


PLATFORMS = (("F", "Основная"), ("S", "Селект"), ("T", "Дискаунтер"))
NO_PLATFORM = "Без площадки"


def platform_of(article):
    """Площадка — по первой букве артикула: F — основная, S — Селект, T — Дискаунтер (reports_model.md §6)."""
    first = (str(article or "").strip()[:1]).upper()
    return dict(PLATFORMS).get(first, NO_PLATFORM)


# Сегмент по металлу — по названию товара (дыра 5 реестра docs/owner_workbook_gaps.md, лист владельца «в т.ч. Серебро - месяц»).
# Проверено 2026-09-23 по августу: серебро по признаку «серебр» или «925» в названии даёт оборот 01.08 142 567,00 и комиссию
# 59 791,90 — ровно строку его листа; всё прочее — золото («золот» или проба 585 / 375 / 750). SKU без названия — «Без признака».
METALS = (("серебр|925", "Серебро"), ("золот|585|375|750", "Золото"))
NO_METAL = "Без признака металла"


def metal_of(product_name):
    import re
    name = str(product_name or "")
    for pattern, label in METALS:
        if re.search(pattern, name, re.IGNORECASE):
            return label
    return NO_METAL


def build_metal_daily(days, buyouts, expenses, kpi_rows, sku2name, unit_cost):
    """{металл: [строки по дням]} — те же колонки, что у площадок, разрез — по названию товара из витрины."""
    names = [label for _p, label in METALS] + [NO_METAL]
    classify = lambda sku: metal_of(sku2name.get(str(sku or ""))) if str(sku or "") else NO_METAL  # noqa: E731
    return build_platform_daily(days, buyouts, expenses, kpi_rows, {}, unit_cost, classify=classify, names=names)


def build_platform_daily(days, buyouts, expenses, kpi_rows, sku2art, unit_cost, classify=None, names=None):
    """{площадка: [строки по дням]} — то, что делится по SKU из базы. classify / names — иной разрез теми же колонками (металл).

    Оборот, комиссия, выручка, СС, маржа — из выкупов по SKU: делятся точно. Логистика, подписка и «прочее
    с эквайрингом» — из marketplace_expenses по SKU (леджер начислений без SKU и по площадкам не делится,
    поэтому эквайринг здесь НЕ отделён от прочего, а компенсаций нет вовсе). Реклама — Performance по SKU
    (ad_spend витрины, решение 5). Расход без SKU и SKU без артикула — отдельная «площадка» «Без площадки»,
    чтобы сумма площадок равнялась целому, а не тихо теряла остаток.
    """
    names = names or [n for _p, n in PLATFORMS] + [NO_PLATFORM]
    rest = names[-1]
    acc = {n: {d: defaultdict(Decimal) for d in days} for n in names}
    where = classify or (lambda sku: platform_of(sku2art.get(str(sku or ""))) if str(sku or "") else NO_PLATFORM)
    for r in buyouts:
        d = r["buyout_date"]
        if d not in acc[rest]:
            continue
        a = acc[where(r["marketplace_sku"])][d]
        qty, units = D(r["buyouts_qty"]), r.get("buyouts_units")
        a["positions"] += qty
        a["turnover"] += D(r["buyouts_amount_seller"]); a["commission"] += D(r["commission_amount"])
        uc = unit_cost(r["marketplace_sku"])
        if uc is not None:
            a["cogs"] += (qty if units is None else D(units)) * uc
    for r in expenses:
        d, t = r["expense_date"], str(r["expense_type"] or "")
        if d not in acc[rest] or t == "commission" or t.startswith("advertising"):
            continue
        a = acc[where(r.get("marketplace_sku"))][d]
        a["logistics" if t == "logistics" else "subscription" if t == "subscription" else "other_with_acq"] += D(r["expense_amount"])
    for r in kpi_rows:
        d = r["kpi_date"]
        if d in acc[rest]:
            acc[where(r.get("marketplace_sku"))][d]["ads_perf"] += D(r.get("ad_spend"))
    out = {}
    for name in names:
        rows = []
        for d in days:
            a, vat = acc[name][d], vat_for(d)
            row = {"date": d, "vat": vat, "positions": a["positions"], "turnover": a["turnover"], "commission": a["commission"],
                   "revenue": (a["turnover"] - a["commission"]) / vat, "cogs": a["cogs"],
                   "logistics": a["logistics"] / vat, "subscription": a["subscription"] / vat, "other_with_acq": a["other_with_acq"] / vat,
                   "ads_perf": a["ads_perf"] / vat}
            row["margin"] = row["revenue"] - row["cogs"]
            row["fin_result"] = row["margin"] - row["logistics"] - row["subscription"] - row["other_with_acq"] - row["ads_perf"]
            row["margin_pct"] = ratio(row["margin"], row["revenue"])
            row["drr_pct"] = ratio(row["ads_perf"], row["turnover"] / vat) if row["turnover"] else None
            row["fin_result_pct"] = ratio(row["fin_result"], row["revenue"])
            rows.append(row)
        out[name] = rows
    return out


PLATFORM_MONEY = ("turnover", "commission", "revenue", "cogs", "margin", "logistics", "subscription", "other_with_acq", "ads_perf", "fin_result", "positions")


def platform_total(rows):
    t = {"date": "Итого"}
    for k in PLATFORM_MONEY:
        t[k] = sum((r[k] for r in rows), Z)
    net = sum((r["turnover"] / r["vat"] for r in rows), Z)
    t["margin_pct"], t["fin_result_pct"] = ratio(t["margin"], t["revenue"]), ratio(t["fin_result"], t["revenue"])
    t["drr_pct"] = ratio(t["ads_perf"], net)
    return t


def platform_addition(total, platform_totals):
    """Сложение площадок в общий лист: [(колонка, Σ площадок, общий лист, разница, объяснение)]."""
    s = lambda k: sum((t[k] for t in platform_totals.values()), Z)  # noqa: E731
    exact = [("Оборот", "turnover"), ("Комиссия", "commission"), ("Выручка", "revenue"), ("Себестоимость", "cogs"), ("Маржа", "margin")]
    out = [(title, s(k), total[k], "те же строки выкупов — обязана быть 0,00") for title, k in exact]
    same_source = "площадки — marketplace_expenses по SKU, общий лист — типы начислений; разница = строки, застрявшие в расходах (upsert не удаляет)"
    out.append(("Логистика", s("logistics"), total["logistics"], same_source))
    out.append(("Подписка", s("subscription"), total["subscription"], same_source))
    both = None if total.get("acquiring") is None else total["other"] + total["acquiring"]
    out.append(("Прочее + Эквайринг", s("other_with_acq"), both, same_source + "; по площадкам эквайринг от прочего не отделён — леджер без SKU"))
    out.append(("Реклама", s("ads_perf"), total.get("ads"), "площадки — Performance по SKU, общий лист — начисления 41 + 54: источники разные по решению 5, разница названа на листе «По SKU»"))
    return [(title, a, b, None if b is None else q(a - b), why) for title, a, b, why in out]


def build_sku(kpi_rows, sku2art, unit_cost, vat):
    """Лист «По SKU»: daily_sku_kpi за месяц, реклама — Performance API (ad_spend)."""
    by = defaultdict(lambda: defaultdict(Decimal))
    names = {}
    for r in kpi_rows:
        sku = str(r.get("marketplace_sku") or "")
        a = by[sku]
        a["positions"] += D(r.get("buyouts_qty")); a["turnover"] += D(r.get("buyouts_amount_seller"))
        a["commission"] += D(r.get("commission_amount")); a["ads_gross"] += D(r.get("ad_spend"))
        a["logistics_gross"] += D(r.get("logistics_amount")); a["other_gross"] += D(r.get("other_expenses_amount"))
        if sku not in names and (r.get("article") or r.get("product_name")):
            names[sku] = (r.get("article") or "", r.get("product_name") or "")
    out = []
    for sku, a in by.items():
        if not any(a.values()):
            continue
        uc = unit_cost(sku) if sku else None
        row = {"sku": sku or "(без SKU)", "article": sku2art.get(sku) or names.get(sku, ("", ""))[0], "name": names.get(sku, ("", ""))[1],
               "positions": a["positions"], "turnover": a["turnover"], "commission": a["commission"],
               "revenue": (a["turnover"] - a["commission"]) / vat,
               "cogs": None if (uc is None and a["positions"]) else a["positions"] * (uc or Z),
               "ads": a["ads_gross"] / vat, "logistics": a["logistics_gross"] / vat, "other": a["other_gross"] / vat}
        row["margin"] = None if row["cogs"] is None else row["revenue"] - row["cogs"]
        row["margin_pct"] = None if row["margin"] is None else ratio(row["margin"], row["revenue"])
        row["drr_pct"] = ratio(row["ads"], row["turnover"] / vat) if row["turnover"] else None
        out.append(row)
    out.sort(key=lambda r: (-r["turnover"], r["sku"]))
    return out


# ---------- ввод-вывод ----------

def fetch(sb, table, select, filters, orders):
    """Постраничное чтение с сортировкой по ПОЛНОМУ уникальному ключу таблицы.

    Без order() PostgREST отдаёт страницы с повторами и пропусками (how-we-work, 15 сентября).
    Сортировка по id при фильтре по дате упирается в statement_timeout 8 с: планировщик идёт по
    индексу id через всю таблицу. Ключ (дата, площадка, sku[, …]) совпадает с уникальным индексом —
    он обслуживает и фильтр, и порядок.
    """
    out, page = [], 0
    while True:
        qb = sb.table(table).select(select)
        for f in filters:
            qb = getattr(qb, f[0])(*f[1:])
        for col in orders:
            qb = qb.order(col)
        res = qb.range(page * 1000, page * 1000 + 999).execute()
        out.extend(res.data)
        if len(res.data) < 1000:
            return out
        page += 1


def load_day_raw(day, today, allow_fetch, counters):
    """Начисления дня: файл → API. Возвращает (accruals | None, откуда)."""
    import json
    path = os.path.join(RAW_DIR, f"{day}.json")
    if os.path.exists(path):
        data = json.load(open(path))
        return (data["accruals"] if isinstance(data, dict) else data), "файл"
    if not allow_fetch:
        return None, "нет сырья"
    import fetch_accrual_postings_raw as raw
    if raw.in_night_window(datetime.now(timezone.utc)):
        print(f"  {day}: сырья нет, а сейчас окно ночного прогона {raw.window_text()} — в API не иду")
        return None, "нет сырья (ночное окно)"
    accruals = raw.fetch_day(day, counters, 60)
    if (date.fromisoformat(today) - date.fromisoformat(day)).days < YOUNG_DAYS:
        return accruals, "API, без файла (моложе двух суток)"
    os.makedirs(RAW_DIR, exist_ok=True)
    json.dump({"date": day, "fetched_at": datetime.now(timezone.utc).isoformat(), "accruals": accruals}, open(path, "w"), ensure_ascii=False)
    return accruals, "API → файл"


LEDGER = "ozon_accrual_daily_types"


def load_types_from_ledger(sb, d1, d2):
    """{день: {type_id: расход, положителен при списании}} из леджера начислений по типам.

    В леджере знак Ozon (списание < 0). Внутри листа типы держатся расходом с плюсом — как type_sums()
    по файлам сырья, — поэтому знак переворачивается здесь, один раз, на входе.
    """
    rows = fetch(sb, LEDGER, "accrual_date,type_id,amount", [("gte", "accrual_date", d1), ("lte", "accrual_date", d2)], ["accrual_date", "type_id"])
    out = defaultdict(dict)
    for r in rows:
        out[r["accrual_date"]][int(r["type_id"])] = -D(r["amount"])
    return dict(out)


def types_differ(a, b):
    """Типы двух источников за один день: {type_id: (a, b)} там, где суммы разные. Нулевая сумма = отсутствию типа."""
    return {t: (a.get(t, Z), b.get(t, Z)) for t in sorted(set(a) | set(b)) if q(a.get(t, Z)) != q(b.get(t, Z))}


def load_costs(sb, snapshot):
    orders = fetch(sb, "marketplace_orders", "id,order_date,order_schema,marketplace_sku,article,orders_qty,orders_amount_seller,"
                                             "cancelled_orders_qty,cancelled_orders_amount_seller,observed_at",
                   [("eq", "marketplace_code", "ozon"), ("gte", "order_date", "2026-03-28")],
                   ["order_date", "marketplace_code", "marketplace_sku", "order_schema"])
    sku_art = defaultdict(lambda: defaultdict(Decimal))
    for r in orders:
        # вес — созданные (подтв. + отм.): после пересборки у полностью отменённых ключей orders_qty = 0, а артикул у них тот же
        sku_art[str(r["marketplace_sku"])][str(r["article"] or "")] += D(r["orders_qty"]) + D(r.get("cancelled_orders_qty"))
    sku2art = {s: max(sorted(a.items()), key=lambda kv: kv[1])[0] for s, a in sku_art.items()}
    norms = sorted({a.lower() for a in sku2art.values() if a})
    cost = {}
    for i in range(0, len(norms), 500):
        res = sb.table("article_unit_costs").select("offer_id_norm,unit_cost").eq("snapshot_date", snapshot).in_("offer_id_norm", norms[i:i + 500]).order("offer_id_norm").execute()
        for r in res.data:
            cost[r["offer_id_norm"]] = D(r["unit_cost"])

    def unit_cost(sku):
        art = sku2art.get(str(sku))
        return cost.get(art.lower()) if art else None
    return sku2art, unit_cost, len(cost), len(norms), orders


def cost_index_note(days, snapshot):
    rates = sorted({cost_index_for(d) for d in days if cost_index_for(d) is not None})
    if not rates:
        return f"Себестоимость — снимок 1С {snapshot}; индекса СС на эти даты нет, колонки «по индексу» пусты."
    return (f"Себестоимость — снимок 1С {snapshot}. Индекс СС " + ", ".join(f"{v:.3f}" for v in rates)
            + " (справочно, пока нет новой выгрузки 1С; источник — лист владельца 22.09, 1–21 сентября: 27 401 217,62 / 23 824 962,63; "
              "у него по дням 1,175 → 1,116). Колонки «по индексу» = снимок × индекс; основные — на снимке.")


def overhead_note(days):
    rates = sorted({(overhead_for(d) or Z) for d in days})
    if rates == [Z]:
        return "Накладные в день на эти даты не заданы — Ebitda пуста, не ноль (OVERHEAD_PER_DAY)."
    return ("Ebitda = Фин. рез. − накладные в день; накладные " + ", ".join(f"{v:,.2f}" for v in rates if v)
            + "/день — из ручного листа за сентябрь (22.09, ячейка T83); откуда число и как оно меняется — владелец скажет. "
              "«Ebitda с компенсациями» сопоставима с «Ebitda» ручного листа, у которого компенсации внутри «Прочего».")


def write_xlsx(path, month, rows, total, sku_rows, notes, sku_notes, platforms=None, addition=None, orders=None, orders_error=None,
               segments=None, types_summary=None):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    money, pct = "#,##0.00", "0.0%"
    bold, young_fill, head_fill = Font(bold=True), PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9E1F2")
    wb = openpyxl.Workbook()
    ws = wb.active
    title = f"Ozon - {MONTHS[int(month[5:7]) - 1]}"
    ws.title = title
    ws["A1"] = "Выкупы Ozon"; ws["A1"].font = bold
    cols = [("Дата реализации", "date", None), ("Оборот - выкупы", "turnover", money), ("Комиссия (с НДС), руб.", "commission", money), ("Комиссия, %", "commission_pct", pct),
            ("Выручка, руб.", "revenue", money), ("Себестоимость, руб.", "cogs", money), ("Маржа, руб.", "margin", money), ("Мар-ть, %", "margin_pct", pct),
            ("Логистика, руб.", "logistics", money), ("% Логистики", "logistics_pct", pct), ("Эквайринг, руб.", "acquiring", money), ("% Эквайринга", "acquiring_pct", pct),
            ("Подписка, руб.", "subscription", money), ("Реклама, руб.", "ads", money), ("% ДРР (от ТО)", "drr_pct", pct), ("Прочее, руб.", "other", money),
            ("Фин. рез., руб.", "fin_result", money), ("% Фин. рез.", "fin_result_pct", pct),
            ("Накладные в день, руб.", "overhead", money), ("Ebitda, руб.", "ebitda", money), ("% Ebitda", "ebitda_pct", pct), (None, None, None),
            ("Компенсации Ozon (доход), руб.", "compensations", money), ("справочно: Фин. рез. с компенсациями", "fin_result_with_comp", money),
            ("% Фин. рез. с компенсациями", "fin_result_with_comp_pct", pct),
            ("справочно: Ebitda с компенсациями", "ebitda_with_comp", money), ("% Ebitda с компенсациями", "ebitda_with_comp_pct", pct),
            ("справочно: СС по индексу", "cogs_index", money), ("Маржа по индексу", "margin_index", money),
            ("Фин. рез. по индексу", "fin_result_index", money), ("Ebitda по индексу", "ebitda_index", money),
            ("справочно: Реклама по образцу", "ads_like_manual", money), ("справочно: Логистика + Прочее по образцу", "log_other_like_manual", money),
            ("Позиций выкупов", "positions", "#,##0"), ("Штук (где не измерены — позиции)", "units", "#,##0"),
            ("строк СС по штукам", "rows_by_units", "#,##0"), ("строк СС по позициям", "rows_by_positions", "#,##0"),
            ("позиций без себестоимости", "no_cost_positions", "#,##0"), ("Примечание", "note", None)]
    for j, (head, _k, _f) in enumerate(cols, 1):
        if head:
            c = ws.cell(row=3, column=j, value=head); c.font = bold; c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for i, r in enumerate(rows + [total], 4):
        r = dict(r)
        drift = r.get("db_minus_raw") or {}
        r["note"] = "; ".join(x for x in ("моложе двух суток — начисления доезжают" if r.get("young") else "",
                                           ("marketplace_expenses расходится с типами начислений (расходы − типы, с НДС): " + ", ".join(f"{k} {v:+,.2f}" for k, v in drift.items())
                                            + " — лист считает по типам") if drift else "",
                                           "" if r.get("has_raw") else ("нет типов начислений (ни в леджере, ни в файле) — эквайринг внутри прочего, рекламы и компенсаций нет" if r["date"] != "Итого" else "есть дни без типов начислений — итог неполон")) if x)
        for j, (head, k, fmt) in enumerate(cols, 1):
            if not head:
                continue
            v = r.get(k)
            if k == "date" and v != "Итого":
                v = date.fromisoformat(v)
            elif isinstance(v, Decimal):
                v = float(v)
            c = ws.cell(row=i, column=j, value=v)
            if fmt:
                c.number_format = fmt
            if k == "date" and r["date"] != "Итого":
                c.number_format = "DD.MM.YYYY"
            if r.get("young"):
                c.fill = young_fill
            if r["date"] == "Итого":
                c.font = bold
    base = 4 + len(rows) + 2
    for n, line in enumerate(notes):
        ws.cell(row=base + n, column=1, value=line)
    ws.freeze_panes = "B4"
    ws.row_dimensions[3].height = 45
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14 if j > 1 else 16
    ws.column_dimensions[get_column_letter(len(cols))].width = 60

    for name, prows in list((platforms or {}).items()) + list((segments or {}).items()):
        ptotal = platform_total(prows)
        if name in (NO_PLATFORM, NO_METAL) and not any(ptotal[k] for k in PLATFORM_MONEY):
            continue
        wsp = wb.create_sheet(name)
        wsp["A1"] = f"Выкупы Ozon — {name}"; wsp["A1"].font = bold
        pcols = [("Дата реализации", "date", None), ("Оборот - выкупы", "turnover", money), ("Комиссия (с НДС), руб.", "commission", money), ("Выручка, руб.", "revenue", money),
                 ("Себестоимость, руб.", "cogs", money), ("Маржа, руб.", "margin", money), ("Мар-ть, %", "margin_pct", pct), ("Логистика, руб.", "logistics", money),
                 ("Подписка, руб.", "subscription", money), ("Реклама (Performance), руб.", "ads_perf", money), ("% ДРР (от ТО)", "drr_pct", pct),
                 ("Прочее с эквайрингом, руб.", "other_with_acq", money), ("Фин. рез., руб.", "fin_result", money), ("% Фин. рез.", "fin_result_pct", pct), ("Позиций", "positions", "#,##0")]
        for j, (head, _k, _f) in enumerate(pcols, 1):
            c = wsp.cell(row=3, column=j, value=head); c.font = bold; c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for i, r in enumerate(prows + [ptotal], 4):
            for j, (_head, k, fmt) in enumerate(pcols, 1):
                v = r.get(k)
                if k == "date" and v != "Итого":
                    v = date.fromisoformat(v)
                elif isinstance(v, Decimal):
                    v = float(v)
                c = wsp.cell(row=i, column=j, value=v)
                if fmt:
                    c.number_format = fmt
                if k == "date" and r["date"] != "Итого":
                    c.number_format = "DD.MM.YYYY"
                if r["date"] == "Итого":
                    c.font = bold
        wsp.cell(row=len(prows) + 7, column=1, value="Площадка — по первой букве артикула (F / S / T). Логистика, подписка и прочее — marketplace_expenses по SKU; эквайринг от прочего не отделён, "
                                                     "компенсаций нет: леджер начислений без SKU. Реклама — Performance по SKU.")
        wsp.freeze_panes = "B4"; wsp.row_dimensions[3].height = 45
        for j in range(1, len(pcols) + 1):
            wsp.column_dimensions[get_column_letter(j)].width = 15
    if addition:
        wsa = wb.create_sheet("Сложение площадок")
        for j, head in enumerate(("Колонка", "Σ листов площадок", "Общий лист", "Разница", "Чем объясняется"), 1):
            c = wsa.cell(row=1, column=j, value=head); c.font = bold; c.fill = head_fill
        for i, (title, a, b, diff, why) in enumerate(addition, 2):
            for j, v in enumerate((title, a, b, diff, why), 1):
                c = wsa.cell(row=i, column=j, value=float(v) if isinstance(v, Decimal) else v)
                if j in (2, 3, 4):
                    c.number_format = money
        for j, w in enumerate((22, 20, 20, 16, 120), 1):
            wsa.column_dimensions[get_column_letter(j)].width = w

    ws2 = wb.create_sheet("По SKU")
    cols2 = [("SKU", "sku", None), ("Артикул", "article", None), ("Товар", "name", None), ("Позиций", "positions", "#,##0"), ("Оборот - выкупы", "turnover", money),
             ("Комиссия (с НДС)", "commission", money), ("Выручка", "revenue", money), ("Себестоимость", "cogs", money), ("Маржа", "margin", money), ("Мар-ть, %", "margin_pct", pct),
             ("Реклама (Performance)", "ads", money), ("% ДРР (от ТО)", "drr_pct", pct), ("Логистика", "logistics", money), ("Прочее (вкл. эквайринг и подписку)", "other", money)]
    for j, (head, _k, _f) in enumerate(cols2, 1):
        c = ws2.cell(row=1, column=j, value=head); c.font = bold; c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    tot = {"sku": "Итого", "article": "", "name": ""}
    for k in ("positions", "turnover", "commission", "revenue", "ads", "logistics", "other"):
        tot[k] = sum((r[k] for r in sku_rows), Z)
    tot["cogs"] = sum((r["cogs"] for r in sku_rows if r["cogs"] is not None), Z)
    tot["margin"] = tot["revenue"] - tot["cogs"]
    tot["margin_pct"] = ratio(tot["margin"], tot["revenue"])
    vat = rows[0]["vat"] if rows else vat_for(month + "-01")
    tot["drr_pct"] = ratio(tot["ads"], tot["turnover"] / vat) if tot["turnover"] else None
    for i, r in enumerate(sku_rows + [tot], 2):
        for j, (_head, k, fmt) in enumerate(cols2, 1):
            v = r.get(k)
            c = ws2.cell(row=i, column=j, value=float(v) if isinstance(v, Decimal) else v)
            if fmt:
                c.number_format = fmt
            if r["sku"] == "Итого":
                c.font = bold
    for n, line in enumerate(sku_notes):
        ws2.cell(row=len(sku_rows) + 4 + n, column=1, value=line)
    ws2.freeze_panes = "D2"
    ws2.row_dimensions[1].height = 45
    for j, w in enumerate([14, 18, 40] + [15] * (len(cols2) - 3), 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    if types_summary:
        wss = wb.create_sheet("Начисления - свод")
        wss["A1"] = "Начисления Ozon по типам за месяц (леджер ozon_accrual_daily_types, знак Ozon: списание < 0)"; wss["A1"].font = bold
        wss["A2"] = ("Группа услуг — как в выгрузке ЛК (сверено с полотном владельца за ноябрь 2025); Вид — справочник владельца «Тип начисления → Вид»; "
                     "статья — наша свёртка TYPE_TO_EXPENSE. Продаж и комиссии здесь нет: в леджере только услуги с type_id.")
        heads = ("type_id", "Тип начисления (ЛК)", "Группа услуг (Ozon)", "Вид (владелец)", "Наша статья", "Σ за месяц, руб.", "Дней", "Строк услуг")
        for j, h in enumerate(heads, 1):
            c = wss.cell(row=4, column=j, value=h); c.font = bold; c.fill = head_fill
        for i, r in enumerate(types_summary, 5):
            for j, k in enumerate(("type_id", "name", "group", "kind", "article", "amount", "days", "lines"), 1):
                v = r.get(k)
                c = wss.cell(row=i, column=j, value=float(v) if isinstance(v, Decimal) else v)
                if k == "amount":
                    c.number_format = money
        for j, w in enumerate((8, 52, 30, 18, 22, 18, 7, 11), 1):
            wss.column_dimensions[get_column_letter(j)].width = w
        wss.freeze_panes = "A5"
    if orders is not None:
        write_orders_sheets(wb, orders, (money, pct, bold, head_fill, young_fill))
    elif orders_error:
        wse = wb.create_sheet("Заказы", 1)
        wse["A1"] = "Лист «Заказы» НЕ СОБРАН: " + orders_error; wse["A1"].font = bold
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return tot


# ---------- лист «Заказы»: прогноз по кривой дозревания ----------

LOG_TABLE = "ozon_posting_status_log"


def other_share_of(ledger_by_day, buyouts):
    """Доля прочих расходов в обороте выкупов: (логистика + эквайринг + подписка + прочее) с НДС / оборот с НДС.

    Считается по дням, которые есть в леджере: оборот берётся за те же дни, иначе доля занижена днями без начислений.
    Реклама (41, 54) и компенсации (25, 10) в прочие не входят; незнакомый тип входит И называется вслух.
    """
    expense, unknown = Z, set()
    for ts in ledger_by_day.values():
        for type_id, v in ts.items():
            if type_id in accrual.AD_TYPE_IDS or type_id in accrual.UNCLASSIFIED_TYPE_IDS:
                continue
            if type_id not in accrual.TYPE_TO_EXPENSE:
                unknown.add(type_id)
            expense += v
    turnover = sum((D(r["buyouts_amount_seller"]) for r in buyouts if r["buyout_date"] in ledger_by_day), Z)
    return (expense / turnover if turnover else None), expense, turnover, sorted(unknown)


def build_orders(sb, days, order_rows, daily_rows, sku2art, unit_cost, tz_name, day_note=None, snapshot=SNAP):
    """Всё для листов «Заказы»: читает лог статусов, выкупы и леджер окна долей; возвращает словарь для записи и печати."""
    d2 = days[-1]
    stamps = defaultdict(list)
    for r in order_rows:
        if r.get("observed_at"):
            stamps[str(r.get("order_schema") or "").lower()].append(r["observed_at"])
    if not any(stamps.get(s) for s in forecast.SCHEMAS):
        raise RuntimeError("в marketplace_orders нет ни одной строки Ozon с observed_at — состояние заказов не датировать")
    # свежесть — по схеме: упал ночью шаг FBS — его заказы на сутки старше, и возраст у них свой
    obs_by_schema = {s: max(forecast.parse_ts(v) for v in set(stamps[s])).astimezone(ZoneInfo(tz_name)) for s in forecast.SCHEMAS if stamps.get(s)}
    obs = max(obs_by_schema.values())
    obs_date = obs.date().isoformat()
    obs_dates = {s: v.date().isoformat() for s, v in obs_by_schema.items()}
    log_from = (date.fromisoformat(obs_date) - timedelta(days=forecast.CURVE_NIGHTS + max(forecast.PLATEAU_AGES) + 1)).isoformat()
    log_rows = fetch(sb, LOG_TABLE, "posting_number,observed_at,schema,order_date,status,amount", [("gte", "order_date", log_from)], ["posting_number", "observed_at"])
    curve = forecast.build_curve(log_rows, tz_name)
    # Чувствительность рублёвой кривой к одной отмене: та же кривая без крупнейшей отмены каждой схемы. Прогноз считается
    # по полной кривой — выбрасывать наблюдение нельзя, — но читатель обязан видеть, сколько на нём держится.
    biggest = {c["biggest_event"][2] for c in curve.values() if c["biggest_event"][2]}
    curve_wo = forecast.build_curve([r for r in log_rows if r["posting_number"] not in biggest], tz_name) if biggest else {}
    sensitivity = {s: {a: (c["r_amt"][a], curve_wo[s]["r_amt"][a]) for a in (1, 7, 14)} for s, c in curve.items() if s in curve_wo}
    w1, w2 = forecast.window_before(d2)
    buyouts_w = fetch(sb, "marketplace_buyouts", "id,buyout_date,marketplace_sku,buyouts_amount_seller,commission_amount",
                      [("eq", "marketplace_code", "ozon"), ("gte", "buyout_date", w1), ("lte", "buyout_date", w2)], ["buyout_date", "marketplace_code", "marketplace_sku"])
    commission_share, commission_base = forecast.shares_by_platform(buyouts_w, lambda sku: platform_of(sku2art.get(str(sku or ""))))
    ledger_w = load_types_from_ledger(sb, w1, w2)
    other_share, other_expense, other_turnover, unknown_types = other_share_of(ledger_w, buyouts_w)
    ads_by_day = {r["date"]: r.get("ads") for r in daily_rows}
    ads_manual_by_day = {r["date"]: r.get("ads_like_manual") for r in daily_rows}     # (41 + 54 + 96 + вся подписка) / НДС — как у владельца
    blocks, said = forecast.build_orders_daily(days, order_rows, curve, obs_dates, commission_share, other_share, ads_by_day, unit_cost, vat_for,
                                               lambda r: platform_of(r.get("article") or sku2art.get(str(r.get("marketplace_sku") or ""))),
                                               ads_manual_by_day)
    young = {r["date"] for r in daily_rows if r.get("young")}
    totals = {}
    for key, rows in blocks.items():
        for r in rows:
            forecast.add_order_ratios(r)
            r["young"] = r["date"] in young
            idx = cost_index_for(r["date"])
            r["cogs_index"] = None if idx is None or r.get("cogs") is None else r["cogs"] * idx
            r["fin_result_index"] = None if idx is None or r.get("fin_result") is None else r["fin_result"] + r["cogs"] - r["cogs_index"]
        totals[key] = forecast.orders_total(rows)
    # сверка плато другим разрезом: доля отмен в самой таблице заказов по дням возраста 21…50 (штуки ТОВАРА, не отправления)
    table_plateau = {}
    for schema in forecast.SCHEMAS:
        acc = defaultdict(Decimal)
        for r in order_rows:
            age = (date.fromisoformat(obs_date) - date.fromisoformat(r["order_date"])).days
            if str(r.get("order_schema") or "").lower() == schema and forecast.MATURE_AGE <= age <= 50:
                acc["conf_q"] += D(r["orders_qty"]); acc["canc_q"] += D(r.get("cancelled_orders_qty"))
                acc["conf_a"] += D(r["orders_amount_seller"]); acc["canc_a"] += D(r.get("cancelled_orders_amount_seller"))
        table_plateau[schema] = (ratio(acc["canc_q"], acc["conf_q"] + acc["canc_q"]), ratio(acc["canc_a"], acc["conf_a"] + acc["canc_a"]))
    return {"blocks": blocks, "totals": totals, "said": said, "curve": curve, "obs": obs, "obs_date": obs_date, "log_rows": len(log_rows),
            "window": (w1, w2), "commission_share": commission_share, "commission_base": commission_base,
            "other_share": other_share, "other_expense": other_expense, "other_turnover": other_turnover, "ledger_days": len(ledger_w),
            "unknown_types": unknown_types, "table_plateau": table_plateau, "sensitivity": sensitivity, "mature": forecast.mature_check(blocks["all"]),
            "day_note": day_note or {}, "snapshot": snapshot}


def orders_notes(o):
    """Шапка и подвал листа «Заказы»: чем измерено каждое число, которое у владельца было константой."""
    pct = lambda v: "—" if v is None else f"{v * 100:.2f} %".replace(".", ",")  # noqa: E731
    curve, names = o["curve"], {"fbo": "FBO", "fbs": "FBS"}
    nights = sorted({n for c in curve.values() for n in c["nights"]})
    head = ("Кривая дозревания: лог статусов отправлений, " + (f"ночи {nights[0]} … {nights[-1]}" if nights else "ночных срезов НЕТ")
            + "".join(f"; {names[s]} — пар соседних ночей {c['pairs']}, отправлений {c['postings']}" for s, c in curve.items())
            + f". Состояние заказов — на ночь {o['obs_date']} ({o['obs'].strftime('%d.%m %H:%M')} МСК). День старше {forecast.MATURE_AGE} суток — факт; моложе — "
              "прогноз: подтверждено сейчас × (1 − доля подтверждённых, которая ещё отменится).")
    w1, w2 = o["window"]
    share = o["commission_share"]
    lines = [f"Комиссия — измеренная доля за {w1} … {w2} из marketplace_buyouts, по площадке товара (в выкупах схемы FBO / FBS нет, а Селект живёт на FBS): "
             + ", ".join(f"{name} {pct(share.get(name))}" for _p, name in PLATFORMS) + f", все {pct(share.get('все'))}. У владельца — 0,41.",
             f"Прочие — измеренная доля за то же окно: (логистика + эквайринг + подписка + прочее, с НДС) {o['other_expense']:,.2f} / оборот выкупов "
             f"{o['other_turnover']:,.2f} = {pct(o['other_share'])} (дней окна в леджере начислений {o['ledger_days']} из 30; компенсации Ozon — доход, сюда не входят). "
             "У владельца — 2,4 % от выручки.",
             "Выкупаемость у владельца — 0,65 на все дни. Здесь — по возрасту дня: доля подтверждённых сейчас, которая ещё отменится (по ₽ / по отправлениям): "
             + "; ".join(f"{names[s]} — 1 сут. {pct(c['r_amt'][1])} / {pct(c['r_cnt'][1])}, 7 сут. {pct(c['r_amt'][7])} / {pct(c['r_cnt'][7])}, 14 сут. {pct(c['r_amt'][14])} / {pct(c['r_cnt'][14])}"
                         for s, c in curve.items()) + ". Вся кривая — на листе «Кривая дозревания».",
             "Рублёвая кривая шумит хвостом цен — крупнейшая отмена в переходах: "
             + "; ".join(f"{names[s]} — {c['biggest_event'][1]:,.0f} ₽ на {c['biggest_event'][0]}-е сутки из {c['events_amt_total']:,.0f} ₽ всех отмен ({pct(ratio(c['biggest_event'][1], c['events_amt_total']))})"
                         for s, c in curve.items() if c["biggest_event"][0] is not None)
             + ". Без неё доля «ещё отменится» по ₽ была бы: "
             + "; ".join(f"{names[s]} — " + ", ".join(f"{a} сут. {pct(wo)} вместо {pct(w)}" for a, (w, wo) in sorted(v.items())) for s, v in o["sensitivity"].items())
             + ". Прогноз считается по полной кривой; с каждой ночью вес одной отмены падает.",
             "Сверка другим разрезом — доля отмен в самой таблице заказов по дням возраста 21…50 суток (штуки товара / ₽): "
             + "; ".join(f"{names[s]} {pct(q_)} / {pct(a_)}" for s, (q_, a_) in o["table_plateau"].items()) + ".",
             owner_multiplier_note(o),
             "Справочный фин. рез. = Маржа (созданных) × 0,65 − Реклама по образцу − Выручка × 0,65 × 0,024; реклама по образцу = (41 + 54 + 96 + вся статья подписки) / НДС "
             "по дню заказа, как в его листе; себестоимость — по созданным штукам. Справочный ДРР = Реклама по образцу / (Выручка × 0,65 × НДС / 0,59) — "
             "у него на листе «Заказы» в знаменателе 0,58, на «Заказы Standard» — 0,73. Колонка модели «Реклама (41 + 54)» и наш фин. рез. прогноз — без подписки и сбора отзывов.",
             cost_index_note([r["date"] for r in o["blocks"]["all"]], o.get("snapshot", SNAP)).replace("Колонки «по индексу» = снимок × индекс; основные — на снимке.",
                                                                                                    "«СС по индексу» и «Фин. рез. прогноз по индексу» = снимок × индекс; основные — на снимке."),
             platform_note(o),
             "Реклама и фин. рез. — только на общем листе: леджер начислений без схемы. Себестоимость — article_unit_costs × прогноз подтверждённых штук товара.",
             f"Приёмка: дозревших дней {o['mature'][0]}, прогноз = факту на {o['mature'][1]}" + (f"; НЕ равен: {', '.join(o['mature'][2])}" if o["mature"][2] else "") + "."]
    if o["unknown_types"]:
        lines.append(f"ВНИМАНИЕ: в окне долей есть незнакомые типы начислений {o['unknown_types']} — вошли в прочие.")
    lines.extend("ВНИМАНИЕ: " + x for x in o["said"])
    return head, lines


def owner_multiplier_note(o):
    """Множители «выручки» владельца по площадкам с датой действия — и рядом измеренная доля комиссии модели."""
    days = [r["date"] for r in o["blocks"]["all"]]
    d1, d2 = days[0], days[-1]
    names = [n for _p, n in PLATFORMS] + [NO_PLATFORM]
    mult = lambda d: ", ".join(f"{n} {str(forecast.owner_after_commission(d, n)).replace('.', ',')}" for n in names)  # noqa: E731
    share = o["commission_share"]
    pct = lambda v: "—" if v is None else f"{v * 100:.1f} %".replace(".", ",")  # noqa: E731
    return (f"Справочная «Выручка» = создано × множитель площадки владельца / НДС; множители по его сентябрьскому листу (проверены сырьём отправлений 09-01 … 09-16 "
            f"по UTC-суткам: 0,530 / 0,900 / 0,530 без разброса): на {d2} — {mult(d2)}" + (f"; на {d1} — {mult(d1)}" if mult(d1) != mult(d2) else "")
            + " (комиссия 47 % / 47 % / 10 %; до 2026-09-01 — 0,59 по подписи июльской книги, июль не перепроверялся). "
            + "Измеренная доля комиссии — в модели: " + ", ".join(f"{n} {pct(share.get(n))}" for _p, n in PLATFORMS) + ". "
            + "Его источник режет сутки по UTC, наша таблица — по МСК: по дням его/наша шумит, по итогу сходится.")


def platform_note(o):
    """Площадки на заказах: доля созданного по площадкам и «без площадки» — вслух, не молча."""
    parts, total = [], o["totals"]["all"]["created_a"] or Z
    for _p, name in PLATFORMS + ((None, NO_PLATFORM),):
        t = o["totals"].get(f"platform:{name}")
        if t is None or (name == NO_PLATFORM and not t["created_a"]):
            continue
        parts.append(f"{name} {t['created_a']:,.2f} ₽ ({(t['created_a'] / total * 100) if total else 0:.1f} %)")
    return ("Площадки — по первой букве артикула, той же картой sku → артикул, что у выкупных листов: создано " + ", ".join(parts)
            + ". «Без площадки» — товары без артикула в карте; в общий лист входят, в листы площадок нет.")


def write_orders_sheets(wb, o, styles):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    money, pct, bold, head_fill, young_fill = styles
    head, lines = orders_notes(o)
    full = [("Дата заказа", "date", None), ("Возраст, сут.", "age", "0"), ("Создано, шт", "created_q", "#,##0"), ("Создано, руб.", "created_a", money),
            ("Подтверждено сейчас, шт", "conf_q", "#,##0"), ("Подтверждено сейчас, руб.", "conf_a", money), ("Отменено сейчас, шт", "canc_q", "#,##0"),
            ("Отменено сейчас, руб.", "canc_a", money), ("Отменено сейчас, % руб.", "canc_pct", pct), ("Ожидаемые отмены ещё, руб.", "expected_cancels_a", money),
            ("Прогноз подтв., шт", "fc_q", "#,##0.0"), ("Прогноз подтв., руб.", "fc_a", money), ("Прогноз отмен всего, % руб.", "fc_canc_pct", pct),
            ("Комиссия прогноз, руб.", "commission", money), ("Комиссия, %", "commission_pct", pct), ("Выручка прогноз, руб.", "revenue", money),
            ("Себестоимость, руб.", "cogs", money), ("Маржа, руб.", "margin", money), ("Мар-ть, %", "margin_pct", pct),
            ("Реклама (41 + 54), руб.", "ads", money), ("Реклама по образцу (41 + 54 + 96 + подписка), руб.", "ads_manual", money),
            ("% ДРР от созданного", "drr_created_pct", pct), ("% ДРР от прогноза подтв.", "drr_fc_pct", pct),
            ("Прочие, руб.", "other", money), ("Фин. рез. прогноз, руб.", "fin_result", money), ("% Фин. рез.", "fin_result_pct", pct), (None, None, None),
            ("СС по индексу, руб.", "cogs_index", money), ("Фин. рез. прогноз по индексу, руб.", "fin_result_index", money),
            ("справочно, формулы владельца: Выручка = создано × множитель площадки владельца / НДС", "owner_revenue", money), ("Маржа (созданных)", "owner_margin", money),
            ("ДРР = Реклама по образцу / (Выручка × 0,65 × НДС / 0,59)", "owner_drr_pct", pct),
            ("Фин. рез. = Маржа × 0,65 − Реклама по образцу − Выручка × 0,65 × 0,024", "owner_fin_result", money),
            ("Наш фин. рез. − владельца", "fin_result_minus_owner", money), ("шт подтв. без себестоимости", "no_cost_q", "#,##0"), ("Примечание", "note", None)]
    only_total = {"ads", "ads_manual", "drr_created_pct", "drr_fc_pct", "fin_result", "fin_result_pct", "fin_result_index", "owner_drr_pct", "owner_fin_result", "fin_result_minus_owner"}
    part_cols = [c for c in full if c[1] not in only_total and c[0] is not None]

    def write_block(ws, title, cols, rows, total, note_schemas, common):
        ws["A1"] = title; ws["A1"].font = bold
        ws["A2"] = head
        for j, (h, _k, _f) in enumerate(cols, 1):
            if h:
                c = ws.cell(row=4, column=j, value=h); c.font = bold; c.fill = head_fill
                c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for i, r in enumerate(rows + [total], 5):
            r = dict(r)
            if r["date"] != "Итого":
                left = "; ".join(f"{s.upper()} ещё −{forecast.remaining_share(o['curve'], s, r['age'], 'amt') * 100:.1f} % ₽".replace(".", ",")
                                 for s in note_schemas if forecast.remaining_share(o["curve"], s, r["age"], "amt") is not None)
                young_note = (o["day_note"].get(r["date"]) or "моложе двух суток — реклама доезжает") if r.get("young") and common else ""
                r["note"] = "; ".join(x for x in ("факт: день дозрел" if r["mature"] else f"прогноз: {left}" if left else "прогноза нет: кривой нет", young_note) if x)
            for j, (h, k, fmt) in enumerate(cols, 1):
                if not h:
                    continue
                v = r.get(k)
                if k == "date" and v != "Итого":
                    v = date.fromisoformat(v)
                elif isinstance(v, Decimal):
                    v = float(v)
                c = ws.cell(row=i, column=j, value=v)
                if fmt:
                    c.number_format = fmt
                if k == "date" and r["date"] != "Итого":
                    c.number_format = "DD.MM.YYYY"
                if r["date"] == "Итого":
                    c.font = bold
                elif not r["mature"]:
                    c.fill = young_fill
        base = 5 + len(rows) + 2
        ws.cell(row=base, column=1, value="Жёлтым — дни моложе 21 суток: в них прогноз, числа изменятся. Без заливки — факт.")
        for n, line in enumerate(lines, 1):
            ws.cell(row=base + n, column=1, value=line)
        ws.freeze_panes = "B5"; ws.row_dimensions[4].height = 60
        for j in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 15
        ws.column_dimensions[get_column_letter(len(cols))].width = 50

    index = 1
    for key, title in (("all", "Заказы"), ("fbo", "Заказы FBO"), ("fbs", "Заказы FBS")):
        ws = wb.create_sheet(title, index); index += 1
        write_block(ws, "Заказы Ozon по дню заказа — прогноз по кривой дозревания" + ("" if key == "all" else f", {key.upper()}"),
                    full if key == "all" else part_cols, o["blocks"][key], o["totals"][key], forecast.SCHEMAS if key == "all" else (key,), key == "all")
    # листы площадок — как у выкупов: та же карта sku → артикул; реклама и фин. рез. только на общем листе (леджер без площадки)
    for _p, name in PLATFORMS + ((None, NO_PLATFORM),):
        key = f"platform:{name}"
        if key not in o["blocks"] or (name == NO_PLATFORM and not o["totals"][key]["created_a"]):
            continue
        ws = wb.create_sheet(f"Заказы {name}", index); index += 1
        write_block(ws, f"Заказы Ozon по дню заказа — {name} (площадка по первой букве артикула)", part_cols, o["blocks"][key], o["totals"][key], forecast.SCHEMAS, False)
    wsc = wb.create_sheet("Кривая дозревания")
    wsc["A1"] = head; wsc["A1"].font = bold
    wsc["A2"] = ("h — доля живых (не отменённых) возраста k, отменённых к следующей ночи; внутри ведра возрастов доля одна. r — доля подтверждённых сейчас, которая ещё "
                 "отменится до 21-х суток: 1 − Π (1 − h). «шт» — отправления, не штуки товара. Срез — справочно: c(a) по молодым заказам и плато по старым; в прогноз не идёт.")
    heads = ["Возраст, сут."]
    for s in o["curve"]:
        heads += [f"{s.upper()}: живых", "отменено к след. ночи", "h, отправления", "h, ₽", "r, отправления", "r, ₽", "срез: c(a), отправления", "срез: c(a), ₽", "срез: r, отправления", "срез: r, ₽"]
    for j, h in enumerate(heads, 1):
        c = wsc.cell(row=4, column=j, value=h); c.font = bold; c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for a in range(0, forecast.MATURE_AGE):
        vals = [a]
        for s, c in o["curve"].items():
            vals += [c["alive"].get(a, (0, Z))[0], c["events"].get(a, (0, Z))[0], c["h_cnt"][a], c["h_amt"][a], c["r_cnt"][a], c["r_amt"][a],
                     c.get("slice_cnt", {}).get(a), c.get("slice_amt", {}).get(a), c.get("slice_r_cnt", {}).get(a), c.get("slice_r_amt", {}).get(a)]
        for j, v in enumerate(vals, 1):
            cell = wsc.cell(row=5 + a, column=j, value=float(v) if isinstance(v, Decimal) else v)
            if j > 1 and (j - 2) % 10 >= 2:
                cell.number_format = "0.00%"
    row = 5 + forecast.MATURE_AGE + 1
    for s, c in o["curve"].items():
        if "slice_cnt_inf" in c:
            wsc.cell(row=row, column=1, value=f"{s.upper()}: плато среза (возрасты 21…29) — отправления {c['slice_cnt_inf'] * 100:.2f} %, ₽ {c['slice_amt_inf'] * 100:.2f} %; ночи {c['nights'][0]} … {c['nights'][-1]}".replace(".", ","))
            row += 1
    wsc.row_dimensions[4].height = 45
    for j in range(1, len(heads) + 1):
        wsc.column_dimensions[get_column_letter(j)].width = 14


def print_orders(o, buyout_total):
    head, lines = orders_notes(o)
    print(f"\nлист «Заказы» (лог статусов: строк {o['log_rows']}):")
    print("  " + head)
    print(f"  {'':8}{'создано ₽':>18}{'подтв. сейчас ₽':>18}{'отменено ₽':>16}{'ещё отменится ₽':>17}{'прогноз подтв. ₽':>18}{'выручка':>16}{'СС':>16}{'маржа':>16}{'прочие':>14}")
    for key, name in (("fbo", "FBO"), ("fbs", "FBS"), ("all", "итого")):
        t = o["totals"][key]
        f = lambda v: "—" if v is None else f"{v:,.2f}"  # noqa: E731
        print(f"  {name:8}{f(t['created_a']):>18}{f(t['conf_a']):>18}{f(t['canc_a']):>16}{f(t['expected_cancels_a']):>17}{f(t['fc_a']):>18}{f(t['revenue']):>16}{f(t['cogs']):>16}{f(t['margin']):>16}{f(t['other']):>14}")
    t = o["totals"]["all"]
    p = lambda v: "—" if v is None else f"{v * 100:.2f} %"  # noqa: E731
    if t["fin_result"] is not None:
        print(f"  реклама {t['ads']:,.2f}; ДРР от созданного {p(t['drr_created_pct'])}, от прогноза подтв. {p(t['drr_fc_pct'])}; фин. рез. прогноз {t['fin_result']:,.2f} ({p(t['fin_result_pct'])} выручки)")
        if t.get("fin_result_index") is not None:
            print(f"  по индексу СС: СС {t['cogs_index']:,.2f}, фин. рез. прогноз {t['fin_result_index']:,.2f}")
        print(f"  реклама по образцу (41 + 54 + 96 + подписка) {t['ads_manual']:,.2f}; по формулам владельца (0,65 / множители площадок / 0,024) на ней: выручка {t['owner_revenue']:,.2f}, "
              f"ДРР {p(t['owner_drr_pct'])}, фин. рез. {t['owner_fin_result']:,.2f}; наш − владельца {t['fin_result_minus_owner']:+,.2f}")
        for _pl, name in PLATFORMS + ((None, NO_PLATFORM),):
            pt = o["totals"].get(f"platform:{name}")
            if pt is not None and pt["created_a"]:
                print(f"  {name:12} создано {f(pt['created_a']):>16}  прогноз подтв. {f(pt['fc_a']):>16}  выручка {f(pt['revenue']):>14}  "
                      f"создано × множитель владельца / НДС {f(pt['owner_revenue']):>14}")
        if buyout_total.get("fin_result") is not None:
            print(f"  рядом — выкупной лист за те же дни: оборот {buyout_total['turnover']:,.2f}, выручка {buyout_total['revenue']:,.2f}, фин. рез. {buyout_total['fin_result']:,.2f} "
                  f"(заказы − выкупы: фин. рез. {t['fin_result'] - buyout_total['fin_result']:+,.2f} — дата заказа против даты реализации, прочие у заказов долей, у выкупов фактом)")
    else:
        print("  реклама за окно неполна (есть дни без типов начислений) — ДРР и фин. рез. в итоге пусты")
    for line in lines:
        print("  · " + line)


# ---------- приёмка против ручного листа ----------

EXACT = (("Оборот", "turnover", "turnover"), ("Комиссия", "commission", "commission"), ("Выручка", "revenue", "revenue"),
         ("Эквайринг", "acquiring", "acquiring"), ("Реклама по образцу = их «Реклама»", "ads_like_manual", "ads"))


def build_types_summary(types_by_day, days):
    """Строки листа «Начисления - свод»: type_id → название ЛК, группа услуг Ozon, Вид владельца, наша статья, Σ за месяц.

    Названия — knowledge/ozon/accrual_types_2026-09-23.json (ответ /v1/finance/accrual/types, `description` = колонка «Тип
    начисления» ЛК); группа и Вид — scripts/accrual_types_owner_check.py (OZON_GROUP по полотну владельца за ноябрь 2025,
    справочник «Тип → Вид» из docs/owner_manual_report_instruction.md). Тип, которого нет ни там, ни там, — «(нет в справочнике)».
    """
    import json
    import accrual_types_owner_check as owner
    path = os.path.join(ROOT, "knowledge", "ozon", "accrual_types_2026-09-23.json")
    names = {t["id"]: t.get("description") or t.get("name") for t in json.load(open(path))} if os.path.exists(path) else {}
    directory = owner.owner_directory() if os.path.exists(owner.DOC) else []
    total, ndays, lines = defaultdict(Decimal), defaultdict(int), defaultdict(int)
    for d in days:
        for type_id, v in (types_by_day.get(d) or {}).items():
            # леджер и type_sums хранят расход со знаком «+» (см. load_types_from_ledger: -amount); в своде — знак Ozon
            total[int(type_id)] += -D(v["amount"] if isinstance(v, dict) else v)
            ndays[int(type_id)] += 1
            lines[int(type_id)] += int(v.get("lines", 0)) if isinstance(v, dict) else 0
    out = []
    for type_id in sorted(total):
        name = names.get(type_id, f"type_id {type_id}")
        kinds = sorted({k for _n, k in owner.match_owner(name, directory, type_id)}) if directory else []
        out.append({"type_id": type_id, "name": name, "group": owner.OZON_GROUP.get(type_id, "(нет в справочнике)"),
                    "kind": " / ".join(kinds) if kinds else "(нет в справочнике)",
                    "article": owner.OUR_ARTICLE.get(type_id, f"unknown_{type_id}"), "amount": total[type_id], "days": ndays[type_id], "lines": lines[type_id]})
    return out


SEGMENT_COLUMNS = (("Оборот (B)", "turnover", ("turnover",)), ("Комиссия (C)", "commission", ("commission",)), ("Логистика (I)", "logistics", ("logistics",)),
                   ("Реклама (K) — у нас Performance по SKU", "ads_perf", ("ads",)), ("Эквайринг + Прочее (M + O) — у нас одной суммой", "other_with_acq", ("acquiring", "other")),
                   ("Себестоимость (F) через индекс", "cogs_index", ("cogs",)))


def read_manual_segment_sheet(path, sheet):
    """Первый блок листа «в т.ч. Серебро - месяц»: строки с 4-й до первой не-даты в колонке A (ниже — «Итого» и блок WB с теми
    же датами; общий read_manual_sheet читал до 60-й строки, и WB-даты затирали Ozon-даты — 2026-09-23)."""
    import datetime
    import warnings
    import openpyxl
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    from reconcile_manual_report_ozon import MANUAL_COLUMNS
    for row in wb[sheet].iter_rows(min_row=4, max_row=400, values_only=True):
        d = row[0]
        if not isinstance(d, datetime.datetime):
            break
        out[d.date().isoformat()] = {c: (None if v is None else Decimal(str(v))) for c, v in zip(MANUAL_COLUMNS, row[1:19])}
    return out


def check_segment(prows, manual):
    """Сегмент (лист «в т.ч. Серебро - месяц» владельца) против нашего сегмента: [(колонка, их Σ, наша Σ, дней 0,00, расхождения)]."""
    out = []
    for title, ours_k, theirs_ks in SEGMENT_COLUMNS:
        t_o = t_t = Z
        bad, missing, n = [], [], 0
        for r in prows:
            m = manual.get(r["date"])
            if not m or any(m.get(k) is None for k in theirs_ks):
                missing.append(r["date"]); continue
            o = r[ours_k] if ours_k != "cogs_index" else r["cogs"] * (cost_index_for(r["date"]) or Decimal(1))
            t = sum((D(m[k]) for k in theirs_ks), Z)
            o, t = q(o), q(t)
            t_o += o; t_t += t; n += 1
            if o != t:
                bad.append((r["date"], o - t))
        out.append({"title": title, "theirs": t_t, "ours": t_o, "diff": t_o - t_t, "equal_days": n - len(bad), "days": n, "bad": bad, "missing": missing})
    return out


def print_check_segment(name, table):
    print(f"\nсегмент «{name}» против листа владельца")
    print(f"{'колонка':56}{'у них':>16}{'у нас':>16}{'разница':>14}  дней 0,00")
    for r in table:
        print(f"{r['title'][:56]:56}{r['theirs']:>16,.2f}{r['ours']:>16,.2f}{r['diff']:>14,.2f}  {r['equal_days']} из {r['days']}"
              + (f"; нет у них: {len(r['missing'])} дн." if r["missing"] else ""))
        if r["bad"]:
            worst = sorted(r["bad"], key=lambda x: -abs(x[1]))[:6]
            print("      расходится: " + ", ".join(f"{d} {v:+,.2f}" for d, v in worst) + (f" … ещё {len(r['bad']) - 6} дн." if len(r["bad"]) > 6 else ""))


def check(rows, manual):
    """Сверка с ручным листом. Возвращает (таблица строк, число нарушений в колонках, обязанных сходиться до копейки)."""
    out, failures = [], 0

    def line(title, ours_fn, theirs_fn, must_match, why=""):
        nonlocal failures
        t_o = t_t = Z
        diffs = []
        for r in rows:
            m = manual.get(r["date"])
            o = ours_fn(r)
            t = theirs_fn(m) if m else None
            if o is None or t is None:
                diffs.append((r["date"], None)); continue
            o, t = q(o), q(t)
            t_o += o; t_t += t
            if o != t:
                diffs.append((r["date"], o - t))
        bad = [(d, v) for d, v in diffs if v is not None]
        missing = [d for d, v in diffs if v is None]
        if must_match and (bad or missing):
            failures += 1
        out.append({"title": title, "theirs": t_t, "ours": t_o, "diff": t_o - t_t, "equal_days": len(rows) - len(bad) - len(missing),
                    "days": len(rows), "bad": bad, "missing": missing, "must": must_match, "why": why})

    g = lambda k: (lambda m: None if m.get(k) is None else D(m[k]))  # noqa: E731
    for title, ours_k, theirs_k in EXACT:
        line(title, lambda r, k=ours_k: r.get(k), g(theirs_k), True)
    their_lo = lambda m: None if m.get("logistics") is None or m.get("other") is None else D(m["logistics"]) + D(m["other"])  # noqa: E731
    line("Л + П по образцу − Компенсации = их Логистика + Прочее",
         lambda r: None if r.get("log_other_like_manual") is None else r["log_other_like_manual"] - r["compensations"], their_lo, False,
         "у владельца компенсации (типы 25, 10) сидят внутри «Прочего» расходом с обратным знаком; у нас — отдельной строкой дохода (решение 8)")
    line("   то же без компенсаций: Л + П по образцу = их Л + П", lambda r: r.get("log_other_like_manual"), their_lo, False,
         "разница — компенсации дня")
    line("Себестоимость (наш снимок 1С против их цен)", lambda r: r.get("cogs"), g("cogs"), False,
         "решение 4 от 2026-09-19: остаёмся на снимке 05-20, расхождение принимается как известное")
    cogs_line = out[-1]
    idx_days = [(r["date"], (D(manual[r["date"]]["cogs"]) / r["cogs"]).quantize(Decimal("0.001")))
                for r in rows if r.get("cogs") and manual.get(r["date"]) and manual[r["date"]].get("cogs") is not None]
    total_idx = (cogs_line["theirs"] / cogs_line["ours"]).quantize(Decimal("0.001")) if cogs_line["ours"] else None
    params = sorted({cost_index_for(r["date"]) for r in rows if cost_index_for(r["date"]) is not None})
    param = params[0] if len(params) == 1 else None
    cogs_line.update({"index_days": idx_days, "index_total": total_idx, "index_param": param,
                      "index_stale": (param is not None and total_idx is not None and abs(total_idx - param) > COST_INDEX_STALE_PCT * param)})
    line("Фин. рез. с компенсациями против их «Фин. рез.»", lambda r: r.get("fin_result_with_comp"), g("fin_result"), False,
         "остаток = разница себестоимости (решение 4) и 09-01; подписка у них внутри рекламы — на итог не влияет")
    line("Ebitda с компенсациями против их «Ebitda»", lambda r: r.get("ebitda_with_comp"), g("ebitda"), False,
         "их − наша = разница себестоимости (решение 4) + 209,02 на 09-01, как у фин. реза: накладные в день одни и те же")
    line("Логистика (наша статья против их строки)", lambda r: r.get("logistics"), g("logistics"), False, "границу не повторяем — решение 7")
    line("Прочее (наше без эквайринга против их строки)", lambda r: r.get("other"), g("other"), False, "границу не повторяем — решение 7")
    line("Реклама (41 + 54) против их «Рекламы»", lambda r: r.get("ads"), g("ads"), False, "у них в рекламе ещё вся подписка (51, 52, 74) и сбор отзывов (96)")
    return out, failures


def read_manual_orders_sheet(path, sheet, columns=("revenue", "margin", "margin_pct", "ads", "drr_pct", "coinvest_pct", "fin_result_pct", "fin_result")):
    """Лист «Заказы» владельца: строки с 3-й; A — день, B выручка, C маржа, D м-ть, E реклама, F ДРР, G соинвест, H % фин. реза, I фин. рез.
    Листы площадок («Заказы Standard», «Ozon Select», «Ozon Дискаунтер») — B выручка, C маржа. Строка-ошибка Excel — None."""
    import datetime
    import warnings
    import openpyxl
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for row in wb[sheet].iter_rows(min_row=3, max_row=80, values_only=True):
        d = row[0]
        if isinstance(d, datetime.datetime) and row[1] is not None:
            out[d.date().isoformat()] = {c: (Decimal(repr(v)) if isinstance(v, (int, float)) else None) for c, v in zip(columns, row[1:1 + len(columns)])}
    return out


def check_orders(rows, manual, platform_rows, manual_platforms, young_days):
    """Сверка листа «Заказы» с листом владельца по дням. Возвращает (таблица, отказов, разложения).

    Обязана сходиться одна колонка — «Реклама по образцу» против его E, и только на днях старше двух суток
    (за D−1 реклама доезжает после ночного сбора). Выручка B и себестоимость — известные разницы, печатаются с
    отношением его/наша по дням; фин. рез. I раскладывается на базу / СС / рекламу / прочие 2,4 % (формулы
    владельца) и на выручку / СС / рекламу / прочие (наш прогноз по кривой). Знак разниц — «наша − его».
    """
    OWN = forecast.OWNER
    table, failures = [], 0
    by_date = {r["date"]: r for r in rows}

    def line(title, ours_fn, theirs_fn, must, why="", with_ratio=False):
        nonlocal failures
        t_o = t_t = Z
        diffs, ratios = [], []
        for r in rows:
            m = manual.get(r["date"])
            o = ours_fn(r)
            t = theirs_fn(m) if m else None
            if o is None or t is None:
                diffs.append((r["date"], None)); continue
            o, t = q(o), q(t)
            t_o += o; t_t += t
            if with_ratio and o:
                ratios.append((r["date"], t / o))
            if o != t:
                diffs.append((r["date"], o - t))
        bad = [(d, v) for d, v in diffs if v is not None]
        missing = [d for d, v in diffs if v is None]
        strict_bad = [(d, v) for d, v in bad if d not in young_days]
        strict_missing = [d for d in missing if d not in young_days]
        if must and (strict_bad or strict_missing):
            failures += 1
        table.append({"title": title, "theirs": t_t, "ours": t_o, "diff": t_o - t_t, "equal_days": len(rows) - len(bad) - len(missing), "days": len(rows),
                      "bad": bad, "missing": missing, "must": must, "why": why, "ratios": ratios, "young_bad": [(d, v) for d, v in bad if d in young_days]})

    g = lambda k: (lambda m: m.get(k))  # noqa: E731
    line("Выручка B vs создано × множитель владельца / НДС", lambda r: r.get("owner_revenue"), g("revenue"), False,
         "его сутки — UTC, наши — МСК: по дням шумит, по итогу ≈ 1; 09-17 у него вбит иначе (+848 179 к правилу) — печатать, не чинить", with_ratio=True)
    line("Маржа C vs Маржа (созданных)", lambda r: r.get("owner_margin"), g("margin"), False, "разница базы и себестоимости (снимок 05-20, решение 4)")
    line("Реклама E vs Реклама по образцу", lambda r: r.get("ads_manual"), g("ads"), True,
         "обязана сходиться на днях старше двух суток; за D−1 реклама доезжает после ночного сбора")
    line("Фин. рез. I vs Фин. рез. по формулам владельца", lambda r: r.get("owner_fin_result"), g("fin_result"), False, "разложение ниже: база / СС / реклама / прочие 2,4 %")
    line("Фин. рез. I vs Фин. рез. прогноз (кривая)", lambda r: r.get("fin_result"), g("fin_result"), False, "разложение ниже: выручка (выкупаемость, комиссия) / СС / реклама / прочие")
    for name, his_rows in manual_platforms.items():
        ours = {r["date"]: r for r in platform_rows.get(name, [])}
        saved = manual
        manual = his_rows
        line(f"Выручка его листа площадки vs {name}: создано × множитель / НДС", lambda r, ours=ours: (ours.get(r["date"]) or {}).get("owner_revenue"), g("revenue"), False,
             "отношение его/наша по дням: ≈ 1, разброс — граница суток (UTC у него, МСК у нас)", with_ratio=True)
        manual = saved

    # разложения — по дням, где есть обе стороны
    dec_owner = defaultdict(Decimal); dec_fc = defaultdict(Decimal); n_owner = n_fc = 0
    sums = defaultdict(Decimal)
    for r in rows:
        m = manual.get(r["date"]) or {}
        if m.get("revenue") is None or m.get("margin") is None or m.get("fin_result") is None:
            continue
        his_cogs = m["revenue"] - m["margin"]
        if r.get("owner_fin_result") is not None and m.get("ads") is not None:
            dB = r["owner_revenue"] - m["revenue"]; dC = r["cogs_created"] - his_cogs; dE = r["ads_manual"] - m["ads"]
            dec_owner["база = 0,65 × Δвыручки"] += OWN["buyout_rate"] * dB
            dec_owner["СС = −0,65 × Δсебестоимости"] += -OWN["buyout_rate"] * dC
            dec_owner["реклама = −Δрекламы по образцу"] += -dE
            dec_owner["прочие = −0,65 × 0,024 × Δвыручки"] += -OWN["buyout_rate"] * OWN["other_rate"] * dB
            dec_owner["итого наш − его"] += r["owner_fin_result"] - m["fin_result"]
            n_owner += 1
        if r.get("fin_result") is not None and m.get("ads") is not None:
            dec_fc["выручка = прогноз выручки − B × 0,65"] += r["revenue"] - m["revenue"] * OWN["buyout_rate"]
            dec_fc["СС = −(СС прогноза − его СС × 0,65)"] += -(r["cogs"] - his_cogs * OWN["buyout_rate"])
            dec_fc["реклама = −(41 + 54 − его E)"] += -(r["ads"] - m["ads"])
            dec_fc["прочие = −(прочие прогноза − B × 0,65 × 0,024)"] += -(r["other"] - m["revenue"] * OWN["buyout_rate"] * OWN["other_rate"])
            dec_fc["итого наш − его"] += r["fin_result"] - m["fin_result"]
            sums["created_a"] += r["created_a"]; sums["fc_a"] += r["fc_a"]; sums["commission"] += r["commission"]
            n_fc += 1
    implied = {"выкупаемость прогноза": ratio(sums["fc_a"], sums["created_a"]), "комиссия прогноза": ratio(sums["commission"], sums["fc_a"])}
    return table, failures, {"owner": (dict(dec_owner), n_owner), "forecast": (dict(dec_fc), n_fc), "implied": implied}


def print_check_orders(table, decomposition, young_days):
    print(f"\n{'лист «Заказы» против листа владельца':58}{'у него':>17}{'у нас':>17}{'наша − его':>15}{'дней 0,00':>11}")
    for t in table:
        print(f"{t['title'][:57]:58}{t['theirs']:>17,.2f}{t['ours']:>17,.2f}{t['diff']:>15,.2f}{t['equal_days']:>7} из {t['days']}"
              + ("   ОБЯЗАНА СХОДИТЬСЯ" if t["must"] and [b for b in t["bad"] if b[0] not in young_days] else ""))
        if t["bad"] and (t["must"] or len(t["bad"]) <= 4):
            print("      расходится: " + ", ".join(f"{d} {v:+,.2f}" + (" (моложе двух суток)" if d in young_days else "") for d, v in t["bad"]))
        if t["missing"]:
            print("      нет значения: " + ", ".join(t["missing"]))
        if t["ratios"]:
            vals = [v for _d, v in t["ratios"]]
            print(f"      его/наша: итого {t['theirs'] / t['ours'] if t['ours'] else 0:.3f}, по дням {min(vals):.3f} … {max(vals):.3f}: "
                  + ", ".join(f"{d[5:]} {v:.3f}" for d, v in t["ratios"]))
        if t["why"] and (t["bad"] or t["missing"]):
            print(f"      чем объясняется: {t['why']}")
    for key, title in (("owner", "Фин. рез. I против формул владельца на наших числах"), ("forecast", "Фин. рез. I против нашего прогноза по кривой")):
        parts, n = decomposition[key]
        print(f"\n  {title} — разложение «наша − его» по {n} дням:")
        for k, v in parts.items():
            print(f"      {k:52}{v:>18,.2f}")
        check = sum((v for k, v in parts.items() if not k.startswith("итого")), Z) - parts.get("итого наш − его", Z)
        print(f"      {'сумма частей − итого':52}{check:>18,.2f}")
    print("  подразумеваемые доли прогноза: " + ", ".join(f"{k} {v * 100:.2f} %" for k, v in decomposition["implied"].items() if v is not None)
          + " (у владельца 65 % и 41 %)")


def created_by_utc_day(postings_by_scheme):
    """{(UTC-день, площадка): создано ₽} из сырых отправлений — цена × количество по всем статусам, площадка по первой букве offer_id.

    Его источник режет сутки по UTC (проверено 2026-09-22: 0,530 / 0,900 / 0,530 без разброса на 16 днях), наша таблица — по МСК;
    сравнивать с его листом честно можно только на сырье с временем заказа. FBO — created_at, FBS — in_process_at (как в загрузчиках).
    """
    import loaders.ozon_fbo_orders_loader as fbo
    out = defaultdict(Decimal)
    for scheme, postings in postings_by_scheme.items():
        field = "created_at" if scheme == "fbo" else "in_process_at"
        for p in postings:
            ts = p.get(field)
            if not ts:
                continue
            day = datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(timezone.utc).date().isoformat()
            for pr in p.get("products") or []:
                out[(day, platform_of(pr.get("offer_id")))] += D(fbo.posting_price(pr)) * D(pr.get("quantity"))
    return dict(out)


def check_orders_utc(manual_platforms, created_utc, days):
    """Его B × НДС / наше созданное по UTC-суткам, по площадкам. Возвращает [(площадка, день, его B, создано, отношение, ожидание)]."""
    rows = []
    for name, his in manual_platforms.items():
        for d in days:
            m = his.get(d)
            created = created_utc.get((d, name))
            if not m or m.get("revenue") is None or not created:
                continue
            rows.append((name, d, m["revenue"], created, (m["revenue"] * vat_for(d) / created).quantize(Decimal("0.001")), forecast.owner_after_commission(d, name)))
    return rows


def print_check_orders_utc(rows):
    print("\n  по UTC-суткам (сырьё отправлений): его B × НДС / наше созданное; ожидание — множитель владельца площадки")
    by = defaultdict(list)
    for name, d, _b, _c, r, exp in rows:
        by[name].append((d, r, exp))
    for name, items in by.items():
        off = [(d, r, exp) for d, r, exp in items if abs(r - exp) > Decimal("0.005")]
        print(f"    {name:12} дней {len(items)}, ровно по множителю {len(items) - len(off)}: " + ", ".join(f"{d[5:]} {r}" for d, r, _e in items))
        if off:
            print(f"    {'':12} отклонения: " + ", ".join(f"{d[5:]} {r} (ожидание {exp})" for d, r, exp in off))


def print_check(table):
    print(f"\n{'колонка':58}{'у них':>17}{'у нас':>17}{'разница':>15}{'дней 0,00':>11}")
    for t in table:
        print(f"{t['title'][:57]:58}{t['theirs']:>17,.2f}{t['ours']:>17,.2f}{t['diff']:>15,.2f}{t['equal_days']:>7} из {t['days']}"
              + ("   ОБЯЗАНА СХОДИТЬСЯ" if t["must"] and (t["bad"] or t["missing"]) else ""))
        if t["bad"] and (t["must"] or len(t["bad"]) <= 6):
            print("      расходится: " + ", ".join(f"{d} {v:+,.2f}" for d, v in t["bad"]))
        if t["missing"]:
            print("      нет значения: " + ", ".join(t["missing"]))
        if t["why"] and (t["bad"] or t["missing"]):
            print(f"      чем объясняется: {t['why']}")
        if t.get("index_days"):
            print(f"      индекс СС (их / наша): по итогу {t['index_total']}, параметр {t['index_param']}; по дням "
                  + ", ".join(f"{d[5:]} {v}" for d, v in t["index_days"]))
            if t.get("index_stale"):
                print(f"      ВНИМАНИЕ: индекс СС протух: параметр {t['index_param']}, по листу {t['index_total']} (больше {COST_INDEX_STALE_PCT * 100:.0f} %)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", required=True, help="YYYY-MM")
    ap.add_argument("--date-to", help="последний день листа; по умолчанию — вчера по местному времени: сегодняшний день ночь ещё не грузила")
    ap.add_argument("--out")
    ap.add_argument("--snapshot", default=SNAP)
    ap.add_argument("--no-fetch", action="store_true", help="в API не ходить ни за чем (= --fetch none)")
    ap.add_argument("--fetch", choices=("none", "young", "all"), default="young",
                    help="в API: none — никогда; young — только за дни моложе двух суток, живым ответом, файл не сохраняется (реклама за D−1 "
                         "доезжает после ночного сбора, в леджере её ещё нет); all — ещё и за дни без файла сырья (в файл). --check подразумевает all")
    ap.add_argument("--check-orders", action="store_true", help="сверить лист «Заказы» с листом владельца (--xlsx, --orders-sheet, --orders-platform-sheets); код 1, если реклама по образцу не сошлась на дне старше двух суток")
    ap.add_argument("--orders-sheet", default="Заказы", help="лист владельца по заказам (строки с 3-й: день, выручка, маржа, …, реклама E, фин. рез. I)")
    ap.add_argument("--orders-platform-sheets", default="Заказы Standard,Ozon Select,Ozon Дискаунтер", help="его листы площадок в порядке Основная, Селект, Дискаунтер")
    ap.add_argument("--day-boundary", choices=("msk", "utc"), default="msk",
                    help="с --check-orders: utc — ещё и таблица по UTC-суткам на сырье отправлений окна (data/postings_raw/history_*, файл есть — в API не идёт; "
                         "иначе сбор теми же методами, что ночь: ~2 обращения на день на схему)")
    ap.add_argument("--types-from", choices=("db", "files"), default="db",
                    help="откуда типы начислений: db — леджер ozon_accrual_daily_types (дня нет в леджере — файл, нет файла — API); "
                         "files — только файлы сырья, как до леджера. С --check читаются ОБА источника и обязаны совпасть")
    ap.add_argument("--summary-json", help="положить сюда итоги книги (для подписи к файлу в Telegram)")
    ap.add_argument("--no-orders", action="store_true", help="лист «Заказы» не собирать (лог статусов и окно долей не читаются)")
    ap.add_argument("--check", action="store_true", help="сверить с ручным листом (--xlsx, --sheet); код возврата 1, если обязанные колонки не сошлись")
    ap.add_argument("--check-segment", help="сверить сегмент по металлу («Серебро» / «Золото») с листом владельца --segment-sheet (--xlsx); только печать")
    ap.add_argument("--segment-sheet", help="лист владельца «в т.ч. Серебро - <месяц>»")
    ap.add_argument("--xlsx"); ap.add_argument("--sheet")
    args = ap.parse_args()
    # «Сегодня» — по местному времени, как у загрузчиков. Сегодняшнего дня в базе нет (ночь грузит по вчера), а в
    # сырье начислений он уже есть: строка с рекламой и эквайрингом без оборота врала бы, поэтому лист кончается вчера.
    today_date = datetime.now(timezone.utc).astimezone(ZoneInfo(os.getenv("APP_TIMEZONE", "Europe/Moscow"))).date()
    today = today_date.isoformat()
    last = args.date_to or (today_date - timedelta(days=1)).isoformat()
    days = [d for d in month_days(args.month, last) if d <= today]
    if not days:
        raise SystemExit("в окне нет ни одного наступившего дня")
    d1, d2 = days[0], days[-1]

    import loaders.ozon_fbo_orders_loader as fbo
    sb = fbo.supabase
    flt = lambda col: [("eq", "marketplace_code", "ozon"), ("gte", col, d1), ("lte", col, d2)]  # noqa: E731
    buyout_cols = "id,buyout_date,marketplace_sku,buyouts_qty,buyouts_amount_seller,commission_amount"
    try:
        buyouts = fetch(sb, "marketplace_buyouts", buyout_cols + ",buyouts_units", flt("buyout_date"), ["buyout_date", "marketplace_code", "marketplace_sku"])
        units_column = True
    except Exception as exc:
        if "buyouts_units" not in str(exc):
            raise
        buyouts = fetch(sb, "marketplace_buyouts", buyout_cols, flt("buyout_date"), ["buyout_date", "marketplace_code", "marketplace_sku"])
        units_column = False
    expenses = fetch(sb, "marketplace_expenses", "id,expense_date,marketplace_sku,expense_type,expense_amount", flt("expense_date"),
                     ["expense_date", "marketplace_code", "marketplace_sku", "expense_type"])
    kpi = fetch(sb, "daily_sku_kpi", "id,kpi_date,marketplace_sku,article,product_name,buyouts_qty,buyouts_amount_seller,commission_amount,ad_spend,logistics_amount,other_expenses_amount",
                flt("kpi_date"), ["kpi_date", "marketplace_code", "marketplace_sku"])
    sku2art, unit_cost, cost_found, cost_asked, order_rows = load_costs(sb, args.snapshot)

    counters, sources, types_by_day, day_note = {"requests": 0, "429": 0}, defaultdict(list), {}, {}
    young_days = {d for d in days if (date.fromisoformat(today) - date.fromisoformat(d)).days < YOUNG_DAYS}
    fetch_mode = "none" if args.no_fetch else ("all" if args.check else args.fetch)
    ledger = load_types_from_ledger(sb, d1, d2) if (args.types_from == "db" or args.check) else {}
    file_types = {}

    def fetch_raw(d, allow):
        """Начисления дня из файла или API; отказ API — не падение листа, а (None, причина): день останется на леджере с примечанием."""
        try:
            return load_day_raw(d, today, allow, counters)
        except Exception as exc:  # noqa: BLE001
            return None, f"живой ответ не получен: {type(exc).__name__}: {str(exc)[:120]}"

    for d in days:
        young = d in young_days
        if args.types_from == "db" and d in ledger and not args.check:
            # Молодой день (моложе двух суток) в леджере лежит ночным снимком, а начисления рекламы за него доезжают к утру
            # (2026-09-22: 41/54 за 09-21 в 00:27 нет, в 07:07 есть): берём живой ответ, файл не сохраняем. Нет ответа — леджер,
            # и в примечании дня «реклама не доехала» — нулём не молчим.
            if young and fetch_mode != "none":
                accruals, where = fetch_raw(d, True)
                if accruals is not None:
                    file_types[d] = type_sums(accruals)
                    types_by_day[d] = file_types[d]
                    sources["живой ответ — моложе двух суток, в леджере рекламы за день ещё нет"].append(d)
                    day_note[d] = "моложе двух суток — реклама живым ответом by-day (в леджере за день её ещё нет)"
                    continue
                day_note[d] = f"реклама не доехала: {where}; типы из леджера — ночной снимок, рекламы за день в нём ещё нет"
            types_by_day[d] = ledger[d]
            sources["леджер"].append(d)
            continue
        # дня нет в леджере (моложе посева и ещё не записан ночью), либо источник — файлы, либо --check сверяет оба;
        # в --check молодой день — живым ответом без требования равенства с леджером.
        accruals, where = fetch_raw(d, fetch_mode == "all" or (young and fetch_mode != "none"))
        if young and accruals is None and d in ledger:
            day_note[d] = f"реклама не доехала: {where}; типы из леджера"
        if accruals is not None:
            file_types[d] = type_sums(accruals)
        if args.types_from == "db" and d in ledger and not (young and accruals is not None):
            types_by_day[d] = ledger[d]
            sources["леджер"].append(d)
        else:
            sources[where + (" — моложе двух суток, живой ответ свежее леджера" if young and d in ledger else "")].append(d)
            if accruals is not None:
                types_by_day[d] = file_types[d]
    source_mismatch = {d: types_differ(ledger[d], file_types[d]) for d in days
                       if d in ledger and d in file_types and d not in young_days and types_differ(ledger[d], file_types[d])}
    young_drift = {d: types_differ(ledger[d], file_types[d]) for d in days if d in young_days and d in ledger and d in file_types and types_differ(ledger[d], file_types[d])}
    both = [d for d in days if d in ledger and d in file_types and d not in young_days]

    rows, unknown = build_daily(days, buyouts, expenses, types_by_day, unit_cost, today)
    for r in rows:
        add_ratios(r)
    total = total_row(rows)
    vat = vat_for(d1)
    sku_rows = build_sku(kpi, sku2art, unit_cost, vat)
    platforms = build_platform_daily(days, buyouts, expenses, kpi, sku2art, unit_cost)
    platform_totals = {name: platform_total(prows) for name, prows in platforms.items()}
    addition = platform_addition(total, platform_totals)
    sku2name = {}
    for r in kpi:
        if r.get("product_name") and str(r["marketplace_sku"]) not in sku2name:
            sku2name[str(r["marketplace_sku"])] = r["product_name"]
    metals = build_metal_daily(days, buyouts, expenses, kpi, sku2name, unit_cost)
    metal_totals = {name: platform_total(prows) for name, prows in metals.items()}
    metal_addition = platform_addition(total, metal_totals)
    types_summary = build_types_summary(types_by_day, days)
    # Лист «Заказы» — отдельная беда от выкупного: не собрался — остальные листы выходят, причина пишется в книгу и вслух.
    orders, orders_error = None, None
    if not args.no_orders:
        try:
            orders = build_orders(sb, days, order_rows, rows, sku2art, unit_cost, os.getenv("APP_TIMEZONE", "Europe/Moscow"), day_note, args.snapshot)
        except Exception as exc:  # noqa: BLE001
            orders_error = f"{type(exc).__name__}: {exc}"

    print(f"окно {d1} … {d2} ({len(days)} дн.): выкупов строк {len(buyouts)}, расходов {len(expenses)}, строк витрины {len(kpi)}; "
          f"sku→article {len(sku2art)}, себестоимость найдена у {cost_found} артикулов из {cost_asked} (снимок {args.snapshot}); НДС {vat}")
    for where, ds in sources.items():
        print(f"  типы начислений — {where}: {len(ds)} дн." + (f" ({ds[0]} … {ds[-1]})" if len(ds) > 1 else f" ({ds[0]})"))
    print(f"  обращений к Seller API {counters['requests']}, 429 — {counters['429']}")
    if unknown:
        print("ВНИМАНИЕ: незнакомые статьи расходов, учтены в «Прочем»:", {k: f"{v:,.2f}" for k, v in unknown.items()})
    if args.check or args.types_from == "files":
        print(f"  два источника типов (леджер и файлы сырья): дат с обоими {len(both)}, совпали до копейки по каждому типу {len(both) - len(source_mismatch)}"
              + (f", в леджере нет {len([d for d in days if d not in ledger])} дн." if any(d not in ledger for d in days) else ""))
        for d, diff in source_mismatch.items():
            print(f"      {d}: " + ", ".join(f"тип {t}: леджер {a:,.2f} / файл {b:,.2f}" for t, (a, b) in diff.items()))
        for d, diff in young_drift.items():
            print(f"      {d} (моложе двух суток, начисления доезжают — не расхождение): "
                  + ", ".join(f"тип {t}: леджер {a:,.2f} / живой {b:,.2f}" for t, (a, b) in diff.items()))
    drifted = [(r["date"], r["db_minus_raw"]) for r in rows if r.get("db_minus_raw")]
    if drifted:
        print("  marketplace_expenses расходится с типами начислений по статьям («+» — в расходах больше). Лист считает по типам; в расходах\n"
              "  застревают строки начислений, которые Ozon потом убрал, — upsert их не удаляет:")
        for d, drift in drifted:
            print(f"      {d}: " + ", ".join(f"{k} {v:+,.2f}" for k, v in drift.items()))
    young = [r["date"] for r in rows if r["young"]]
    if young:
        print(f"  моложе двух суток (начисления доезжают, в листе помечены): {', '.join(young)}")
    # статья subscription против типа 51: «Подписка» листа — вся статья; если в ней не только Premium, сказать
    sub_article = sum((D(r["expense_amount"]) for r in expenses if r["expense_type"] == "subscription"), Z)
    sub_51 = sum((ts.get(PREMIUM_TYPE, Z) for ts in types_by_day.values()), Z)
    if types_by_day and len(types_by_day) == len(days) and q(sub_article) != q(sub_51):
        print(f"  подписка: статья {sub_article:,.2f} против типа 51 в сырье {sub_51:,.2f}, разница {sub_article - sub_51:,.2f} — в статье есть типы 52 / 74 либо начисления доехали после сбора")

    ads_sheet1 = total["ads"]
    ads_sku = sum((r["ads"] for r in sku_rows), Z)
    # разница рекламы двух листов — по дням, чтобы объяснение стояло рядом с числом
    perf_by_day = defaultdict(Decimal)
    for r in kpi:
        perf_by_day[r["kpi_date"]] += D(r.get("ad_spend")) / vat_for(r["kpi_date"])
    ad_gaps = sorted(((r["date"], perf_by_day.get(r["date"], Z) - r["ads"]) for r in rows if r.get("ads") is not None),
                     key=lambda kv: -abs(kv[1]))
    ad_gaps = [(d, v) for d, v in ad_gaps if abs(v) >= 1]
    sku_notes = [f"Реклама на этом листе — Performance API (ad_spend витрины), без НДС: {ads_sku:,.2f}.",
                 (f"На листе «Ozon - …» реклама из начислений (типы 41 + 54), без НДС: {ads_sheet1:,.2f}; разница {ads_sku - ads_sheet1:+,.2f} — "
                  "Performance против начислений: день без сбора Performance даёт недобор (за 2026-09-14 в Performance пусто — известная дыра), "
                  "начисления ведут учёт по своей дате.") if ads_sheet1 is not None else "Реклама из начислений за окно неполна (есть дни без сырья) — разницу не считаю.",
                 ("Разница по дням (Performance − начисления, без НДС): " + "; ".join(f"{d} {v:+,.2f}" for d, v in ad_gaps[:8])
                  + (f"; ещё {len(ad_gaps) - 8} дн. на {sum((v for _d, v in ad_gaps[8:]), Z):+,.2f}" if len(ad_gaps) > 8 else "")) if ad_gaps else "По дням реклама двух источников сходится до рубля.",
                 "Прочее здесь — other_expenses_amount витрины: эквайринг, подписка и внешнее продвижение внутри; строка «(без SKU)» — расходы без товара."]
    notes = [f"Источники: marketplace_buyouts, ozon_accrual_daily_types (типы начислений; статьи — свёрткой типов), article_unit_costs (снимок {args.snapshot}, позиции выкупов). НДС {vat}.",
             "Компенсации Ozon (типы 25 и 10) — отдельной строкой дохода, в расходы не входят; «Фин. рез. с компенсациями» сопоставим с «Фин. рез.» ручного листа.",
             "Реклама = начисления 41 + 54. «Реклама по образцу» = (41 + 54 + 96 + вся статья подписки: 51, 52, 74) / НДС — так считает ручной лист. "
             "«Логистика + Прочее по образцу» = (логистика + прочее − тип 1 − тип 96) / НДС.",
             overhead_note(days),
             cost_index_note(days, args.snapshot),
             "Граница свёртки типов начислений — 24.08.2026: с 25.08 расходы лежат по справочнику владельца (типы 16, 17, 45, 62, 78, 82 — прочее; "
             "46, 63 — логистика), до неё — старой свёрткой; до пересборки истории логистика / прочее по обе стороны границы не сопоставимы.",
             "Даты, залитые жёлтым, моложе двух суток: начисления ещё доезжают, числа вырастут.",
             f"Собрано {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}; scripts/report_ozon_month.py."]
    out = args.out or os.path.join(OUT_DIR, f"ozon_{args.month}.xlsx")
    cogs_note = (f"Себестоимость: по штукам — {total['rows_by_units']} строк выкупов, по позициям — {total['rows_by_positions']} строк"
                 + ("." if units_column else " (колонки marketplace_buyouts.buyouts_units ещё нет — миграция не применена)."))
    notes.insert(1, cogs_note)
    write_xlsx(out, args.month, rows, total, sku_rows, notes, sku_notes, platforms, addition, orders, orders_error, segments=metals, types_summary=types_summary)
    print("  " + cogs_note)
    print(f"\nплощадки (по первой букве артикула):")
    print(f"  {'площадка':14}{'оборот':>18}{'выручка':>18}{'СС':>18}{'маржа':>16}{'логистика':>13}{'прочее+экв.':>14}{'реклама Perf.':>16}{'фин. рез.':>16}")
    for name, t in platform_totals.items():
        print(f"  {name:14}{t['turnover']:>18,.2f}{t['revenue']:>18,.2f}{t['cogs']:>18,.2f}{t['margin']:>16,.2f}{t['logistics']:>13,.2f}{t['other_with_acq']:>14,.2f}{t['ads_perf']:>16,.2f}{t['fin_result']:>16,.2f}")
    print(f"  сложение площадок в общий лист:")
    for title, a, b, diff, why in addition:
        print(f"    {title:20}{a:>18,.2f}{('—' if b is None else f'{b:,.2f}'):>18}{('—' if diff is None else f'{diff:,.2f}'):>14}" + (f"   {why}" if diff else ""))
    print(f"\nметалл (по названию товара из витрины; SKU без названия — «{NO_METAL}»):")
    print(f"  {'металл':22}{'оборот':>18}{'комиссия':>16}{'выручка':>16}{'СС':>16}{'логистика':>13}{'прочее+экв.':>14}{'реклама Perf.':>16}{'фин. рез.':>16}{'позиций':>9}")
    for name, t in metal_totals.items():
        print(f"  {name:22}{t['turnover']:>18,.2f}{t['commission']:>16,.2f}{t['revenue']:>16,.2f}{t['cogs']:>16,.2f}{t['logistics']:>13,.2f}{t['other_with_acq']:>14,.2f}{t['ads_perf']:>16,.2f}{t['fin_result']:>16,.2f}{t['positions']:>9,.0f}")
    print(f"  сложение металлов в общий лист (обязаны сходиться там же, где площадки):")
    for title, a, b, diff, why in metal_addition:
        print(f"    {title:20}{a:>18,.2f}{('—' if b is None else f'{b:,.2f}'):>18}{('—' if diff is None else f'{diff:,.2f}'):>14}" + (f"   {why}" if diff else ""))
    print(f"лист «Начисления - свод»: типов {len(types_summary)}, из них без группы ЛК {sum(1 for r in types_summary if r['group'] == '(нет в справочнике)')}, "
          f"без Вида владельца {sum(1 for r in types_summary if r['kind'] == '(нет в справочнике)')}; Σ {sum((r['amount'] for r in types_summary), Z):,.2f}")
    print(f"\nитого: оборот {total['turnover']:,.2f}, комиссия {total['commission']:,.2f}, выручка {total['revenue']:,.2f}, СС {total['cogs']:,.2f}, маржа {total['margin']:,.2f}")
    if total["fin_result"] is not None:
        print(f"       логистика {total['logistics']:,.2f}, эквайринг {total['acquiring']:,.2f}, подписка {total['subscription']:,.2f}, реклама {total['ads']:,.2f}, "
              f"прочее {total['other']:,.2f}, фин. рез. {total['fin_result']:,.2f}"
              + (f", накладные {total['overhead']:,.2f}, Ebitda {total['ebitda']:,.2f}" if total.get("ebitda") is not None else ", Ebitda — (накладные не заданы)")
              + (f"; по индексу СС: СС {total['cogs_index']:,.2f}, фин. рез. {total['fin_result_index']:,.2f}" if total.get("fin_result_index") is not None else ""))
    print(f"лист «По SKU»: строк {len(sku_rows)}, реклама Performance {ads_sku:,.2f}" + (f" против начислений {ads_sheet1:,.2f}, разница {ads_sku - ads_sheet1:+,.2f}" if ads_sheet1 is not None else ""))
    if ad_gaps:
        print("       по дням (Performance − начисления, без НДС): " + "; ".join(f"{d} {v:+,.2f}" for d, v in ad_gaps[:8])
              + (f"; ещё {len(ad_gaps) - 8} дн. на {sum((v for _d, v in ad_gaps[8:]), Z):+,.2f}" if len(ad_gaps) > 8 else ""))
    if orders is not None:
        print_orders(orders, total)
    elif orders_error:
        print(f"\nЛИСТ «ЗАКАЗЫ» НЕ СОБРАН: {orders_error}")
    print(f"записано: {out}")

    if args.summary_json:
        import json
        num = lambda v: None if v is None else str(q(v))  # noqa: E731
        no_types = [r["date"] for r in rows if not r["has_raw"]]
        summary = {"month": args.month, "date_from": d1, "date_to": d2,
                   "buyouts": {k: num(total.get(k)) for k in ("turnover", "revenue", "fin_result", "ebitda", "fin_result_index")},
                   "warnings": ([f"нет типов начислений за {', '.join(no_types)} — реклама, эквайринг и фин. рез. за эти дни пусты, итог неполон"] if no_types else [])
                               + ([f"лист «Заказы» не собран: {orders_error}"] if orders_error else [])}
        if orders is not None:
            t, nights = orders["totals"]["all"], sorted({n for c in orders["curve"].values() for n in c["nights"]})
            summary["orders"] = {"created": num(t["created_a"]), "forecast_confirmed": num(t["fc_a"]), "fin_result": num(t["fin_result"]),
                                 "drr_created": None if t["drr_created_pct"] is None else str(t["drr_created_pct"].quantize(Decimal("0.0001"))),
                                 "drr_forecast": None if t["drr_fc_pct"] is None else str(t["drr_fc_pct"].quantize(Decimal("0.0001"))),
                                 "curve_nights": f"{nights[0]} … {nights[-1]}" if nights else "нет", "mature_days": orders["mature"][0],
                                 "forecast_days": len(days) - orders["mature"][0]}
            summary["warnings"] += list(orders["said"])
            for sch, c in orders["curve"].items():
                share = ratio(c["biggest_event"][1], c["events_amt_total"])
                if share is not None and share >= Decimal("0.10") and sch in orders["sensitivity"]:
                    w, wo = orders["sensitivity"][sch][7]
                    rub = f"{c['biggest_event'][1]:,.0f}".replace(",", " ")
                    p1 = lambda v: f"{v * 100:.1f} %".replace(".", ",")  # noqa: E731
                    summary["warnings"].append(f"кривая {sch.upper()} по ₽ держится на одной отмене {rub} ₽ ({p1(share)} всех отмен в переходах): "
                                               f"на 7-е сутки «ещё отменится» {p1(w)} против {p1(wo)} без неё — прогноз {sch.upper()} пессимистичен")
        json.dump(summary, open(args.summary_json, "w"), ensure_ascii=False, indent=1)

    code = 2 if orders_error else 0
    if args.check:
        if not args.xlsx or not args.sheet:
            raise SystemExit("--check требует --xlsx и --sheet")
        manual = {d: {k: (None if v is None else Decimal(v)) for k, v in row.items()} for d, row in read_manual_sheet(args.xlsx, args.sheet).items()}
        table, failures = check(rows, manual)
        print_check(table)
        missing_in_ledger = [d for d in days if d not in ledger]
        no_file = [d for d in days if d not in file_types]
        sources_ok = not source_mismatch and not missing_in_ledger and not no_file
        print(f"\nисточники типов: леджер против файлов сырья — дат {len(days)}, из них моложе двух суток {len(young_days)} (живой ответ, равенства не требую"
              + (f"; доехало на {sorted(young_drift)}" if young_drift else "; доехавших типов нет") + f"); с обоими источниками {len(both)}, "
              f"совпали по каждому типу {len(both) - len(source_mismatch)}"
              + (f"; нет в леджере: {missing_in_ledger}" if missing_in_ledger else "") + (f"; нет файла: {no_file}" if no_file else "")
              + ("" if sources_ok else "   ОБЯЗАНЫ СОВПАСТЬ"))
        failures += 0 if sources_ok else 1
        print(f"приёмка: колонок, обязанных сходиться до копейки, — {len(EXACT)}, плюс совпадение источников; не сошлось {failures}")
        code = 1 if failures else code
    if args.check_segment:
        if not args.xlsx or not args.segment_sheet:
            raise SystemExit("--check-segment требует --xlsx и --segment-sheet")
        if args.check_segment not in metals:
            raise SystemExit(f"сегмент {args.check_segment!r} не строится; есть: {list(metals)}")
        manual_s = read_manual_segment_sheet(args.xlsx, args.segment_sheet)
        print_check_segment(args.check_segment, check_segment(metals[args.check_segment], manual_s))
    if args.check_orders:
        if not args.xlsx:
            raise SystemExit("--check-orders требует --xlsx")
        if orders is None:
            raise SystemExit(f"лист «Заказы» не собран: {orders_error}")
        manual_o = read_manual_orders_sheet(args.xlsx, args.orders_sheet)
        sheet_names = [x.strip() for x in args.orders_platform_sheets.split(",") if x.strip()]
        manual_p = {name: read_manual_orders_sheet(args.xlsx, sheet, ("revenue", "margin")) for sheet, (_p, name) in zip(sheet_names, PLATFORMS)}
        platform_rows = {name: orders["blocks"].get(f"platform:{name}", []) for _p, name in PLATFORMS}
        table_o, failures_o, decomposition = check_orders(orders["blocks"]["all"], manual_o, platform_rows, manual_p, young_days)
        print_check_orders(table_o, decomposition, young_days)
        if args.day_boundary == "utc":
            # сырьё окна — теми же функциями, что пересборка истории: файл history_<схема>_<от>_<до>.json есть — в API не идёт.
            # Последний UTC-день кончается в 21:00 UTC по МСК-окну, поэтому окно берётся до d2 + 1 день (не позже сегодня).
            import rebuild_ozon_orders_history as hist
            to = min(date.fromisoformat(d2) + timedelta(days=1), today_date).isoformat()
            raw = {scheme: hist.fetch_history(scheme, d1, to) for scheme in ("fbo", "fbs")}
            utc_rows = check_orders_utc(manual_p, created_by_utc_day(raw), days)
            print_check_orders_utc(utc_rows)
        print(f"приёмка листа «Заказы»: обязанная колонка одна (реклама по образцу, дни старше двух суток); не сошлось {failures_o}")
        code = 1 if failures_o else code
    print("db_writes = 0")
    sys.exit(code)


if __name__ == "__main__":
    main()
