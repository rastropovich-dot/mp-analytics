#!/usr/bin/env python3
"""Лист «WB - <месяц>» по образцу ручного листа владельца — из отчёта реализации WB.

    venv/bin/python3 scripts/report_wb_month.py --month 2026-09                     → data/reports/wb_2026-09_to_<дата>.xlsx
    venv/bin/python3 scripts/report_wb_month.py --month 2026-09 --date-to 2026-09-21 --rows-from-files data/wb_sales_report_raw
    venv/bin/python3 scripts/report_wb_month.py --month 2026-09 --date-to 2026-09-21 --check \\
        --xlsx ../mp-analytics/data/manual_report_september_20260922_v2.xlsx --sheet "WB - сентябрь"

ИСТОЧНИК — строки отчёта реализации (wb_sales_report_rows, зерно rrdId; или те же
строки из файлов сырья бэкфилла — --rows-from-files, пока таблицы нет). ДЕНЬ СТРОКИ —
saleDt в московском времени (у строки без saleDt — rrDate): так датирует владелец,
проверено 21 из 21 по обороту, комиссии, выручке, логистике и прочему за 1…21.09
(WB-5 §3). rrDate отстаёт от даты продажи у 65 строк из 32 009 (02-01…09-22), обычно
на день, поэтому строки читаются по rrDate с запасом LAG_DAYS, а строки продаж с датой
вне листа считаются и печатаются. Карта колонок — docs/wb_report_model.md.

КОЛОНКИ (по дню, суммы за день; знак: Продажа +, Возврат −):
    Оборот              Σ retailPriceWithDisc (продажа − возврат) — = marketplace_buyouts, до рубля
    Комиссия            Σ retailPriceWithDisc × commissionPercent / 100 (кВВ строки, знак возврата) — его C
                        (решение советника 2026-09-23: колонка листа — как у владельца)
    Удержано из выплаты Оборот − Σ forPay± («К перечислению продавцу»), по отчёту 44,4 % — справочно:
                        всё, что WB удержал из выплаты (комиссия, скидки, доплаты), не одна комиссия
    НДС за возмещение   (Σ ppvzReward + Σ rebillLogisticCost) × (НДС − 1) / НДС — колонка AB его «свода»
    Выручка             (Оборот − Комиссия) / НДС − НДС за возмещение — его E, до копейки
    Себестоимость       снимок 1С (article_unit_costs) по базовому артикулу (unit_cost_uniform), для
                        безразмерных — точная строка; --cogs-by variant — по варианту артикул-размер
    Логистика           Σ deliveryService / НДС — его I до копейки при дате saleDt МСК
    Реклама             Σ updSum за день (updTime по МСК) / НДС — у владельца реклама без НДС (июль 1 694 178 =
                        Σ updSum / 1,22 до рубля); источник по умолчанию — таблица wb_ad_spend_daily,
                        --ads api — живой adv/v1/upd (1 обращение, ручной путь), --ads none — пусто, не ноль
    Эквайринг           Σ acquiringFee (продажа +, возврат −) / НДС — его M; 09-17 он возвраты не вычел
    Прочее              Σ paidStorage / НДС + Σ penalty (штрафы без НДС) + Σ deduction / НДС — его O
    Фин. рез.           Маржа − Логистика − Реклама − Эквайринг − Прочее
    Накладные в день    OVERHEAD_PER_DAY["wb"] с датой действия (T81 листа владельца: 74 002,00 с 09-01)
    Ebitda              Фин. рез. − Накладные

НДС — VAT_RATES из scripts/report_ozon_month.py (один параметр на обе площадки).
Дни моложе двух суток — жёлтые: строки дня появляются на D+1 и могут доезжать.
Только чтение; db_writes = 0.
"""
import argparse
import glob
import json
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(ROOT, ".env"))

from loaders import stale_keys  # noqa: E402
import loaders.wb_sales_report_loader as loader  # noqa: E402
import loaders.wb_ads_loader as ads_loader  # noqa: E402
from reconcile_manual_report_ozon import read_manual_sheet  # noqa: E402
from report_ozon_month import MONTHS, VAT_RATES, vat_for  # noqa: E402,F401  — тот же НДС, что у Ozon

Z = Decimal(0)
C = Decimal("0.01")
SNAP = "2026-05-20"
OUT_DIR = os.path.join("data", "reports")
YOUNG_DAYS = 2
MSK = timezone(timedelta(hours=3))
LAG_DAYS = 7   # rrDate позже даты продажи: 60 строк из 32 009 на 1…3 дня, 5 — на 16…115 (те на лист своего месяца не попадут, печатаются)
# Накладные в день по площадке с датой действия — ячейка T81 листа «WB - сентябрь» владельца (22.09).
OVERHEAD_PER_DAY = {"wb": (("2026-09-01", Decimal("74002.00")),)}
# Расхождения с листом на стороне владельца (WB-5 §3): печатаем, не чиним, в отказ не считаем.
# 09-17: эквайринг двух возвратов (441,25 + 936,24 с НДС) он не вычел, на 09-03 / 09-07 / 09-14 — вычел.
# Прежние три «известных дня» (эквайринг 09-03 и 09-07, прочее 09-02) сняты 2026-09-23: это были наши
# формулы — знак возврата в эквайринге и штрафы без НДС, — а не его сторона.
KNOWN_MANUAL_DIFFS = {("acquiring", "2026-09-17")}
MONEY_FIELDS = ("retail_price_with_disc", "retail_amount", "for_pay", "commission_percent", "ppvz_reward", "rebill_logistic_cost",
                "delivery_service", "acquiring_fee", "paid_storage", "penalty", "deduction", "cashback_discount")
SELECT = ("rrd_id,rr_date,sale_dt,seller_oper_name,doc_type,vendor_code,tech_size,nm_id,quantity," + ",".join(MONEY_FIELDS))


def D(v):
    return Decimal(str(v)) if v not in (None, "") else Z


def q(v):
    return Decimal(v).quantize(C)


def ratio(a, b):
    return (Decimal(a) / Decimal(b)).quantize(Decimal("0.0001")) if b else None


def overhead_for(day, platform="wb"):
    for valid_from, value in OVERHEAD_PER_DAY.get(platform, ()):
        if day >= valid_from:
            return value
    return None


def month_days(month, date_to=None):
    import calendar
    y, m = int(month[:4]), int(month[5:7])
    last = date(y, m, calendar.monthrange(y, m)[1])
    if date_to:
        last = min(last, date.fromisoformat(date_to))
    out, d = [], date(y, m, 1)
    while d <= last:
        out.append(d.isoformat()); d += timedelta(days=1)
    return out


def row_day(r):
    """День строки на листе: saleDt в московском времени; без saleDt (или с датой-заглушкой) — rrDate.
    Правило владельца: оборот, комиссия, выручка, логистика и прочее сошлись 21 из 21 только так (WB-5 §3);
    по rrDate оборот расходился 09-09 / 09-10 (четыре продажи 09-09 16:50–18:44 UTC с rrDate 09-10)."""
    value = r.get("sale_dt")
    if value:
        try:
            ts = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            ts = None
        if ts is not None and ts.year >= 2000:
            if ts.tzinfo is None:
                ts = ts.replace(tzinfo=timezone.utc)
            return ts.astimezone(MSK).date().isoformat()
    return str(r["rr_date"])[:10]


def window_end(d2):
    """Последний rrDate, который читаем ради строк с датой продажи ≤ d2."""
    return (date.fromisoformat(d2) + timedelta(days=LAG_DAYS)).isoformat()


# ---------- источники ----------

def load_rows_db(sb, d1, d2):
    """Строки по rrDate d1 … d2 + LAG_DAYS; день строки на листе даёт row_day."""
    return stale_keys.read_window_rows(sb, loader.TABLE, SELECT, [("gte", "rr_date", d1), ("lte", "rr_date", window_end(d2))], ["rr_date", "rrd_id"])


def load_rows_files(raw_dir, d1, d2):
    """Те же строки из файлов сырья бэкфилла, через loader.build_row — колонки как в таблице."""
    rows, seen = [], set()
    d2x = window_end(d2)
    for path in sorted(glob.glob(os.path.join(raw_dir, "daily_*.json"))):
        name = os.path.basename(path)[len("daily_"):-len(".json")]
        f1, f2 = name.split("_")
        if f2 < d1 or f1 > d2x:
            continue
        for item in json.load(open(path, encoding="utf-8"), parse_float=Decimal):
            day = str(item["rrDate"])[:10]
            if not (d1 <= day <= d2x):
                continue
            rid = int(item["rrdId"])
            if rid in seen:
                raise RuntimeError(f"rrdId {rid} встречается дважды в файлах сырья")
            seen.add(rid)
            rows.append(loader.build_row(item, "files"))
    return rows


def load_costs(sb, snapshot=SNAP):
    """Снимок 1С целиком (offer_id_norm, unit_cost, unit_cost_uniform): точные строки и «единая» по базовому артикулу."""
    rows = stale_keys.read_window_rows(sb, "article_unit_costs", "offer_id_norm,unit_cost,unit_cost_uniform",
                                       [("eq", "snapshot_date", snapshot)], ["offer_id_norm"])
    exact, uniform = {}, {}
    for r in rows:
        exact[r["offer_id_norm"]] = D(r["unit_cost"])
        base = r["offer_id_norm"].split("-", 1)[0]
        if r.get("unit_cost_uniform") is not None and base not in uniform:
            uniform[base] = D(r["unit_cost_uniform"])
    return exact, uniform


def unit_cost_for(exact, uniform, vendor_code, tech_size, mode="base"):
    """(себестоимость, источник): variant — точный вариант артикул-размер; plain — безразмерная точная строка;
    uniform — «единая» по базовому артикулу; none — в снимке нет."""
    code = str(vendor_code or "").strip().lower()
    if not code:
        return None, "none"
    size = str(tech_size or "").strip()
    if mode == "variant" and size and size != "0":
        v = exact.get(f"{code}-{size}")
        if v is not None:
            return v, "variant"
    if code in exact:
        return exact[code], "plain"
    if code in uniform:
        return uniform[code], "uniform"
    return None, "none"


def fetch_ads(d1, d2):
    """Ручной путь: живой advert-api /adv/v1/upd (интервал ≤ 31 день), одно обращение; день — updTime по МСК."""
    by, undated = defaultdict(Decimal), Z
    for x in ads_loader.request_upd(d1, d2):
        when = x.get("updTime")
        if when:
            by[ads_loader.upd_day(when)] += D(x.get("updSum"))
        else:
            undated += D(x.get("updSum"))
    return dict(by), undated


def load_ads_db(sb, d1, d2):
    """Списания из wb_ad_spend_daily по upd_day (updSum с НДС). Таблицы нет — RuntimeError, реклама остаётся неизвестной."""
    try:
        rows = stale_keys.read_window_rows(sb, ads_loader.TABLE, "advert_id,upd_time,upd_day,upd_sum",
                                           [("gte", "upd_day", d1), ("lte", "upd_day", d2)], ["upd_day", "advert_id", "upd_time"])
    except Exception as error:
        raise RuntimeError(f"не прочитать {ads_loader.TABLE}: {str(error)[:160]}")
    by = defaultdict(Decimal)
    for r in rows:
        by[str(r["upd_day"])] += D(r["upd_sum"])
    return dict(by), Z


# ---------- расчёт ----------

def build_daily(rows, days, cost_fn, ads_by_day, today, ads_known=True, outside=None):
    """Строки листа по дням. cost_fn(vendor_code, tech_size) → (себестоимость, источник).
    outside (если передан словарь) получает продажи/возвраты с датой продажи вне листа: {"before"|"after": [строк, Σ оборот]}."""
    by = {d: defaultdict(Decimal) for d in days}
    cnt = {d: defaultdict(int) for d in days}
    if outside is not None:
        outside.setdefault("before", [0, Z]); outside.setdefault("after", [0, Z])
    for r in rows:
        d = row_day(r)
        op = r["seller_oper_name"]
        if d not in by:
            if outside is not None and op in ("Продажа", "Возврат") and days:
                side = outside["before" if d < days[0] else "after"]
                side[0] += 1; side[1] += (1 if op == "Продажа" else -1) * D(r["retail_price_with_disc"])
            continue
        s = by[d]
        if op in ("Продажа", "Возврат"):
            sign = 1 if op == "Продажа" else -1
            price = D(r["retail_price_with_disc"])
            s["turnover"] += sign * price
            s["for_pay"] += sign * D(r["for_pay"])
            pct = r.get("commission_percent")
            if pct in (None, ""):
                cnt[d]["no_pct_rows"] += 1
            s["commission"] += sign * price * D(pct) / 100
            s["acquiring_raw"] += sign * D(r["acquiring_fee"])
            qty = int(r.get("quantity") or 1)
            cost, source = cost_fn(r.get("vendor_code"), r.get("tech_size"))
            cnt[d]["positions"] += sign * qty
            cnt[d]["cost_" + source] += qty
            if cost is None:
                cnt[d]["no_cost_positions"] += qty
                s["no_cost_turnover"] += sign * price   # доля оборота без СС — вслух (задача WB-5 §3)
            else:
                s["cogs"] += sign * cost * qty
        else:
            s["acquiring_other"] += D(r["acquiring_fee"])   # эквайринг вне продаж/возвратов — в лист не входит, печатается, если не ноль
        s["reward"] += D(r["ppvz_reward"]); s["rebill"] += D(r["rebill_logistic_cost"])
        s["delivery"] += D(r["delivery_service"])
        s["storage"] += D(r["paid_storage"]); s["penalty"] += D(r["penalty"]); s["deduction"] += D(r.get("deduction"))
        cnt[d]["rows"] += 1
    out = []
    for d in days:
        s, c, vat = by[d], cnt[d], vat_for(d)
        vat_refund = (s["reward"] + s["rebill"]) * (vat - 1) / vat
        commission = s["commission"]                    # как у владельца: кВВ строк (решение советника 2026-09-23)
        withheld = s["turnover"] - s["for_pay"]         # справочно: удержано из выплаты всего
        revenue = (s["turnover"] - commission) / vat - vat_refund
        cogs = s["cogs"]
        ads = (ads_by_day.get(d, Z) / vat) if ads_known else None   # updSum с НДС → без, как на листе владельца
        logistics, acquiring = s["delivery"] / vat, s["acquiring_raw"] / vat
        other = s["storage"] / vat + s["penalty"] + s["deduction"] / vat   # штрафы у владельца без НДС — 21 из 21 только так
        margin = revenue - cogs
        fin = margin - logistics - (ads or Z) - acquiring - other
        overhead = overhead_for(d)
        young = (date.fromisoformat(d) >= today - timedelta(days=YOUNG_DAYS - 1))
        row = {"date": d, "vat": vat, "rows": c["rows"], "positions": c["positions"], "no_cost_positions": c["no_cost_positions"], "no_pct_rows": c["no_pct_rows"],
               "cost_sources": {k[5:]: v for k, v in c.items() if k.startswith("cost_")},
               "turnover": s["turnover"], "commission": commission, "commission_pct": ratio(commission, s["turnover"]),
               "withheld": withheld, "withheld_pct": ratio(withheld, s["turnover"]), "vat_refund": vat_refund, "revenue": revenue,
               "cogs": cogs, "margin": margin, "margin_pct": ratio(margin, revenue), "logistics": logistics, "logistics_pct": ratio(logistics, revenue),
               "ads": ads, "drr_pct": ratio(ads, s["turnover"]) if ads is not None else None, "acquiring": acquiring, "acquiring_pct": ratio(acquiring, revenue),
               "other": other, "fin_result": fin, "fin_result_pct": ratio(fin, revenue), "overhead": overhead,
               "ebitda": (fin - overhead) if overhead is not None else None, "young": young,
               "reward": s["reward"], "rebill": s["rebill"], "storage": s["storage"], "penalty": s["penalty"], "deduction": s["deduction"],
               "acquiring_other": s["acquiring_other"], "no_cost_turnover": s["no_cost_turnover"]}
        row["ebitda_pct"] = ratio(row["ebitda"], revenue) if row["ebitda"] is not None else None
        out.append(row)
    return out


def total_row(rows):
    t = {"date": "Итого", "young": False, "vat": rows[-1]["vat"] if rows else vat_for("2026-01-01")}
    for k in ("rows", "positions", "no_cost_positions", "no_pct_rows"):
        t[k] = sum(r[k] for r in rows)
    for k in ("turnover", "commission", "withheld", "vat_refund", "revenue", "cogs", "margin", "logistics", "acquiring", "other", "fin_result", "reward", "rebill", "storage", "penalty", "deduction", "acquiring_other", "no_cost_turnover"):
        t[k] = sum((r[k] for r in rows), Z)
    t["ads"] = sum((r["ads"] for r in rows if r["ads"] is not None), Z) if any(r["ads"] is not None for r in rows) else None
    t["overhead"] = sum((r["overhead"] for r in rows if r["overhead"] is not None), Z) if any(r["overhead"] is not None for r in rows) else None
    t["ebitda"] = (t["fin_result"] - t["overhead"]) if t["overhead"] is not None else None
    t["commission_pct"] = ratio(t["commission"], t["turnover"]); t["withheld_pct"] = ratio(t["withheld"], t["turnover"])
    t["margin_pct"] = ratio(t["margin"], t["revenue"])
    t["logistics_pct"] = ratio(t["logistics"], t["revenue"]); t["acquiring_pct"] = ratio(t["acquiring"], t["revenue"])
    t["drr_pct"] = ratio(t["ads"], t["turnover"]) if t["ads"] is not None else None
    t["fin_result_pct"] = ratio(t["fin_result"], t["revenue"]); t["ebitda_pct"] = ratio(t["ebitda"], t["revenue"]) if t["ebitda"] is not None else None
    srcs = defaultdict(int)
    for r in rows:
        for k, v in r["cost_sources"].items():
            srcs[k] += v
    t["cost_sources"] = dict(srcs)
    return t


# ---------- лист «Заказы WB» (блок J–P листа «Заказы» владельца; WB-7 §4) ----------

OWNER_ORDERS_AFTER_COMMISSION = Decimal("0.58")   # K владельца = orderSum воронки × 0,58 / НДС (WB-6 §2: 16 из 21 до копейки)
DISCOUNTER_LETTER = "t"                            # «WB Дискаунтер» = артикулы `t…` (бренд «КОЮЗ Топаз», подтверждено владельцем 09-24)
PLATFORMS = (("all", "Заказы WB"), ("standard", "Заказы WB Standard"), ("discounter", "Заказы WB Дискаунтер"))
FUNNEL_TABLE = "wb_funnel_products_daily"
FUNNEL_SELECT = "day,nm_id,vendor_code,order_count,order_sum,buyout_count,buyout_sum,cancel_count,cancel_sum"
COINVEST_TOLERANCE = Decimal("0.005")               # ожидание задачи: P в пределах 0,005 — 15 из 22


def platform_of(vendor_code):
    """Площадка по первой букве артикула: `t` — Дискаунтер, остальное — Standard ('Заказы Standard'!K + 'WB Дискаунтер'!B = K)."""
    return "discounter" if str(vendor_code or "").strip()[:1].lower() == DISCOUNTER_LETTER else "standard"


def load_funnel_db(sb, d1, d2):
    """Воронка по товарам из wb_funnel_products_daily (день заказа МСК). Таблицы нет — RuntimeError, листы заказов не собираются."""
    try:
        return stale_keys.read_window_rows(sb, FUNNEL_TABLE, FUNNEL_SELECT, [("gte", "day", d1), ("lte", "day", d2)], ["day", "nm_id"])
    except Exception as error:
        raise RuntimeError(f"не прочитать {FUNNEL_TABLE}: {str(error)[:160]}")


def load_funnel_files(raw_dir, d1, d2):
    """Те же строки из сырья бэкфилла воронки (data/wb_funnel_raw/funnel_<день>.json, все страницы дня)."""
    rows = []
    for path in sorted(glob.glob(os.path.join(raw_dir, "funnel_*.json"))):
        day = os.path.basename(path)[len("funnel_"):-len(".json")]
        if not (d1 <= day <= d2):
            continue
        payload = json.load(open(path, encoding="utf-8"))
        for p in payload.get("products") or []:
            product, st = p.get("product") or {}, (p.get("statistic") or {}).get("selected") or {}
            if product.get("nmId") in (None, ""):
                continue
            rows.append({"day": day, "nm_id": int(product["nmId"]), "vendor_code": product.get("vendorCode"),
                         "order_count": st.get("orderCount"), "order_sum": st.get("orderSum"), "buyout_count": st.get("buyoutCount"),
                         "buyout_sum": st.get("buyoutSum"), "cancel_count": st.get("cancelCount"), "cancel_sum": st.get("cancelSum")})
    return rows


def coinvest_by_day(report_rows):
    """P владельца (рабочее правило советника, WB-7): (Σ retailPriceWithDisc − Σ retailAmount) / Σ retailPriceWithDisc
    по строкам «Продажа» дня (saleDt МСК) — всего и по площадке. {день: {all|standard|discounter: [Σ цена, Σ оплачено]}}."""
    acc = defaultdict(lambda: defaultdict(lambda: [Z, Z]))
    for r in report_rows:
        if r["seller_oper_name"] != "Продажа":
            continue
        d, pf = row_day(r), platform_of(r.get("vendor_code"))
        for key in ("all", pf):
            acc[d][key][0] += D(r["retail_price_with_disc"]); acc[d][key][1] += D(r["retail_amount"])
    return acc


def build_orders_daily(funnel_rows, days, cost_fn, ads_by_day, coinvest, today, ads_known=True, platform="all"):
    """Строки листа заказов по дням: K = Σ orderSum × 0,58 / НДС; L = K − Σ orderCount × СС (базовый артикул);
    реклама — только на общем листе (списания по кампаниям на площадки не делятся); P — из отчёта реализации."""
    by = {d: defaultdict(Decimal) for d in days}
    cnt = {d: defaultdict(int) for d in days}
    for r in funnel_rows:
        d = str(r["day"])
        if d not in by:
            continue
        pf = platform_of(r.get("vendor_code"))
        if platform != "all" and pf != platform:
            continue
        s, c = by[d], cnt[d]
        qty = int(r.get("order_count") or 0)
        s["orders_qty"] += qty; s["orders_sum"] += D(r.get("order_sum"))
        s["funnel_buyouts_sum"] += D(r.get("buyout_sum")); s["funnel_cancel_sum"] += D(r.get("cancel_sum"))
        c["funnel_buyouts_qty"] += int(r.get("buyout_count") or 0); c["funnel_cancel_qty"] += int(r.get("cancel_count") or 0)
        c["cards"] += 1
        if qty:
            cost, source = cost_fn(r.get("vendor_code"), None)
            c["cost_" + source] += qty
            if cost is None:
                c["no_cost_qty"] += qty; s["no_cost_sum"] += D(r.get("order_sum"))
            else:
                s["cogs"] += cost * qty
    out = []
    for d in days:
        s, c, vat = by[d], cnt[d], vat_for(d)
        revenue = s["orders_sum"] * OWNER_ORDERS_AFTER_COMMISSION / vat
        margin = revenue - s["cogs"]
        ads = (ads_by_day.get(d, Z) / vat) if (ads_known and platform == "all") else None
        ci = coinvest.get(d, {}).get(platform)
        coinvest_pct = ((ci[0] - ci[1]) / ci[0]).quantize(Decimal("0.0001")) if ci and ci[0] else None
        young = (date.fromisoformat(d) >= today - timedelta(days=YOUNG_DAYS - 1))
        out.append({"date": d, "vat": vat, "cards": c["cards"], "orders_qty": int(s["orders_qty"]), "orders_sum": s["orders_sum"],
                    "revenue": revenue, "cogs": s["cogs"], "margin": margin, "margin_pct": ratio(margin, revenue),
                    "ads": ads, "drr_pct": ratio(ads, s["orders_sum"]) if ads is not None else None, "coinvest_pct": coinvest_pct,
                    "funnel_buyouts_sum": s["funnel_buyouts_sum"], "funnel_buyouts_qty": c["funnel_buyouts_qty"],
                    "funnel_cancel_sum": s["funnel_cancel_sum"], "funnel_cancel_qty": c["funnel_cancel_qty"],
                    "no_cost_qty": c["no_cost_qty"], "no_cost_sum": s["no_cost_sum"],
                    "cost_sources": {k[5:]: v for k, v in c.items() if k.startswith("cost_")}, "young": young})
    return out


def orders_total_row(rows):
    t = {"date": "Итого", "young": False, "vat": rows[-1]["vat"] if rows else vat_for("2026-01-01"), "coinvest_pct": None}
    for k in ("cards", "orders_qty", "funnel_buyouts_qty", "funnel_cancel_qty", "no_cost_qty"):
        t[k] = sum(r[k] for r in rows)
    for k in ("orders_sum", "revenue", "cogs", "margin", "funnel_buyouts_sum", "funnel_cancel_sum", "no_cost_sum"):
        t[k] = sum((r[k] for r in rows), Z)
    t["ads"] = sum((r["ads"] for r in rows if r["ads"] is not None), Z) if any(r["ads"] is not None for r in rows) else None
    t["margin_pct"] = ratio(t["margin"], t["revenue"])
    t["drr_pct"] = ratio(t["ads"], t["orders_sum"]) if t["ads"] is not None else None
    srcs = defaultdict(int)
    for r in rows:
        for k, v in r["cost_sources"].items():
            srcs[k] += v
    t["cost_sources"] = dict(srcs)
    return t


ORDERS_COLS = [("День", "date", None), ("Заказано, шт", "orders_qty", "int"), ("Заказано на сумму (с НДС), руб.", "orders_sum", "money"),
               ("Выручка (× 0,58 / НДС), руб.", "revenue", "money"), ("Себестоимость (заказы × СС снимка), руб.", "cogs", "money"),
               ("Маржа, руб.", "margin", "money"), ("М-ть, %", "margin_pct", "pct"), ("Реклама, руб.", "ads", "money"), ("% ДРР (от заказов)", "drr_pct", "pct"),
               ("Соинвест, % = (Σ цена − Σ оплачено) / Σ цена по продажам дня (saleDt МСК)", "coinvest_pct", "pct"),
               (None, None, None), ("справочно: выкупы воронки по дню заказа, руб.", "funnel_buyouts_sum", "money"), ("выкупы воронки, шт", "funnel_buyouts_qty", "int"),
               ("отмены и возвраты воронки, руб.", "funnel_cancel_sum", "money"), ("отмены воронки, шт", "funnel_cancel_qty", "int"),
               ("карточек", "cards", "int"), ("заказов без СС, шт", "no_cost_qty", "int"), ("Примечание", "note", None)]


def _write_orders_sheet(wb, title, rows, total, notes):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fmt = {"money": "#,##0.00", "pct": "0.0%", "int": "#,##0"}
    bold, young_fill, head_fill = Font(bold=True), PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9E1F2")
    ws = wb.create_sheet(title)
    ws["A1"] = f"Статистика созданных заказов Wildberries — {title}"; ws["A1"].font = bold
    for j, (head, _k, _f) in enumerate(ORDERS_COLS, 1):
        if head:
            c = ws.cell(row=3, column=j, value=head); c.font = bold; c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for i, r in enumerate(rows + [total], 4):
        r = dict(r)
        r["note"] = "; ".join(x for x in (
            "моложе двух суток — воронка пересматривает день" if r.get("young") else "",
            f"без СС {r['no_cost_qty']} заказов на {r['no_cost_sum']:,.0f}" if r.get("no_cost_qty") else "",
            "реклама на площадки не делится" if r.get("ads") is None and r["date"] != "Итого" and title != PLATFORMS[0][1] else "") if x)
        for j, (head, k, f) in enumerate(ORDERS_COLS, 1):
            if not head:
                continue
            v = r.get(k)
            if k == "date" and v != "Итого":
                v = date.fromisoformat(v)
            elif isinstance(v, Decimal):
                v = float(v)
            c = ws.cell(row=i, column=j, value=v)
            if f:
                c.number_format = fmt[f]
            if k == "date" and r["date"] != "Итого":
                c.number_format = "DD.MM.YYYY"
            if r.get("young"):
                c.fill = young_fill
            if r["date"] == "Итого":
                c.font = bold
    base = 4 + len(rows) + 2
    for n, line in enumerate(notes):
        ws.cell(row=base + n, column=1, value=line)
    ws.freeze_panes = "B4"; ws.row_dimensions[3].height = 60
    for j in range(1, len(ORDERS_COLS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15 if j > 1 else 12
    ws.column_dimensions[get_column_letter(len(ORDERS_COLS))].width = 60


# ---------- книга ----------

COLS = [("Дата реализации", "date", None), ("Оборот (с НДС), руб.", "turnover", "money"), ("Комиссия (с НДС), руб.", "commission", "money"),
        ("Комиссия, %", "commission_pct", "pct"), ("справочно: удержано из выплаты всего (оборот − forPay), руб.", "withheld", "money"),
        ("справочно: удержано, %", "withheld_pct", "pct"), ("Выручка, руб.", "revenue", "money"),
        ("Себестоимость, руб.", "cogs", "money"), ("Маржа, руб.", "margin", "money"), ("Мар-ть, %", "margin_pct", "pct"),
        ("Логистика, руб.", "logistics", "money"), ("% Логистики", "logistics_pct", "pct"), ("Реклама, руб.", "ads", "money"),
        ("% ДРР (от ТО)", "drr_pct", "pct"), ("Эквайринг, руб.", "acquiring", "money"), ("% Эквайринга", "acquiring_pct", "pct"),
        ("Прочее, руб.", "other", "money"), ("Фин. рез., руб.", "fin_result", "money"), ("% Фин. рез.", "fin_result_pct", "pct"),
        ("Накладные в день, руб.", "overhead", "money"), ("Ebitda, руб.", "ebitda", "money"), ("% Ebitda", "ebitda_pct", "pct"),
        (None, None, None), ("НДС за возмещение (вычет), руб.", "vat_refund", "money"), ("Возмещение ПВЗ (с НДС)", "reward", "money"),
        ("Возмещение перемещения (с НДС)", "rebill", "money"), ("Хранение (с НДС)", "storage", "money"), ("Штрафы (без НДС, как в отчёте)", "penalty", "money"),
        ("Удержания (с НДС)", "deduction", "money"),
        ("Позиций (продажи − возвраты)", "positions", "int"), ("позиций без себестоимости", "no_cost_positions", "int"),
        ("строк отчёта", "rows", "int"), ("Примечание", "note", None)]


def write_xlsx(path, month, rows, total, notes, orders=None):
    """Книга месяца: лист «WB - <месяц>» и, если переданы, листы заказов (orders — [(название, строки, итог, подвал)])."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fmt = {"money": "#,##0.00", "pct": "0.0%", "int": "#,##0"}
    bold, young_fill, head_fill = Font(bold=True), PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9E1F2")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"WB - {MONTHS[int(month[5:7]) - 1]}"
    ws["A1"] = "Выкупы WB"; ws["A1"].font = bold
    ws["E2"] = "Вычитается НДС за возмещение из выручки! (как на листе владельца)"
    for j, (head, _k, _f) in enumerate(COLS, 1):
        if head:
            c = ws.cell(row=3, column=j, value=head); c.font = bold; c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for i, r in enumerate(rows + [total], 4):
        r = dict(r)
        r["note"] = "; ".join(x for x in (
            "моложе двух суток — строки отчёта доезжают" if r.get("young") else "",
            f"без СС {r['no_cost_positions']} из {r['positions']} позиций" if r.get("no_cost_positions") else "",
            "реклама не запрашивалась" if r.get("ads") is None else "") if x)
        for j, (head, k, f) in enumerate(COLS, 1):
            if not head:
                continue
            v = r.get(k)
            if k == "date" and v != "Итого":
                v = date.fromisoformat(v)
            elif isinstance(v, Decimal):
                v = float(v)
            c = ws.cell(row=i, column=j, value=v)
            if f:
                c.number_format = fmt[f]
            if k == "date" and r["date"] != "Итого":
                c.number_format = "DD.MM.YYYY"
            if r.get("young"):
                c.fill = young_fill
            if r["date"] == "Итого":
                c.font = bold
    base = 4 + len(rows) + 2
    for n, line in enumerate(notes):
        ws.cell(row=base + n, column=1, value=line)
    ws.freeze_panes = "B4"; ws.row_dimensions[3].height = 48
    for j in range(1, len(COLS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15 if j > 1 else 16
    ws.column_dimensions[get_column_letter(len(COLS))].width = 60
    for o_title, o_rows, o_total, o_notes in (orders or []):
        _write_orders_sheet(wb, o_title, o_rows, o_total, o_notes)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


# ---------- приёмка ----------

CHECK_LINES = (
    # (заголовок, наше поле, их колонка, обязана сходиться, чем объясняется расхождение)
    ("Оборот (с НДС)", "turnover", "turnover", True, ""),
    ("Выручка (оборот − комиссия) / НДС − НДС за возмещение", "revenue", "revenue", True, ""),
    ("Эквайринг (acquiringFee, возврат со знаком минус) / НДС", "acquiring", "acquiring", True,
     "09-17 владелец не вычел эквайринг двух возвратов (441,25 + 936,24 с НДС = 1 129,09 без) — на 09-03, 09-07, 09-14 вычел; его сторона"),
    ("Прочее (хранение / НДС + штрафы без НДС + удержания / НДС)", "other", "other", True, ""),
    ("Комиссия (Σ цена × кВВ строки, знак возврата)", "commission", "commission", True, ""),
    ("Логистика (deliveryService / НДС, по дате продажи МСК)", "logistics", "logistics", True, ""),
    ("справочно: удержано из выплаты всего (оборот − forPay) против его комиссии", "withheld", "commission", False,
     "удержано всё: комиссия по кВВ + согласованная скидка + доплаты; по отчёту 44,4 % против кВВ ~42 % — не одно и то же"),
    ("Себестоимость (снимок 1С 05-20 по базовому артикулу против его цен)", "cogs", "cogs", False, "снимок 05-20 и стыковка по базовому артикулу (WB-5 §4)"),
    ("Фин. рез. против его", "fin_result", "fin_result", False, "разница — СС (и эквайринг 09-17)"),
)


def check(rows, manual, young_days=()):
    out, failures = [], 0
    for title, ours_k, theirs_k, must, why in CHECK_LINES:
        t_o = t_t = Z
        bad, missing, known = [], [], []
        for r in rows:
            if r["date"] in young_days:
                continue
            m = manual.get(r["date"])
            o = r.get(ours_k)
            t = None if not m or m.get(theirs_k) is None else Decimal(m[theirs_k])
            if o is None or t is None:
                missing.append(r["date"]); continue
            o, t = q(o), q(t)
            t_o += o; t_t += t
            if o != t:
                (known if (ours_k, r["date"]) in KNOWN_MANUAL_DIFFS else bad).append((r["date"], o - t))
        days = len([r for r in rows if r["date"] not in young_days])
        if must and (bad or missing):
            failures += 1
        out.append({"title": title, "theirs": t_t, "ours": t_o, "diff": t_o - t_t, "equal_days": days - len(bad) - len(known) - len(missing),
                    "days": days, "bad": bad, "known": known, "missing": missing, "must": must, "why": why})
    return out, failures


def print_check(table):
    print(f"\n{'колонка':66}{'у него':>17}{'у нас':>17}{'разница':>15}{'дней 0,00':>11}")
    for t in table:
        print(f"{t['title'][:65]:66}{t['theirs']:>17,.2f}{t['ours']:>17,.2f}{t['diff']:>15,.2f}{t['equal_days']:>7} из {t['days']}"
              + ("   ОБЯЗАНА СХОДИТЬСЯ" if t["must"] and (t["bad"] or t["missing"]) else ""))
        if t["known"]:
            print("      известные расхождения (его сторона, не чиним): " + ", ".join(f"{d} {v:+,.2f}" for d, v in t["known"]))
        if t["bad"] and (t["must"] or len(t["bad"]) <= 8):
            print("      расходится: " + ", ".join(f"{d} {v:+,.2f}" for d, v in t["bad"]))
        elif t["bad"]:
            print(f"      расходится на {len(t['bad'])} днях: от {min(v for _d, v in t['bad']):+,.2f} до {max(v for _d, v in t['bad']):+,.2f}")
        if t["missing"]:
            print("      нет значения: " + ", ".join(t["missing"]))
        if t["why"] and (t["bad"] or t["known"] or t["missing"]):
            print(f"      чем объясняется: {t['why']}")


# ---------- приёмка листов заказов ----------

ORDERS_MANUAL_JP = ["revenue", "margin", "margin_pct", "ads", "drr_pct", "coinvest_pct"]   # K…P блока «Заказы» и «Заказы Standard»
ORDERS_MANUAL_BD = ["revenue", "margin", "margin_pct"]                                      # B…D листа «WB Дискаунтер»


def read_manual_orders(path, sheet, date_col, first_col, names):
    """Блок заказов книги владельца: дата в колонке date_col (0 = A), значения с first_col по именам names. Файл большой — read_only."""
    import datetime
    import warnings
    import openpyxl
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for row in wb[sheet].iter_rows(min_row=3, max_row=60, values_only=True):
        d = row[date_col] if date_col < len(row) else None
        if isinstance(d, datetime.datetime):
            out[d.date().isoformat()] = {n: (repr(row[first_col + i]) if first_col + i < len(row) and row[first_col + i] is not None else None)
                                         for i, n in enumerate(names)}
    return out


ORDERS_CHECK_LINES = (
    # (площадка, заголовок, наше поле, их поле, вид сравнения, ожидание задачи, чем объясняется)
    ("all", "«Заказы» K Выручка = воронка × 0,58 / НДС", "revenue", "revenue", "exact", "16 из 21", "пять дней у владельца выше любой нашей воронки — его снимки (WB-6 §2), не искать"),
    ("all", "«Заказы» L Маржа = K − СС", "margin", "margin", "exact", "0 из 21", "L — по его ценам 1С, у нас снимок 05-20 по базовому артикулу"),
    ("all", "«Заказы» N Реклама", "ads", "ads", "exact", "21 из 21", ""),
    ("all", "«Заказы» P Соинвест = (Σ цена − Σ оплачено) / Σ цена по продажам дня", "coinvest_pct", "coinvest_pct", "tolerance", "15 из 22 в пределах 0,005", "09-03 и 09-14 у владельца — рука (решение советника)"),
    ("standard", "«Заказы Standard» K Выручка", "revenue", "revenue", "exact", "—", "площадка = не `t`"),
    ("standard", "«Заказы Standard» P Соинвест", "coinvest_pct", "coinvest_pct", "tolerance", "—", ""),
    ("discounter", "«WB Дискаунтер» B Выручка (артикулы `t`)", "revenue", "revenue", "exact", "9 из 21", "сдвиги даты и разница воронки с supplier/orders по товару (WB-6 §2)"),
    ("discounter", "«WB Дискаунтер» C Маржа", "margin", "margin", "exact", "—", "его цены 1С"),
)


def check_orders(sheets, manuals, young_days=()):
    """sheets — {площадка: строки листа}; manuals — {площадка: {день: {поле: repr}}}. Возвращает таблицу строк приёмки."""
    out = []
    for platform, title, ours_k, theirs_k, kind, expectation, why in ORDERS_CHECK_LINES:
        rows, manual = sheets.get(platform, []), manuals.get(platform, {})
        t_o = t_t = Z
        equal, bad, missing = 0, [], []
        for r in rows:
            if r["date"] in young_days:
                continue
            m = manual.get(r["date"])
            o = r.get(ours_k)
            t = None if not m or m.get(theirs_k) is None else Decimal(m[theirs_k])
            if o is None or t is None:
                missing.append(r["date"]); continue
            if kind == "tolerance":
                o4, t4 = Decimal(o).quantize(Decimal("0.0001")), Decimal(t).quantize(Decimal("0.0001"))
                if abs(o4 - t4) <= COINVEST_TOLERANCE:
                    equal += 1
                else:
                    bad.append((r["date"], o4 - t4))
            else:
                o, t = q(o), q(t)
                t_o += o; t_t += t
                if o == t:
                    equal += 1
                else:
                    bad.append((r["date"], o - t))
        days = len([r for r in rows if r["date"] not in young_days])
        out.append({"title": title, "kind": kind, "theirs": t_t, "ours": t_o, "diff": t_o - t_t, "equal_days": equal, "days": days,
                    "bad": bad, "missing": missing, "expectation": expectation, "why": why})
    return out


def print_check_orders(table):
    print(f"\n{'лист / колонка':70}{'у него':>17}{'у нас':>17}{'разница':>15}{'дней 0,00':>11}   ожидание задачи")
    for t in table:
        money = t["kind"] == "exact"
        theirs = "{:,.2f}".format(t["theirs"]) if money else "—"
        ours = "{:,.2f}".format(t["ours"]) if money else "—"
        diff = "{:,.2f}".format(t["diff"]) if money else "—"
        print(f"{t['title'][:69]:70}{theirs:>17}{ours:>17}{diff:>15}{t['equal_days']:>7} из {t['days']}   {t['expectation']}")
        if t["bad"]:
            fmt = (lambda v: f"{v:+,.2f}") if money else (lambda v: f"{v:+.4f}")
            print("      расходится: " + ", ".join(f"{d} {fmt(v)}" for d, v in t["bad"][:12]) + (" …" if len(t["bad"]) > 12 else ""))
        if t["missing"]:
            print("      нет значения: " + ", ".join(t["missing"]))
        if t["why"] and (t["bad"] or t["missing"]):
            print(f"      чем объясняется: {t['why']}")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", required=True, help="YYYY-MM")
    ap.add_argument("--date-to", help="последний день листа; по умолчанию — вчера (строки дня появляются на D+1)")
    ap.add_argument("--out")
    ap.add_argument("--rows-from-files", metavar="DIR", help="строки отчёта из файлов сырья бэкфилла, а не из таблицы")
    ap.add_argument("--snapshot", default=SNAP)
    ap.add_argument("--cogs-by", choices=("base", "variant"), default="base", help="СС по базовому артикулу (решение WB-5) или по варианту артикул-размер")
    ap.add_argument("--ads", choices=("table", "api", "none"), default="table",
                    help="реклама: table — из wb_ad_spend_daily (по умолчанию); api — живой adv/v1/upd (ручной путь); none — пусто")
    ap.add_argument("--fetch", dest="ads", action="store_const", const="api", help="то же, что --ads api")
    ap.add_argument("--no-fetch", dest="ads", action="store_const", const="none", help="то же, что --ads none")
    ap.add_argument("--check", action="store_true"); ap.add_argument("--xlsx"); ap.add_argument("--sheet")
    ap.add_argument("--funnel-from-files", metavar="DIR", help="воронка по товарам из сырья бэкфилла (data/wb_funnel_raw), а не из таблицы")
    ap.add_argument("--check-orders", action="store_true", help="приёмка листов заказов по книге владельца (--xlsx)")
    ap.add_argument("--orders-sheet", default="Заказы"); ap.add_argument("--orders-standard-sheet", default="Заказы Standard")
    ap.add_argument("--orders-discounter-sheet", default="WB Дискаунтер")
    args = ap.parse_args(argv)

    today = date.today()
    d2 = args.date_to or (today - timedelta(days=1)).isoformat()
    days = month_days(args.month, d2)
    d1 = days[0]
    sb = loader._client()
    rows = load_rows_files(args.rows_from_files, d1, days[-1]) if args.rows_from_files else load_rows_db(sb, d1, days[-1])
    print(f"строк отчёта по rrDate за {d1} … {window_end(days[-1])} (запас {LAG_DAYS} дн. под дату продажи): {len(rows)} ("
          + ("файлы " + args.rows_from_files if args.rows_from_files else loader.TABLE) + ")")
    exact, uniform = load_costs(sb, args.snapshot)
    print(f"снимок 1С {args.snapshot}: точных строк {len(exact)}, базовых артикулов с единой СС {len(uniform)}; стыковка — {args.cogs_by}")
    ads_by_day, ads_known, ads_note = {}, False, "реклама не запрашивалась (--ads none)"
    if args.ads != "none":
        try:
            ads_by_day, undated = fetch_ads(d1, days[-1]) if args.ads == "api" else load_ads_db(sb, d1, days[-1])
            ads_known = True
            source = "adv/v1/upd (живой)" if args.ads == "api" else ads_loader.TABLE
            ads_note = f"реклама: {source} за {d1}…{days[-1]} — списаний {sum(ads_by_day.values(), Z):,.2f} с НДС" + (f", без даты {undated:,.2f}" if undated else "")
        except Exception as error:  # реклама не обязана останавливать лист; пусто — не ноль
            ads_note = f"реклама не получена: {error}"
    print(ads_note)
    cost_fn = lambda code, size: unit_cost_for(exact, uniform, code, size, args.cogs_by)  # noqa: E731
    outside = {}
    daily = build_daily(rows, days, cost_fn, ads_by_day, today, ads_known, outside)
    total = total_row(daily)
    young = [r["date"] for r in daily if r["young"]]
    outside_text = (f"продаж/возвратов с датой продажи раньше листа {outside['before'][0]} (оборот {outside['before'][1]:,.2f}), "
                    f"позже {outside['after'][0]} (оборот {outside['after'][1]:,.2f})")
    print(outside_text)
    if total["acquiring_other"]:
        print(f"ВНИМАНИЕ: acquiringFee вне строк продаж/возвратов — {total['acquiring_other']:,.2f} с НДС, в лист не вошёл")
    if total["no_pct_rows"]:
        print(f"ВНИМАНИЕ: строк продаж/возвратов без commissionPercent — {total['no_pct_rows']}, в комиссии по образцу они как 0 %")
    notes = [f"Источник: отчёт реализации WB (finance-api sales-reports/detailed, period=daily); день строки — saleDt в московском времени, без saleDt — rrDate. "
             f"Строк {len(rows)}; {outside_text}.",
             f"Комиссия = Σ цена × кВВ строки (commissionPercent), как у владельца; справочно «удержано из выплаты всего» = Оборот − forPay "
             f"(комиссия + согласованная скидка + доплаты). НДС {total['vat']}.",
             f"Себестоимость: снимок 1С {args.snapshot}, стыковка {args.cogs_by}; позиций по источникам: " + ", ".join(f"{k} {v}" for k, v in sorted(total["cost_sources"].items()))
             + f"; без СС {total['no_cost_positions']} из {total['positions']} позиций — оборот {total['no_cost_turnover']:,.2f}"
             + (f" ({ratio(total['no_cost_turnover'], total['turnover']):.1%} оборота)." if total["turnover"] else "."),
             ads_note, "Логистика = deliveryService / НДС по дате продажи МСК — с листом владельца до копейки; эквайринг — возврат со знаком минус; "
             "прочее = хранение / НДС + штрафы без НДС + удержания / НДС (docs/wb_report_model.md).",
             f"Жёлтым — дни моложе {YOUNG_DAYS} суток: строки доезжают." + (f" Таких дней: {', '.join(young)}." if young else "")]
    # Листы заказов — из воронки по товарам (таблица или сырьё); нет источника — листов нет, сказано вслух.
    funnel_rows, funnel_note = [], ""
    try:
        if args.funnel_from_files:
            funnel_rows = load_funnel_files(args.funnel_from_files, d1, days[-1])
            funnel_note = f"воронка по товарам: файлы {args.funnel_from_files}, строк {len(funnel_rows)}"
        else:
            funnel_rows = load_funnel_db(sb, d1, days[-1])
            funnel_note = f"воронка по товарам: {FUNNEL_TABLE}, строк {len(funnel_rows)}"
    except RuntimeError as error:
        funnel_note = f"воронка по товарам не прочитана: {error}; листы «Заказы WB» не собраны"
    print(funnel_note)
    orders, sheets = [], {}
    if funnel_rows:
        coinvest = coinvest_by_day(rows)
        for platform, title in PLATFORMS:
            o_rows = build_orders_daily(funnel_rows, days, cost_fn, ads_by_day, coinvest, today, ads_known, platform)
            o_total = orders_total_row(o_rows)
            sheets[platform] = o_rows
            o_notes = [f"Источник: воронка продаж WB по товарам (sales-funnel/products, день заказа МСК) — {funnel_note}. Площадка — первая буква артикула: "
                       f"`{DISCOUNTER_LETTER}` — Дискаунтер (КОЮЗ Топаз), остальное — Standard; лист «{title}».",
                       f"Выручка = Σ orderSum × {OWNER_ORDERS_AFTER_COMMISSION} / НДС (как K владельца); Себестоимость = Σ orderCount × СС снимка {args.snapshot} по базовому артикулу; "
                       f"без СС {o_total['no_cost_qty']} из {o_total['orders_qty']} заказов ({o_total['no_cost_sum']:,.0f} с НДС); позиций по источникам: "
                       + ", ".join(f"{k} {v}" for k, v in sorted(o_total["cost_sources"].items())) + ".",
                       "Соинвест, % — рабочее правило: (Σ retailPriceWithDisc − Σ retailAmount) / Σ retailPriceWithDisc по строкам «Продажа» отчёта реализации за день (saleDt МСК); точного правила владельца нет (WB-6 §2).",
                       "Реклама — только на общем листе (списания по кампаниям на площадки не делятся); ДРР — от суммы заказов с НДС.",
                       f"Жёлтым — дни моложе {YOUNG_DAYS} суток: воронка пересматривает день задним числом."]
            orders.append((title, o_rows, o_total, o_notes))
    out = args.out or os.path.join(OUT_DIR, f"wb_{args.month}_to_{days[-1]}.xlsx")
    write_xlsx(out, args.month, daily, total, notes, orders=orders)
    for title, _r, o_total, _n in orders:
        print(f"{title} {d1}…{days[-1]}: заказов {o_total['orders_qty']} на {o_total['orders_sum']:,.2f}, выручка {o_total['revenue']:,.2f}, "
              f"СС {o_total['cogs']:,.2f}, маржа {o_total['margin']:,.2f}" + (f", реклама {o_total['ads']:,.2f}" if o_total["ads"] is not None else "")
              + f"; без СС {o_total['no_cost_qty']} заказов")
    ads_text = "—" if total["ads"] is None else f"{total['ads']:,.2f}"
    print(f"итого {d1}…{days[-1]}: оборот {total['turnover']:,.2f}, комиссия {total['commission']:,.2f} ({total['commission_pct']:.2%}), выручка {total['revenue']:,.2f}, "
          f"СС {total['cogs']:,.2f}, маржа {total['margin']:,.2f}, логистика {total['logistics']:,.2f}, реклама {ads_text}, "
          f"эквайринг {total['acquiring']:,.2f}, прочее {total['other']:,.2f}, фин. рез. {total['fin_result']:,.2f}"
          + (f", накладные {total['overhead']:,.2f}, Ebitda {total['ebitda']:,.2f}" if total["ebitda"] is not None else ""))
    print(f"без СС: {total['no_cost_positions']} из {total['positions']} позиций, оборот {total['no_cost_turnover']:,.2f}"
          + (f" — {ratio(total['no_cost_turnover'], total['turnover']):.2%} оборота" if total["turnover"] else ""))
    print(f"справочно: удержано из выплаты всего (оборот − forPay) {total['withheld']:,.2f} ({total['withheld_pct']:.2%}); НДС за возмещение {total['vat_refund']:,.2f}")
    print(f"записано: {out}")
    code = 0
    if args.check:
        if not args.xlsx or not args.sheet:
            raise SystemExit("--check требует --xlsx и --sheet")
        manual = read_manual_sheet(args.xlsx, args.sheet)
        table, failures = check(daily, manual, set(young))
        print_check(table)
        print(f"приёмка: колонок, обязанных сходиться до копейки, — {sum(1 for l in CHECK_LINES if l[3])}; не сошлось {failures}; дни моложе двух суток исключены: {young or 'нет'}")
        code = 1 if failures else 0
    if args.check_orders:
        if not args.xlsx:
            raise SystemExit("--check-orders требует --xlsx")
        if not sheets:
            print("приёмка листов заказов: листы не собраны (нет воронки по товарам) — пропущена")
        else:
            manuals = {"all": read_manual_orders(args.xlsx, args.orders_sheet, 9, 10, ORDERS_MANUAL_JP),
                       "standard": read_manual_orders(args.xlsx, args.orders_standard_sheet, 9, 10, ORDERS_MANUAL_JP),
                       "discounter": read_manual_orders(args.xlsx, args.orders_discounter_sheet, 0, 1, ORDERS_MANUAL_BD)}
            young_orders = {r["date"] for r in sheets["all"] if r["young"]}
            print_check_orders(check_orders(sheets, manuals, young_orders))
            print(f"приёмка листов заказов: обязанных сходиться колонок нет (K — его снимки на пяти днях, L — его цены); "
                  f"дни моложе двух суток исключены: {sorted(young_orders) or 'нет'}")
    print("db_writes = 0")
    return code


if __name__ == "__main__":
    sys.exit(main())
