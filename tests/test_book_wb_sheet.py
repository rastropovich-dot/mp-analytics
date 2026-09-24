"""Лист «WB - месяц» в утренней книге Ozon: строки — функциями report_wb_month, запись — вторым листом; отказ не роняет доставку."""
import importlib.util
import os
import tempfile
import unittest
from decimal import Decimal
from unittest import mock

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def load(name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(ROOT, "scripts", f"{name}.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


book = load("book_wb_sheet")
send = load("send_ozon_month_report")


def wb_row(day, turnover, young=False):
    return {"date": day, "turnover": Decimal(turnover), "commission": Decimal("10"), "commission_pct": Decimal("0.1"), "withheld": Decimal("0"),
            "withheld_pct": Decimal("0"), "revenue": Decimal("50"), "cogs": Decimal("20"), "margin": Decimal("30"), "margin_pct": Decimal("0.6"),
            "logistics": Decimal("3"), "logistics_pct": Decimal("0.06"), "ads": None, "drr_pct": None, "acquiring": Decimal("1"), "acquiring_pct": Decimal("0.02"),
            "other": Decimal("1"), "fin_result": Decimal("25"), "fin_result_pct": Decimal("0.5"), "overhead": None, "ebitda": None, "ebitda_pct": None,
            "vat_refund": Decimal("0"), "reward": Decimal("0"), "rebill": Decimal("0"), "storage": Decimal("0"), "penalty": Decimal("0"), "deduction": Decimal("0"),
            "positions": 2, "no_cost_positions": 0, "rows": 3, "young": young}


class WriteSheet(unittest.TestCase):
    def test_sheet_lands_second_with_report_wb_month_columns(self):
        wb = openpyxl.Workbook(); wb.active.title = "Ozon - сентябрь"; wb.create_sheet("Заказы")
        daily = [wb_row("2026-09-01", "100"), wb_row("2026-09-02", "200", young=True)]
        total = dict(wb_row("Итого", "300")); total["date"] = "Итого"
        book.write_wb_sheet(wb, "2026-09", daily, total, ["заметка"])
        self.assertEqual(wb.sheetnames, ["Ozon - сентябрь", "WB - сентябрь", "Заказы"])
        ws = wb["WB - сентябрь"]
        heads = [c.value for c in ws[3] if c.value]
        self.assertEqual(heads[:3], ["Дата реализации", "Оборот (с НДС), руб.", "Комиссия (с НДС), руб."])
        self.assertEqual(ws.cell(row=4, column=2).value, 100.0)
        self.assertEqual(ws.cell(row=6, column=1).value, "Итого")
        self.assertEqual(ws.cell(row=5, column=2).fill.fgColor.rgb[-6:], "FFF2CC")       # молодой день — жёлтый, как у report_wb_month
        self.assertEqual(ws.cell(row=8, column=1).value, "заметка")

    def test_existing_wb_sheet_is_replaced_not_duplicated(self):
        wb = openpyxl.Workbook(); wb.active.title = "Ozon - сентябрь"; wb.create_sheet("WB - сентябрь")
        total = dict(wb_row("Итого", "1")); total["date"] = "Итого"
        book.write_wb_sheet(wb, "2026-09", [wb_row("2026-09-01", "1")], total, [])
        self.assertEqual(wb.sheetnames.count("WB - сентябрь"), 1)


class DeliveryHook(unittest.TestCase):
    def test_wb_line_goes_into_the_caption(self):
        caption = send.build_caption("2026-09", "2026-09-24", {"buyouts": {"turnover": 1e6}}, False, "WB: оборот 16,9 млн")
        self.assertIn("WB: оборот 16,9 млн", caption)
        self.assertLess(caption.index("Выкупы"), caption.index("WB:"))

    def test_failure_to_add_wb_sheet_is_a_caption_line_not_an_exception(self):
        fake = mock.Mock(); fake.add_wb_sheet.side_effect = RuntimeError("нет таблицы")
        with mock.patch.dict("sys.modules", {"book_wb_sheet": fake}), mock.patch("builtins.print"):
            line = send.add_wb_sheet_nonfatal("/tmp/x.xlsx", "2026-09", "2026-09-24")
        self.assertIn("не собран", line); self.assertIn("нет таблицы", line)

    def test_success_returns_wb_totals_line(self):
        fake = mock.Mock(); fake.add_wb_sheet.return_value = ({"turnover": 16877663.0, "revenue": 7.9e6, "fin_result": 1.9e6, "ebitda": 1e5,
                                                                "commission": 1, "logistics": 1}, {"rows": 15053, "days": 24, "young": ["2026-09-23", "2026-09-24"]})
        with mock.patch.dict("sys.modules", {"book_wb_sheet": fake}), mock.patch("builtins.print"):
            line = send.add_wb_sheet_nonfatal("/tmp/x.xlsx", "2026-09", "2026-09-24")
        self.assertEqual(line, "WB: оборот 16,9 млн, выручка 7,9 млн, фин. рез. 1,9 млн, Ebitda 0,1 млн")


if __name__ == "__main__":
    unittest.main()
