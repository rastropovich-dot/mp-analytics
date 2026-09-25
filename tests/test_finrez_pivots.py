"""Тридцать восьмая §4: пересадка сводных владельца — чистые преобразования XML; интеграция — только если оригиналы владельца лежат в data/ (вне git)."""
import os
import sys
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import finrez_pivots as fp  # noqa: E402


class Transforms(unittest.TestCase):
    def test_inline_strings_strip_styles_and_notes(self):
        xml = ('<worksheet><cols><col min="2" max="2" width="22" style="7"/></cols><sheetData><row r="3"><c r="A3" s="5" t="s"><v>1</v></c>'
               '<c r="B3" t="s"><v>0</v></c><c r="C3" s="2"><v>12.5</v></c></row></sheetData><legacyDrawing r:id="rId2"/><pageSetup paperSize="9" r:id="rId9"/></worksheet>')
        out = fp.inline_strings(xml, ["Статус", "Бренд & Ко"])
        self.assertIn('<c r="A3" t="inlineStr"><is><t xml:space="preserve">Бренд &amp; Ко</t></is></c>', out)
        self.assertIn('<c r="B3" t="inlineStr"><is><t xml:space="preserve">Статус</t></is></c>', out)
        self.assertIn('<c r="C3"><v>12.5</v></c>', out)
        self.assertNotIn('style="7"', out); self.assertNotIn("legacyDrawing", out); self.assertNotIn("pageSetup", out)

    def test_merge_styles_and_remap(self):
        ours = '<styleSheet><numFmts count="1"><numFmt numFmtId="164" formatCode="0.0%"/></numFmts><fonts count="1"><font/></fonts><cellXfs count="1"><xf/></cellXfs><tableStyles count="0"/></styleSheet>'
        owner = ('<styleSheet><numFmts count="2"><numFmt numFmtId="3" formatCode="#,##0"/><numFmt numFmtId="165" formatCode="0.0%"/></numFmts>'
                 '<dxfs count="2"><dxf><numFmt numFmtId="165" formatCode="0.0%"/></dxf><dxf><fill/></dxf></dxfs></styleSheet>')
        merged, offset, fmap = fp.merge_styles(ours, owner)
        self.assertEqual((offset, fmap), (0, {165: 165}))
        self.assertIn('<numFmts count="2"><numFmt numFmtId="164" formatCode="0.0%"/><numFmt numFmtId="165" formatCode="0.0%"/></numFmts>', merged)
        self.assertIn('<dxfs count="2"><dxf><numFmt numFmtId="165"', merged)
        merged2, offset2, fmap2 = fp.merge_styles(merged, owner)          # второй файл владельца: свои dxf — со сдвигом, numFmt 165 → 166
        self.assertEqual((offset2, fmap2), (2, {165: 166}))
        self.assertIn('<dxfs count="4">', merged2)
        pt = fp.remap_formats('<format dxfId="1"/><dataField numFmtId="165"/><dataField numFmtId="3"/>', offset2, fmap2)
        self.assertEqual(pt, '<format dxfId="3"/><dataField numFmtId="166"/><dataField numFmtId="3"/>')

    def test_col_letter(self):
        self.assertEqual([fp.col_letter(n) for n in (1, 15, 26, 27, 30)], ["A", "O", "Z", "AA", "AD"])


@unittest.skipUnless(all(os.path.exists(os.path.join(ROOT, s["owner"])) for s in fp.SPECS), "оригиналы владельца лежат вне git")
class Integration(unittest.TestCase):
    def test_transplant_into_a_small_book_and_check(self):
        import tempfile
        from decimal import Decimal as D
        import report_finrez as fr
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "book.xlsx")
            wb = openpyxl.Workbook(write_only=True)
            for title, cols, row in (("Данные Ozon выкупы", fr.LONG_COLS, [1, "2026-09-01", "F1", "11", 1, "Товарооборот", "KARATOV", "сен", 0, "KARATOV", "кольца", 100, "Кольцо", 1, 9, 0, "Ozon"]),
                                     ("Данные WB выкупы", [(h, k, None) for h, k, f in fr.wbfin.DATA_COLS], [None] * len(fr.wbfin.DATA_COLS)),
                                     ("Данные заказы", fr.ORDER_DATA_COLS, [None] * len(fr.ORDER_DATA_COLS))):
                ws = wb.create_sheet(title); ws.append([h for h, _k, _f in cols]); ws.append(row)
            wb.save(path)
            rep = fp.transplant(path, path, row_counts={"Данные Ozon выкупы": 2, "Данные WB выкупы": 2, "Данные заказы": 2})
            self.assertEqual([r["ref"] for r in rep], ["A1:O2", "A1:O2", "A1:AD2"])
            errors, info = fp.check(path)
            self.assertEqual(errors, [], errors)
            self.assertEqual(openpyxl.load_workbook(path, read_only=True).sheetnames[-3:], ["Сводная Ozon выкупы", "Сводная WB выкупы", "Сводная заказы"])


if __name__ == "__main__":
    unittest.main()
