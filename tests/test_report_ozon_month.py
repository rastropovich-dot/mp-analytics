"""Генератор листа «Ozon - <месяц>»: формулы ручного листа на наших источниках.

Формулы сняты с ячеек листа владельца (data/manual_report_september.xlsx, 2026-09-19):
выручка = (оборот − комиссия) / НДС, ДРР = реклама / (оборот / НДС),
фин. рез. = маржа − логистика − реклама − эквайринг − прочее (у нас ещё − подписка).
Живая приёмка против листа — режим --check самого скрипта; здесь — арифметика и правила.
"""
import importlib.util
import os
import tempfile
import unittest
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_spec = importlib.util.spec_from_file_location("report_ozon_month", os.path.join(ROOT, "scripts", "report_ozon_month.py"))
rep = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(rep)

D = Decimal
DAY = "2026-09-10"


def buyout(amount="1220", commission="0", qty=1, sku="11", day=DAY):
    return {"buyout_date": day, "marketplace_sku": sku, "buyouts_qty": qty, "buyouts_amount_seller": amount, "commission_amount": commission}


def expense(kind, amount, day=DAY):
    return {"expense_date": day, "expense_type": kind, "expense_amount": amount}


# расход положителен, как внутри листа: 32 — логистика; 1, 96, 38 — статья other; 51 — подписка; 25 — компенсация нам (в леджере +61)
TYPES = {1: D("12.2"), 41: D("61"), 54: D("61"), 51: D("24.4"), 96: D("12.2"), 38: D("12.2"), 32: D("122"), 25: D("-61")}
EXPENSES = [expense("logistics", "122"), expense("other", "36.6"), expense("subscription", "24.4"),
            expense("advertising_clicks", "999"), expense("commission", "777")]


def one_day(buyouts=None, expenses=None, types=TYPES, cost=lambda sku: D("300"), today="2026-09-19"):
    rows, unknown = rep.build_daily([DAY], buyouts or [buyout()], EXPENSES if expenses is None else expenses,
                                    {} if types is None else {DAY: types}, cost, today)
    return rep.add_ratios(rows[0]), unknown


class VatTests(unittest.TestCase):
    def test_rate_has_an_effective_date(self):
        self.assertEqual(rep.vat_for("2025-12-31"), D("1.20"))
        self.assertEqual(rep.vat_for("2026-01-01"), D("1.22"))
        self.assertEqual(rep.vat_for("2026-09-10"), D("1.22"))


class DailyFormulaTests(unittest.TestCase):
    def test_columns_follow_the_manual_sheet_formulas(self):
        r, unknown = one_day()
        self.assertEqual(unknown, {})
        self.assertEqual((r["turnover"], r["commission"], r["revenue"]), (D("1220"), D("0"), D("1000")))
        self.assertEqual((r["cogs"], r["margin"]), (D("300"), D("700")))
        self.assertEqual(r["logistics"], D("100"))                 # статья logistics / НДС
        self.assertEqual(r["acquiring"], D("10"))                  # тип 1 / НДС
        self.assertEqual(r["subscription"], D("20"))
        self.assertEqual(r["ads"], D("100"))                       # 41 + 54, без 51 и 96
        self.assertEqual(r["other"], D("20"))                      # (other − тип 1) / НДС
        self.assertEqual(r["fin_result"], D("700") - 100 - 10 - 20 - 100 - 20)
        self.assertEqual(r["compensations"], D("50"))              # типы 25 + 10: доход, «+» — деньги нам
        self.assertEqual(r["fin_result_with_comp"], r["fin_result"] + 50)
        self.assertEqual(r["drr_pct"], D("0.1"))                   # реклама / (оборот / НДС)
        self.assertEqual(r["margin_pct"], D("0.7"))

    def test_reference_columns_reproduce_the_manual_grouping(self):
        r, _ = one_day()
        self.assertEqual(r["ads_like_manual"], D("130"))           # (41 + 54 + 51 + 96) / НДС
        self.assertEqual(r["log_other_like_manual"], D("110"))     # (logistics + other − тип 1 − тип 96) / НДС
        # у владельца компенсация сидит внутри «Прочего» расходом с обратным знаком: их Л + П = наше − компенсации
        self.assertEqual(r["log_other_like_manual"] - r["compensations"], D("60"))

    def test_performance_ads_and_commission_expenses_stay_out_of_sheet_one(self):
        with_extra, _ = one_day()
        without, _ = one_day(expenses=[e for e in EXPENSES if e["expense_type"] not in ("advertising_clicks", "commission")])
        self.assertEqual(with_extra["fin_result"], without["fin_result"])

    def test_commission_comes_from_buyouts(self):
        r, _ = one_day(buyouts=[buyout("1220", "244")])
        self.assertEqual(r["revenue"], D("800"))
        self.assertEqual(r["commission_pct"], D("0.2"))

    def test_unknown_type_goes_to_other_and_is_named(self):
        r, unknown = one_day(types={**TYPES, 777: D("12.2")})
        self.assertEqual(unknown, {"тип 777": D("12.2")})
        self.assertEqual(r["other"], D("30"))

    def test_expense_columns_come_from_accrual_types_not_from_the_expenses_table(self):
        """В marketplace_expenses застревают строки, которые Ozon убрал: лист считает по типам, базу только сверяет."""
        clean, _ = one_day()
        stale, _ = one_day(expenses=EXPENSES + [expense("other", "841.77")])
        self.assertEqual((stale["other"], stale["fin_result"]), (clean["other"], clean["fin_result"]))
        self.assertEqual(stale["db_minus_raw"], {"other": D("841.77")})
        self.assertEqual(clean["db_minus_raw"], {})

    def test_no_types_means_empty_not_zero(self):
        r, _ = one_day(types=None)
        for k in ("acquiring", "ads", "compensations", "fin_result", "fin_result_with_comp", "drr_pct"):
            self.assertIsNone(r[k], k)
        self.assertEqual(r["other"], D("30"))                      # запасной путь — статья из расходов; эквайринг не отделить
        self.assertIsNone(rep.total_row([r])["fin_result"])        # итог неполон, и это видно

    def test_position_without_cost_is_counted_not_priced_at_zero(self):
        r, _ = one_day(buyouts=[buyout(qty=2, sku="no-cost")], cost=lambda sku: None)
        self.assertEqual((r["cogs"], r["no_cost_positions"]), (D("0"), 2))

    def test_young_dates_are_flagged(self):
        self.assertTrue(one_day(today="2026-09-11")[0]["young"])
        self.assertFalse(one_day(today="2026-09-12")[0]["young"])

    def test_ledger_types_are_flipped_to_expense_positive_once(self):
        class Q:
            def __init__(self, rows): self.rows = rows
            def select(self, *_a): return self
            def gte(self, *_a): return self
            def lte(self, *_a): return self
            def order(self, *_a): return self
            def range(self, *_a): return self
            def execute(self): return type("R", (), {"data": self.rows})()
        sb = type("SB", (), {"table": lambda self, name: Q([{"accrual_date": DAY, "type_id": 41, "amount": -61.0}, {"accrual_date": DAY, "type_id": 25, "amount": 61.0}])})()
        self.assertEqual(rep.load_types_from_ledger(sb, DAY, DAY), {DAY: {41: D("61.0"), 25: D("-61.0")}})

    def test_two_sources_are_compared_type_by_type(self):
        self.assertEqual(rep.types_differ({1: D("10"), 38: D("0")}, {1: D("10.004")}), {})          # ноль = отсутствию, копейки — по округлению
        self.assertEqual(rep.types_differ({1: D("10")}, {1: D("10.01"), 96: D("5")}), {1: (D("10"), D("10.01")), 96: (D("0"), D("5"))})


class TotalsTests(unittest.TestCase):
    def test_total_percentages_are_computed_from_sums(self):
        days = ["2026-09-10", "2026-09-11"]
        rows, _ = rep.build_daily(days, [buyout("1220", day=days[0]), buyout("2440", day=days[1])], [],
                                  {d: {41: D("122")} for d in days}, lambda sku: D("0"), "2026-09-19")
        total = rep.total_row([rep.add_ratios(r) for r in rows])
        self.assertEqual((total["turnover"], total["ads"]), (D("3660"), D("200")))
        self.assertEqual(total["drr_pct"], D("200") / D("3000"))

    def test_month_days_respects_date_to(self):
        self.assertEqual(len(rep.month_days("2026-09")), 30)
        self.assertEqual(rep.month_days("2026-09", "2026-09-15")[-1], "2026-09-15")


class CheckTests(unittest.TestCase):
    def setUp(self):
        self.row, _ = one_day()
        self.manual = {DAY: {"turnover": D("1220"), "commission": D("0"), "revenue": D("1000"), "acquiring": D("10"), "ads": D("130"),
                             "logistics": D("95"), "other": D("-35"), "cogs": D("350"), "fin_result": D("450")}}

    def test_exact_columns_pass_and_known_boundaries_do_not_fail_the_check(self):
        table, failures = rep.check([self.row], self.manual)
        self.assertEqual(failures, 0)
        by = {t["title"]: t for t in table}
        self.assertEqual(by["Л + П по образцу − Компенсации = их Логистика + Прочее"]["diff"], D("0.00"))   # 110 − 50 против 95 − 35
        self.assertEqual(by["   то же без компенсаций: Л + П по образцу = их Л + П"]["diff"], D("50.00"))
        self.assertEqual(by["Себестоимость (наш снимок 1С против их цен)"]["diff"], D("-50.00"))
        self.assertEqual(by["Логистика (наша статья против их строки)"]["diff"], D("5.00"))               # границу не повторяем — не провал

    def test_a_kopeck_off_in_an_exact_column_fails(self):
        self.manual[DAY]["turnover"] = D("1220.01")
        _table, failures = rep.check([self.row], self.manual)
        self.assertEqual(failures, 1)

    def test_a_day_missing_from_the_manual_sheet_fails(self):
        _table, failures = rep.check([self.row], {})
        self.assertEqual(failures, len(rep.EXACT))


class WorkbookTests(unittest.TestCase):
    def test_workbook_has_both_sheets_total_row_and_notes(self):
        import openpyxl
        row, _ = one_day(today="2026-09-11")
        total = rep.total_row([row])
        sku_rows = rep.build_sku([{"marketplace_sku": "11", "article": "F11", "product_name": "товар", "buyouts_qty": 1, "buyouts_amount_seller": "1220",
                                   "commission_amount": "0", "ad_spend": "122", "logistics_amount": "122", "other_expenses_amount": "61"}],
                                 {"11": "F11"}, lambda sku: D("300"), D("1.22"))
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "reports", "ozon_2026-09.xlsx")
            rep.write_xlsx(path, "2026-09", [row], total, sku_rows, ["источники"], ["про рекламу"])
            wb = openpyxl.load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Ozon - сентябрь", "По SKU"])
            ws = wb["Ozon - сентябрь"]
            self.assertEqual(ws["A5"].value, "Итого")
            self.assertEqual(ws["B4"].value, 1220)
            heads = {c.value: c.column for c in ws[3] if c.value}
            self.assertIn("моложе двух суток", ws.cell(row=4, column=heads["Примечание"]).value)
            self.assertEqual(ws.cell(row=4, column=heads["Компенсации Ozon (доход), руб."]).value, 50)
            self.assertEqual(ws.cell(row=4, column=heads["справочно: Фин. рез. с компенсациями"]).value, 500)
            sku = wb["По SKU"]
            self.assertEqual((sku["A2"].value, sku["B2"].value, sku["K2"].value), ("11", "F11", 100))
            self.assertEqual(sku["A3"].value, "Итого")


if __name__ == "__main__":
    unittest.main()
