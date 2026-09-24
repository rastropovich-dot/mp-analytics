"""Лист «WB - месяц»: формулы владельца на строках отчёта реализации, СС по базовому артикулу, приёмка.

Держим: день строки — saleDt в московском времени (без saleDt — rrDate); оборот = продажа − возврат
по retailPriceWithDisc; комиссия = Σ цена × кВВ строки со знаком возврата (как у владельца, решение
советника 2026-09-23); справочно «удержано из выплаты всего» = оборот − forPay; НДС за возмещение = (ppvzReward + rebillLogisticCost) × 0,22/1,22; эквайринг —
возврат со знаком минус, / 1,22; прочее = хранение / 1,22 + штрафы без НДС + удержания / 1,22;
Ebitda = фин. рез. − 74 002 с 09-01; СС — по базовому артикулу через unit_cost_uniform, безразмерные —
точной строкой; строки с датой продажи вне листа считаются, не теряются молча; известные дни в
приёмке печатаются и в отказ не считаются.
"""
import os
import tempfile
import unittest
from datetime import date
from decimal import Decimal

import scripts.report_wb_month as rep

D = Decimal


def row(day, oper="Продажа", code="f000283615", size="17,5", **extra):
    base = {"rrd_id": extra.pop("rrd_id", 1), "rr_date": day, "sale_dt": f"{day}T10:00:00Z", "seller_oper_name": oper,
            "doc_type": "Продажа" if oper == "Продажа" else "", "vendor_code": code, "tech_size": size, "nm_id": 1, "quantity": 1,
            "retail_price_with_disc": None, "retail_amount": None, "for_pay": None, "commission_percent": None, "ppvz_reward": None,
            "rebill_logistic_cost": None, "delivery_service": None, "acquiring_fee": None, "paid_storage": None, "penalty": None,
            "deduction": None, "cashback_discount": None}
    base.update(extra)
    return base


ROWS = [
    row("2026-09-01", retail_price_with_disc="1000", for_pay="556", commission_percent="42", ppvz_reward="18115.294", rebill_logistic_cost="8463.49", acquiring_fee="122", rrd_id=1),
    row("2026-09-01", oper="Возврат", retail_price_with_disc="100", for_pay="55.6", commission_percent="43", acquiring_fee="12.2", rrd_id=2),
    row("2026-09-01", oper="Доставка", delivery_service="244", rrd_id=3),
    row("2026-09-01", oper="Хранение", paid_storage="1220", rrd_id=4),
    row("2026-09-01", oper="Штраф", penalty="244", rrd_id=5),
    row("2026-09-01", oper="Удержание", deduction="122", rrd_id=6),
]
EXACT = {"f000283615-17,5": D("100"), "f000283615-16": D("120"), "f000009483": D("50")}
UNIFORM = {"f000283615": D("110"), "f000009483": D("50")}


def cost_base(code, size):
    return rep.unit_cost_for(EXACT, UNIFORM, code, size, "base")


class FormulaTests(unittest.TestCase):
    def test_owner_formulas_on_one_day(self):
        [r] = rep.build_daily(ROWS, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        self.assertEqual(r["turnover"], D("900"))                       # 1000 − 100
        self.assertEqual(r["commission"], D("377.00"))                  # 1000 × 42 % − 100 × 43 % — кВВ строки, знак возврата
        self.assertEqual(r["withheld"], D("900") - D("500.4"))          # справочно: оборот − forPay± (556 − 55.6)
        self.assertEqual(rep.q(r["vat_refund"]), rep.q((D("18115.294") + D("8463.49")) * 22 / 122))  # как колонка AB владельца
        self.assertEqual(rep.q(r["revenue"]), rep.q((D("900") - D("377")) / D("1.22") - r["vat_refund"]))   # его E
        self.assertEqual(r["logistics"], D("200"))                      # 244 / 1,22
        self.assertEqual(r["acquiring"], D("90"))                       # (122 − 12,2) / 1,22 — возврат со знаком минус
        self.assertEqual(r["other"], D("1344"))                         # 1220 / 1,22 + 244 (штраф без НДС) + 122 / 1,22
        self.assertEqual(r["deduction"], D("122"))
        self.assertEqual(r["cogs"], D("110") - D("110"))                # продажа + возврат одного артикула по единой СС
        self.assertEqual(r["overhead"], D("74002.00"))
        self.assertEqual(r["ebitda"], r["fin_result"] - D("74002.00"))
        self.assertEqual(r["positions"], 0)
        self.assertEqual(r["no_pct_rows"], 0)
        self.assertFalse(r["young"])

    def test_young_days_are_the_last_two(self):
        rows = rep.build_daily([], ["2026-09-08", "2026-09-09", "2026-09-10"], cost_base, {}, date(2026, 9, 10))
        self.assertEqual([r["young"] for r in rows], [False, True, True])

    def test_overhead_is_empty_before_its_start(self):
        [r] = rep.build_daily([], ["2026-08-31"], cost_base, {}, date(2026, 9, 10))
        self.assertIsNone(r["overhead"]); self.assertIsNone(r["ebitda"])

    def test_ads_unknown_stays_empty_not_zero(self):
        [r] = rep.build_daily([], ["2026-09-01"], cost_base, {}, date(2026, 9, 10), ads_known=False)
        self.assertIsNone(r["ads"]); self.assertIsNone(r["drr_pct"])

    def test_ads_are_taken_without_vat_like_the_owner(self):
        [r] = rep.build_daily([], ["2026-09-01"], cost_base, {"2026-09-01": D("122")}, date(2026, 9, 10), ads_known=True)
        self.assertEqual(r["ads"], D("100"))
        [r] = rep.build_daily([], ["2026-09-02"], cost_base, {"2026-09-01": D("122")}, date(2026, 9, 10), ads_known=True)
        self.assertEqual(r["ads"], D(0))                          # известно и ноль — ноль, не пусто

    def test_a_sale_without_commission_percent_is_counted_not_hidden(self):
        rows = [row("2026-09-01", retail_price_with_disc="1000", for_pay="556")]
        [r] = rep.build_daily(rows, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        self.assertEqual((r["no_pct_rows"], r["commission"]), (1, D(0)))


class DayRuleTests(unittest.TestCase):
    def test_sale_day_is_sale_dt_in_moscow_time(self):
        self.assertEqual(rep.row_day({"rr_date": "2026-09-10", "sale_dt": "2026-09-09T16:50:53Z"}), "2026-09-09")   # четыре продажи 09-09 с rrDate 09-10
        self.assertEqual(rep.row_day({"rr_date": "2026-02-01", "sale_dt": "2026-01-31T23:08:23Z"}), "2026-02-01")   # 02:08 МСК
        self.assertEqual(rep.row_day({"rr_date": "2026-09-10", "sale_dt": "2026-09-09T16:50:53+00:00"}), "2026-09-09")  # формат timestamptz из таблицы

    def test_rows_without_sale_dt_fall_back_to_rr_date(self):
        self.assertEqual(rep.row_day({"rr_date": "2026-09-02", "sale_dt": None}), "2026-09-02")
        self.assertEqual(rep.row_day({"rr_date": "2026-09-02", "sale_dt": "0001-01-01T00:00:00Z"}), "2026-09-02")
        self.assertEqual(rep.row_day({"rr_date": "2026-09-02", "sale_dt": "не дата"}), "2026-09-02")

    def test_sales_dated_outside_the_sheet_are_counted_not_lost(self):
        rows = [row("2026-09-02", sale_dt="2026-08-31T10:00:00Z", retail_price_with_disc="500", for_pay="300", commission_percent="42", rrd_id=1),
                row("2026-09-30", sale_dt="2026-10-01T10:00:00Z", oper="Возврат", retail_price_with_disc="70", for_pay="40", commission_percent="42", rrd_id=2),
                row("2026-09-02", oper="Хранение", sale_dt="2026-08-31T10:00:00Z", paid_storage="122", rrd_id=3)]
        outside = {}
        daily = rep.build_daily(rows, ["2026-09-01", "2026-09-02"], cost_base, {}, date(2026, 9, 10), outside=outside)
        self.assertEqual(outside, {"before": [1, D("500")], "after": [1, D("-70")]})
        self.assertEqual(sum(r["turnover"] for r in daily), D(0))
        self.assertEqual(sum(r["storage"] for r in daily), D(0))       # хранение с датой вне листа тоже не на листе, но в счётчик продаж не входит

    def test_window_end_adds_the_lag(self):
        self.assertEqual(rep.window_end("2026-09-30"), "2026-10-07")


class CostTests(unittest.TestCase):
    def test_base_mode_uses_uniform_cost_and_plain_rows(self):
        self.assertEqual(rep.unit_cost_for(EXACT, UNIFORM, "F000283615", "17,5", "base"), (D("110"), "uniform"))
        self.assertEqual(rep.unit_cost_for(EXACT, UNIFORM, "f000009483", "0", "base"), (D("50"), "plain"))
        self.assertEqual(rep.unit_cost_for(EXACT, UNIFORM, "t000000001", "17", "base"), (None, "none"))

    def test_variant_mode_prefers_the_exact_size(self):
        self.assertEqual(rep.unit_cost_for(EXACT, UNIFORM, "f000283615", "17,5", "variant"), (D("100"), "variant"))
        self.assertEqual(rep.unit_cost_for(EXACT, UNIFORM, "f000283615", "19", "variant"), (D("110"), "uniform"))

    def test_missing_cost_is_counted_not_zeroed(self):
        rows = [row("2026-09-01", code="t000000001", retail_price_with_disc="10", for_pay="5", commission_percent="42")]
        [r] = rep.build_daily(rows, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        self.assertEqual((r["no_cost_positions"], r["cogs"], r["no_cost_turnover"]), (1, D(0), D("10")))


ACQ_TITLE = "Эквайринг (acquiringFee, возврат со знаком минус) / НДС"
LOG_TITLE = "Логистика (deliveryService / НДС, по дате продажи МСК)"


class CheckTests(unittest.TestCase):
    def manual_like_ours(self, daily):
        r = daily[0]
        return {"2026-09-01": {"turnover": "900", "revenue": str(rep.q(r["revenue"])), "acquiring": "90", "other": "1344",
                               "commission": "377", "logistics": "200", "cogs": "0", "fin_result": "0"}}

    def test_six_mandatory_columns_match_and_known_days_are_printed_not_failures(self):
        daily = rep.build_daily(ROWS, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        manual = self.manual_like_ours(daily)
        table, failures = rep.check(daily, manual)
        self.assertEqual(failures, 0)
        by = {t["title"]: t for t in table}
        self.assertEqual([t["title"] for t in table if t["must"]],
                         ["Оборот (с НДС)", "Выручка (оборот − комиссия) / НДС − НДС за возмещение", ACQ_TITLE,
                          "Прочее (хранение / НДС + штрафы без НДС + удержания / НДС)", "Комиссия (Σ цена × кВВ строки, знак возврата)", LOG_TITLE])
        self.assertTrue(all(by[t]["equal_days"] == 1 for t in by if by[t]["must"]))
        rep.KNOWN_MANUAL_DIFFS.add(("acquiring", "2026-09-01"))
        try:
            manual["2026-09-01"]["acquiring"] = "80"
            table, failures = rep.check(daily, manual)
            self.assertEqual(failures, 0)
            self.assertEqual({t["title"]: t for t in table}[ACQ_TITLE]["known"], [("2026-09-01", D("10.00"))])
        finally:
            rep.KNOWN_MANUAL_DIFFS.discard(("acquiring", "2026-09-01"))

    def test_a_mandatory_mismatch_fails(self):
        daily = rep.build_daily(ROWS, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        manual = {"2026-09-01": {"turnover": "901", "revenue": "1", "acquiring": "90", "other": "1344", "commission": "377", "logistics": "200"}}
        _table, failures = rep.check(daily, manual)
        self.assertGreaterEqual(failures, 1)

    def test_logistics_is_mandatory_now(self):
        daily = rep.build_daily(ROWS, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        manual = self.manual_like_ours(daily)
        manual["2026-09-01"]["logistics"] = "150"
        table, failures = rep.check(daily, manual)
        self.assertEqual(failures, 1)
        self.assertEqual({t["title"]: t for t in table}[LOG_TITLE]["bad"], [("2026-09-01", D("50.00"))])

    def test_young_days_are_excluded_from_the_check(self):
        daily = rep.build_daily(ROWS, ["2026-09-01"], cost_base, {}, date(2026, 9, 1))
        _table, failures = rep.check(daily, {}, young_days={"2026-09-01"})
        self.assertEqual(failures, 0)


class BookTests(unittest.TestCase):
    def test_workbook_is_written_with_the_owner_columns(self):
        import openpyxl
        daily = rep.build_daily(ROWS, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        total = rep.total_row(daily)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "wb.xlsx")
            rep.write_xlsx(path, "2026-09", daily, total, ["note"])
            ws = openpyxl.load_workbook(path)["WB - сентябрь"]
            heads = [ws.cell(row=3, column=j).value for j in range(1, 6)]
            self.assertEqual(heads[:5], ["Дата реализации", "Оборот (с НДС), руб.", "Комиссия (с НДС), руб.", "Комиссия, %",
                                         "справочно: удержано из выплаты всего (оборот − forPay), руб."])
            self.assertEqual(ws.cell(row=5, column=1).value, "Итого")
            self.assertAlmostEqual(ws.cell(row=4, column=2).value, 900.0)


if __name__ == "__main__":
    unittest.main()


class OrdersSheetTests(unittest.TestCase):
    """Листы «Заказы WB» (WB-7 §4): K = Σ orderSum × 0,58 / НДС, L = K − Σ orderCount × СС, площадка по букве `t`,
    соинвест из отчёта реализации по продажам дня, реклама только на общем листе, приёмка по блоку J–P владельца."""

    FUNNEL = [
        {"day": "2026-09-01", "nm_id": 1, "vendor_code": "F000283615", "order_count": 2, "order_sum": 200000, "buyout_count": 1, "buyout_sum": 90000, "cancel_count": 0, "cancel_sum": 0},
        {"day": "2026-09-01", "nm_id": 2, "vendor_code": "t000000001", "order_count": 1, "order_sum": 50000, "buyout_count": 0, "buyout_sum": 0, "cancel_count": 1, "cancel_sum": 50000},
        {"day": "2026-09-01", "nm_id": 3, "vendor_code": "F000009483", "order_count": 0, "order_sum": 0, "buyout_count": 0, "buyout_sum": 0, "cancel_count": 0, "cancel_sum": 0},
        {"day": "2026-09-02", "nm_id": 4, "vendor_code": "X1", "order_count": 1, "order_sum": 1000, "buyout_count": 0, "buyout_sum": 0, "cancel_count": 0, "cancel_sum": 0},
    ]
    REPORT = [
        row("2026-09-01", retail_price_with_disc="100000", retail_amount="57000", code="f000283615", rrd_id=1, sale_dt="2026-09-01T10:00:00Z"),
        row("2026-09-01", retail_price_with_disc="10000", retail_amount="6000", code="t000000001", rrd_id=2, sale_dt="2026-09-01T10:00:00Z"),
        row("2026-09-01", oper="Возврат", retail_price_with_disc="5000", retail_amount="3000", code="f000283615", rrd_id=3, sale_dt="2026-09-01T10:00:00Z"),
    ]

    def test_platform_is_the_first_letter_t(self):
        self.assertEqual((rep.platform_of("t000000001"), rep.platform_of("T0001"), rep.platform_of("F000283615"), rep.platform_of(None)), ("discounter", "discounter", "standard", "standard"))

    def test_owner_k_l_and_coinvest_on_one_day(self):
        coinvest = rep.coinvest_by_day(self.REPORT)
        self.assertEqual(coinvest["2026-09-01"]["all"], [D("110000"), D("63000")])          # возврат в соинвест не входит
        self.assertEqual(coinvest["2026-09-01"]["discounter"], [D("10000"), D("6000")])
        [r1, r2] = rep.build_orders_daily(self.FUNNEL, ["2026-09-01", "2026-09-02"], cost_base, {"2026-09-01": D("1220")}, coinvest, date(2026, 9, 10), True, "all")
        self.assertEqual((r1["orders_qty"], r1["orders_sum"], r1["cards"]), (3, D("250000"), 3))
        self.assertEqual(rep.q(r1["revenue"]), rep.q(D("250000") * D("0.58") / D("1.22")))   # 118 852,46 — K владельца
        self.assertEqual(r1["cogs"], D("110") * 2)                                            # СС только у F000283615 (uniform 110); t… нет в снимке
        self.assertEqual((r1["no_cost_qty"], r1["no_cost_sum"]), (1, D("50000")))
        self.assertEqual(r1["margin"], r1["revenue"] - D("220"))
        self.assertEqual(r1["ads"], D("1000"))                                                 # updSum с НДС / 1,22
        self.assertEqual(r1["coinvest_pct"], D("0.4273"))                                      # (110 000 − 63 000) / 110 000
        self.assertEqual((r1["funnel_buyouts_sum"], r1["funnel_cancel_qty"]), (D("90000"), 1))
        self.assertEqual((r2["orders_qty"], r2["coinvest_pct"], r2["ads"]), (1, None, rep.Z))

    def test_platform_sheets_split_the_total_and_carry_no_ads(self):
        coinvest = rep.coinvest_by_day(self.REPORT)
        [s] = rep.build_orders_daily(self.FUNNEL, ["2026-09-01"], cost_base, {"2026-09-01": D("1220")}, coinvest, date(2026, 9, 10), True, "standard")
        [d] = rep.build_orders_daily(self.FUNNEL, ["2026-09-01"], cost_base, {"2026-09-01": D("1220")}, coinvest, date(2026, 9, 10), True, "discounter")
        [a] = rep.build_orders_daily(self.FUNNEL, ["2026-09-01"], cost_base, {"2026-09-01": D("1220")}, coinvest, date(2026, 9, 10), True, "all")
        self.assertEqual(s["revenue"] + d["revenue"], a["revenue"])                            # 'Заказы Standard'!K + 'WB Дискаунтер'!B = K
        self.assertEqual((s["orders_sum"], d["orders_sum"]), (D("200000"), D("50000")))
        self.assertIsNone(s["ads"]); self.assertIsNone(d["ads"])
        self.assertEqual(d["coinvest_pct"], D("0.4000"))
        total = rep.orders_total_row([a])
        self.assertEqual((total["orders_qty"], total["revenue"], total["ads"]), (3, a["revenue"], D("1000")))

    def test_funnel_files_are_read_like_the_table(self):
        import json
        with tempfile.TemporaryDirectory() as tmp:
            payload = {"day": "2026-09-01", "products": [
                {"product": {"nmId": 7, "vendorCode": "t1"}, "statistic": {"selected": {"orderCount": 2, "orderSum": 300, "buyoutCount": 1, "buyoutSum": 100, "cancelCount": 0, "cancelSum": 0}}},
                {"product": {"vendorCode": "no-nm"}, "statistic": {"selected": {"orderCount": 1, "orderSum": 1}}}]}
            json.dump(payload, open(os.path.join(tmp, "funnel_2026-09-01.json"), "w"))
            json.dump({"day": "2026-09-05", "products": []}, open(os.path.join(tmp, "funnel_2026-09-05.json"), "w"))
            rows = rep.load_funnel_files(tmp, "2026-09-01", "2026-09-02")
        self.assertEqual(rows, [{"day": "2026-09-01", "nm_id": 7, "vendor_code": "t1", "order_count": 2, "order_sum": 300, "buyout_count": 1,
                                 "buyout_sum": 100, "cancel_count": 0, "cancel_sum": 0}])

    def test_manual_block_is_read_and_checked_exact_and_within_tolerance(self):
        import datetime
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "owner.xlsx")
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Заказы"
            ws.append(["Статистика"]); ws.append(["День", "Выручка"])
            ws.cell(row=3, column=10, value=datetime.datetime(2026, 9, 1)); ws.cell(row=3, column=11, value=118852.4590); ws.cell(row=3, column=12, value=100000.0)
            ws.cell(row=3, column=14, value=0); ws.cell(row=3, column=16, value=0.4298)
            ws.cell(row=4, column=10, value=datetime.datetime(2026, 9, 2)); ws.cell(row=4, column=11, value=999.0); ws.cell(row=4, column=16, value=0.5)
            wb.save(path)
            manual = rep.read_manual_orders(path, "Заказы", 9, 10, rep.ORDERS_MANUAL_JP)
        self.assertEqual(manual["2026-09-01"]["revenue"], "118852.459")
        self.assertIsNone(manual["2026-09-01"]["drr_pct"])
        sheets = {"all": [
            {"date": "2026-09-01", "revenue": D("118852.4590163934"), "margin": D("5"), "ads": rep.Z, "coinvest_pct": D("0.4273"), "young": False},
            {"date": "2026-09-02", "revenue": D("1000"), "margin": D("5"), "ads": rep.Z, "coinvest_pct": D("0.4900"), "young": False}]}
        table = rep.check_orders(sheets, {"all": manual})
        by = {t["title"]: t for t in table}
        k = by["«Заказы» K Выручка = воронка × 0,58 / НДС"]
        self.assertEqual((k["equal_days"], k["days"], k["bad"]), (1, 2, [("2026-09-02", D("1.00"))]))
        p = by["«Заказы» P Соинвест = (Σ цена − Σ оплачено) / Σ цена по продажам дня"]
        self.assertEqual((p["equal_days"], p["bad"]), (1, [("2026-09-02", D("-0.0100"))]))   # 09-01: |0,4273 − 0,4298| ≤ 0,005
        n = by["«Заказы» N Реклама"]
        self.assertEqual((n["equal_days"], n["missing"]), (1, ["2026-09-02"]))

    def test_workbook_gets_the_month_sheet_and_three_orders_sheets(self):
        import openpyxl
        coinvest = rep.coinvest_by_day(self.REPORT)
        month_rows = rep.build_daily(ROWS, ["2026-09-01"], cost_base, {}, date(2026, 9, 10))
        orders = []
        for platform, title in rep.PLATFORMS:
            o_rows = rep.build_orders_daily(self.FUNNEL, ["2026-09-01"], cost_base, {}, coinvest, date(2026, 9, 10), True, platform)
            orders.append((title, o_rows, rep.orders_total_row(o_rows), ["подвал"]))
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "b.xlsx")
            rep.write_xlsx(path, "2026-09", month_rows, rep.total_row(month_rows), ["n"], orders=orders)
            wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ["WB - сентябрь", "Заказы WB", "Заказы WB Standard", "Заказы WB Дискаунтер"])
        ws = wb["Заказы WB"]
        self.assertEqual(ws.cell(row=3, column=4).value, "Выручка (× 0,58 / НДС), руб.")
        self.assertAlmostEqual(ws.cell(row=4, column=4).value, float(D("250000") * D("0.58") / D("1.22")), places=2)
        self.assertEqual(ws.cell(row=5, column=1).value, "Итого")
