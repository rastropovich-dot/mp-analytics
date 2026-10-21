# -*- coding: utf-8 -*-
"""WB-18: лист «Хранение WB по артикулам» книги «Фин рез» из wb_paid_storage_daily.

Держим: строка = (месяц, nmId) с ярлыками из «Данных WB выкупы» (артикул / бренд / категория владельца), штук-дни без сторно,
сторно — со знаком и отдельной колонкой, без НДС по ставке дня; продажи — «Данные» того же месяца, без продаж — пусто, не 0;
свод по месяцам: Σ таблицы = «Хранение» формы за те же дни (разница 0,00), «за весь месяц» — вся форма; блоки × бренд / × категория
складываются в общий итог; читатель таблицы — индексными запросами по ключу и не теряет строки одного nmId на границе страницы;
лист в книге пишется после «Данных WB выкупы», ошибка модуля — лист с одной строкой, книга цела.
"""
import io
import os
import tempfile
import unittest
from contextlib import redirect_stdout
from decimal import Decimal

import scripts.report_finrez_wb as fr
from tests.test_report_finrez_wb import FakeSb

D = Decimal


def srow(day, nm, price, cnt=1, calc="короба: товары свыше базы", chrt=1, barcode="b", vendor="f000283615", subject="Серьги", brand="KARATOV"):
    return {"day": day, "nm_id": nm, "chrt_id": chrt, "calc_type": calc, "barcode": barcode, "warehouse_price": price, "barcodes_count": cnt,
            "vendor_code": vendor, "subject": subject, "brand": brand}


def drow(day, nm, sales="0", storage="0", article="F000283615", brand="KARATOV", category="Серьги"):
    return {"date": day, "nm_id": nm, "sales": D(sales), "storage": D(storage), "article": article, "brand": brand, "category": category, "title": "t"}


class BuildStorageTests(unittest.TestCase):
    def data(self):
        storage = [srow("2026-10-07", 1, "10.00", 2), srow("2026-10-07", 1, "2.00", 1, chrt=2), srow("2026-10-08", 1, "-1.50", 1, calc=fr.STORAGE_STORNO),
                   srow("2026-10-08", 2, "5.00", 3, vendor="s000000001", subject="Кольца", brand="SOKOLOV"),
                   srow("2026-09-30", 2, "4.00", 1, vendor="s000000001", subject="Кольца", brand="SOKOLOV"),
                   srow("2026-10-09", 3, "7.00", 1, vendor="x000000009", subject="Упаковки для украшений", brand="X")]
        data = [drow("2026-10-07", 1, sales="1000", storage="0"), drow("2026-10-07", None, storage="12.00", article=fr.NO_PRODUCT),
                drow("2026-10-08", None, storage="3.50", article=fr.NO_PRODUCT), drow("2026-10-09", None, storage="7.00", article=fr.NO_PRODUCT),
                drow("2026-09-30", None, storage="4.00", article=fr.NO_PRODUCT), drow("2026-09-15", None, storage="100.00", article=fr.NO_PRODUCT),
                drow("2026-10-08", 2, sales="-200", article="S000000001", brand="SOKOLOV", category="Кольца")]
        return storage, data

    def test_article_rows_labels_money_and_percent(self):
        st = fr.build_storage(*self.data())
        by = {(r["ym"], r["nm_id"]): r for r in st["articles"]}
        r1 = by[("2026-10", 1)]
        self.assertEqual((r1["days"], r1["unit_days"]), (2, 3))                       # штук-дни — без строки сторно
        self.assertEqual((r1["storage"], r1["storno"]), (D("10.50"), D("-1.50")))
        self.assertEqual(fr.q(r1["storage_net"]), fr.q(D("12.00") / fr.wbm.vat_for("2026-10-07") + D("-1.50") / fr.wbm.vat_for("2026-10-08")))
        self.assertEqual((r1["article"], r1["brand"], r1["category"], r1["subject"]), ("F000283615", "KARATOV", "Серьги", "Серьги"))
        self.assertEqual((r1["sales"], r1["storage_pct"]), (D("1000"), D("0.0105")))
        r2 = by[("2026-10", 2)]
        self.assertEqual((r2["sales"], r2["storage_pct"]), (D("-200"), None))          # возвраты больше продаж — пусто, не 0
        self.assertIsNone(by[("2026-09", 2)]["sales"]); self.assertIsNone(by[("2026-09", 2)]["storage_pct"])   # продаж нет — пусто
        r3 = by[("2026-10", 3)]                                                       # nmId без строк «Данных» — ярлыки из таблицы
        self.assertEqual((r3["article"], r3["category"]), ("X000000009", fr.category_of("Упаковки для украшений")))
        self.assertEqual(st["nm_without_labels"], [3])
        self.assertEqual([r["nm_id"] for r in st["articles"] if r["ym"] == "2026-10"], [1, 3, 2])   # по убыванию хранения внутри месяца

    def test_month_summary_equals_form_storage_on_the_same_days(self):
        st = fr.build_storage(*self.data())
        m = st["months"]["2026-10"]
        self.assertEqual((m["days"], m["nm_count"], m["unit_days"]), (3, 3, 7))                 # 3 + 3 + 1 (сторно не считается)
        self.assertEqual((m["storage"], m["storage_rounded"], m["form_storage"], m["diff"]), (D("22.50"), D("22.50"), D("22.50"), D("0.00")))
        self.assertEqual((m["form_storage_all"], m["form_days"]), (D("22.50"), 3))
        s = st["months"]["2026-09"]                                                   # в таблице один день сентября, форма — два
        self.assertEqual((s["storage"], s["form_storage"], s["diff"], s["form_storage_all"], s["form_days"]), (D("4.00"), D("4.00"), D("0.00"), D("104.00"), 2))
        self.assertEqual((m["sales"], m["storage_pct"]), (D("800"), D("0.0281")))
        self.assertEqual(st["check"][0], ("2026-09", 1, D("4.00"), D("4.00"), D("0.00"), D("104.00"), 2))
        self.assertEqual(st["total"]["storage"], D("26.50")); self.assertEqual(st["total"]["diff"], D("0.00"))
        self.assertEqual(st["by_day"]["2026-10-08"]["storage"], D("3.50")); self.assertIsNone(st["by_day"]["2026-10-08"]["storage_pct"])

    def test_groups_add_up_to_the_total(self):
        st = fr.build_storage(*self.data())
        for key in ("by_brand", "by_category"):
            self.assertEqual(sum((g["Итого"]["storage"] for g in st[key].values()), D(0)), st["total"]["storage"])
            self.assertEqual(sum((g["Итого"]["unit_days"] for g in st[key].values()), 0), st["total"]["unit_days"])
        self.assertEqual(st["by_brand"]["SOKOLOV"]["2026-10"]["sales"], D("-200"))
        self.assertEqual(set(st["by_category"]), {"Серьги", "Кольца", fr.category_of("Упаковки для украшений")})

    def test_identity_rounds_each_day_to_the_kopeck_like_the_realisation_report(self):
        storage = [srow("2026-10-07", 1, "1.00400"), srow("2026-10-07", 2, "2.00300"), srow("2026-10-08", 1, "3.00400")]   # дни 3,007 и 3,004
        data = [drow("2026-10-07", None, storage="3.01", article=fr.NO_PRODUCT), drow("2026-10-08", None, storage="3.00", article=fr.NO_PRODUCT)]
        m = fr.build_storage(storage, data)["months"]["2026-10"]
        self.assertEqual((m["storage"], m["storage_rounded"], m["form_storage"], m["diff"]), (D("6.01100"), D("6.01"), D("6.01"), D("0.00")))

    def test_empty_table_gives_empty_sheet_parts(self):
        st = fr.build_storage([], [drow("2026-10-07", 1, sales="10", storage="5")])
        self.assertEqual((st["articles"], st["months"], st["check"], st["rows"]), ([], {}, [], 0))
        self.assertEqual(st["total"]["storage"], D(0))


class ReadStorageRowsTests(unittest.TestCase):
    def rows(self):
        out = []
        for day in ("2026-10-07", "2026-10-09"):
            for nm in (1, 2, 3, 4):
                for chrt in (1, 2, 3):
                    out.append(srow(day, nm, "1.00", chrt=chrt))
        return out

    def test_reads_everything_with_small_pages_and_does_not_lose_rows_of_one_nm_id(self):
        sb = FakeSb({fr.STORAGE_TABLE: self.rows()})
        got = fr.read_storage_rows(sb, "2026-10-01", "2026-10-31", page=5)
        key = lambda r: (r["day"], r["nm_id"], r["chrt_id"])
        self.assertEqual(sorted(map(key, got)), sorted(map(key, self.rows())))
        self.assertEqual(len(got), 24)
        self.assertTrue(all(c[0] == fr.STORAGE_TABLE for c in sb.calls))
        self.assertTrue(all(any(k[0] in ("gte", "eq") and k[1] == "day" for k in c[1]) for c in sb.calls))   # каждый запрос — по индексу дня

    def test_window_without_rows_costs_one_request(self):
        sb = FakeSb({fr.STORAGE_TABLE: self.rows()})
        self.assertEqual(fr.read_storage_rows(sb, "2026-11-01", "2026-11-30", page=5), [])
        self.assertEqual(len(sb.calls), 1)

    def test_nm_with_rows_not_less_than_a_page_is_loud(self):
        rows = [srow("2026-10-07", 1, "1.00", chrt=i) for i in range(6)]
        with self.assertRaises(RuntimeError):
            fr.read_storage_rows(FakeSb({fr.STORAGE_TABLE: rows}), "2026-10-01", "2026-10-31", page=5)


class BookSheetTests(unittest.TestCase):
    def test_sheet_is_written_after_wb_data_and_survives_a_module_error(self):
        import openpyxl
        import scripts.report_finrez as book
        days = ["2026-10-07", "2026-10-08"]
        storage, data = BuildStorageTests().data()
        st = fr.build_storage([r for r in storage if r["day"] in days], [r for r in data if r["date"] in days], days=days)
        for parts in (st, {"error": "boom"}, None):
            wb = openpyxl.Workbook(write_only=True)
            n = book.write_wb_storage_sheet(wb, parts, days)
            with tempfile.TemporaryDirectory() as d:
                path = os.path.join(d, "b.xlsx"); wb.save(path)
                ws = openpyxl.load_workbook(path, read_only=True)[book.WB_STORAGE_SHEET]
                cells = [[c for c in row] for row in ws.iter_rows(values_only=True)]
            if parts is st:
                self.assertEqual(n, 2)                                                 # (окт, 1), (окт, 2) — nmId 3 хранится 09.10, вне окна
                self.assertEqual(cells[1][:3], tuple(h for h, _k, _f in fr.STORAGE_MONTH_COLS[:3]) if isinstance(cells[1], tuple) else list(h for h, _k, _f in fr.STORAGE_MONTH_COLS[:3]))
                self.assertEqual(cells[2][0], "окт"); self.assertEqual(cells[3][0], "07.окт"); self.assertEqual(cells[5][0], "Общий итог")
                self.assertAlmostEqual(cells[2][4], 15.5)                              # Σ октября за два дня: 10 + 2 − 1,5 + 5
                self.assertAlmostEqual(cells[2][9], 0.0)                               # разница с формой (колонка J)
                flat = [c for row in cells for c in row]
                self.assertIn("По брендам", flat); self.assertIn("По категориям", flat); self.assertIn("F000283615", flat)
            else:
                self.assertEqual(n, 0); self.assertIn("Лист хранения WB не собран", str(cells[0][0]))


if __name__ == "__main__":
    unittest.main()
